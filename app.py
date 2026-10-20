from __future__ import annotations

import csv
import io
import json
import os
import re
import secrets
import unicodedata
import zipfile
import time
from datetime import timedelta, datetime, date
from functools import wraps
from urllib.parse import urlparse, parse_qsl, urlencode

import qrcode
import requests
from PIL import Image
from qrcode.constants import ERROR_CORRECT_H
from authlib.integrations.flask_client import OAuth
from flask import (
    Flask, Response, abort, current_app, flash, jsonify, redirect,
    render_template, request, send_file, session, url_for
)
from flask_wtf.csrf import CSRFProtect
from sqlalchemy import func, text, DateTime, inspect, or_, and_, select, union_all, literal
from werkzeug.middleware.proxy_fix import ProxyFix
from sqlalchemy.orm import aliased

from config import Config
from models import (
    db, User, OAuthAccount, YouTubeChannel, Playlist, Video, PlaylistVideo, Course, CourseVideo, CoursePlaylist, ClassGroup,
    SchoolClass, Student, ClassEnrollment, CourseClass, Participant, QrCode, QrTargetHistory, ScanLog, WatchSession, WatchEvent,
    HubSubscription, HubEvent, ErrorLog, ExperimentRun, FunctionalTestRun,
    ResearchSetting, SurveyResult, ConnectedGoogleAccount, HomeroomChange, EditHistory, CourseTransfer, utcnow
)
from security import encrypt_json, decrypt_json
from services.youtube import sync_my_channel, sync_video_ids, reconcile_user, sync_all_connected_channels
from services.websub import subscribe_channel, verify_callback, parse_notification

csrf = CSRFProtect()
oauth = OAuth()


# Danh mục nhiệm vụ đối chứng được chuẩn hóa theo đúng thiết kế thực nghiệm:
# chỉ so sánh thủ công - Web App cho nhiệm vụ tạo QR riêng và tạo QR hàng loạt.
RESEARCH_COMPARISON_TASKS = [
    {
        "key": "TAO_QR_RIENG",
        "name": "Tạo và lưu mã QR riêng",
        "manual_task": "Tạo và lưu mã QR cho 01 video theo quy trình thủ công",
    },
    {
        "key": "TAO_QR_HANG_LOAT",
        "name": "Tạo và lưu mã QR hàng loạt",
        "manual_task": "Tạo và lưu mã QR cho một nhóm video/playlist theo quy trình thủ công",
    },
]
RESEARCH_COMPARISON_TASK_MAP = {item["key"]: item for item in RESEARCH_COMPARISON_TASKS}


STUDENT_STATUS_LABELS = {
    "dang_hoc": "Đang học",
    "bao_luu": "Bảo lưu",
    "nghi_hoc": "Nghỉ học",
    "chuyen_truong": "Chuyển trường",
    "chuyen_lop": "Chuyển lớp",
    "khac": "Khác",
}


def normalize_student_code(value) -> str:
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return re.sub(r"\s+", "", str(value or "").strip()).upper()[:64]


def student_code_tail(code) -> str:
    """Dùng đúng ba chữ số cuối khi trình bày, không dùng làm khóa định danh."""
    match = re.search(r"(\d{3})$", str(code or ""))
    return match.group(1) if match else "—"


def enrollment_active_on(enrollment: ClassEnrollment, on_date: date) -> bool:
    return bool(enrollment.started_on <= on_date and (enrollment.ended_on is None or enrollment.ended_on > on_date))


def class_size_at(school_class: SchoolClass, on_date: date | None = None) -> int:
    on_date = on_date or utcnow().date()
    return len({e.student_id for e in school_class.enrollments if enrollment_active_on(e, on_date)})


def student_current_status_in_class(student: Student, school_class: SchoolClass, on_date: date | None = None):
    on_date = on_date or utcnow().date()
    rows = sorted((e for e in student.enrollments if e.school_class_id == school_class.id), key=lambda e: (e.started_on, e.id), reverse=True)
    for e in rows:
        if enrollment_active_on(e, on_date):
            return "dang_hoc", e
    if rows:
        return (rows[0].end_status or "khac"), rows[0]
    return "khac", None


def audit_edit(user, kind, entity_id, field, before, after):
    previous, updated = str(before or ""), str(after or "")
    if previous != updated:
        db.session.add(EditHistory(entity_kind=kind, entity_id=entity_id, field=field,
                                   old_value=previous[:500], new_value=updated[:500],
                                   changed_by=user.id))


def shared_class_by_name(name):
    return SchoolClass.query.filter(
        func.lower(func.trim(SchoolClass.name)) == name.strip().lower(),
        SchoolClass.active.is_(True)).first()


def create_or_find_school_class(name, user, declared_size=None):
    name = re.sub(r"\s+", " ", (name or "").strip())
    if not name or len(name) > 255:
        raise ValueError("Tên lớp không hợp lệ.")
    if db.engine.dialect.name == "postgresql":
        db.session.execute(text("SELECT pg_advisory_xact_lock(48191271)"))
    existing = shared_class_by_name(name)
    if existing:
        return existing, False
    obj = SchoolClass(code=f"TMP-{secrets.token_hex(8)}", name=name,
                      declared_size=declared_size)
    db.session.add(obj)
    db.session.flush()
    assign_auto_school_class_code(obj)
    return obj, True


def class_management_url(user, school_class):
    anchor = f"#lop-{school_class.id}"
    if school_class.homeroom_teacher_id == user.id:
        return url_for("my_classes_page") + anchor
    if db.session.query(CourseClass).join(Course, Course.id == CourseClass.course_id).filter(
        CourseClass.school_class_id == school_class.id, CourseClass.active.is_(True),
        Course.owner_id == user.id, Course.active.is_(True)).first():
        return url_for("teaching_classes_page") + anchor
    return url_for("classes_page") if user.role == "admin" else url_for("teaching_classes_page") + anchor


def can_seed_school_class(user, school_class):
    """GV dạy lớp có thể nhập danh sách BAN ĐẦU, nếu lớp chưa có bất cứ SV nào."""
    if school_class.homeroom_teacher_id == user.id:
        return True
    # GV bộ môn (kể cả tài khoản quản trị đang giảng dạy) chỉ được nhập danh sách ban đầu.
    return (ClassEnrollment.query.filter_by(school_class_id=school_class.id).first() is None and
            db.session.query(CourseClass).join(Course, Course.id == CourseClass.course_id)
            .filter(CourseClass.school_class_id == school_class.id, CourseClass.active.is_(True),
                    Course.owner_id == user.id, Course.active.is_(True)).first() is not None)


def can_manage_school_class(user: User, school_class: SchoolClass) -> bool:
    return bool(user and (user.role == "admin" or school_class.homeroom_teacher_id == user.id))


def active_enrollments_for_student(student: Student, on_date: date | None = None):
    on_date = on_date or utcnow().date()
    return [e for e in student.enrollments if enrollment_active_on(e, on_date)]


def resolve_student_class_for_qr(student_code: str, qr: QrCode, on_date: date | None = None):
    """Xác định lớp của SV theo mã Nhà trường và thời điểm quét.

    Nếu QR thuộc môn học/ mô đun, SV phải đang thuộc một lớp được gán cho môn học/ mô đun đó.
    """
    on_date = on_date or utcnow().date()
    code = normalize_student_code(student_code)
    student = Student.query.filter_by(student_code=code).first()
    if not student:
        return None, None, "Mã sinh viên chưa có trong danh sách lớp của hệ thống. Vui lòng kiểm tra lại mã hoặc liên hệ giảng viên chủ nhiệm."
    enrollments = active_enrollments_for_student(student, on_date)
    if qr.course_id:
        allowed_ids = {link.school_class_id for link in CourseClass.query.filter_by(course_id=qr.course_id, active=True).all()}
        if not allowed_ids:
            return student, None, "Môn học/ mô đun này chưa được giảng viên gán lớp áp dụng. Vui lòng liên hệ giảng viên phụ trách."
        enrollments = [e for e in enrollments if e.school_class_id in allowed_ids]
    if not enrollments:
        return student, None, "Mã sinh viên không thuộc lớp đang áp dụng cho môn học/ mô đun tại thời điểm hiện tại. Vui lòng liên hệ giảng viên phụ trách."
    if len(enrollments) > 1:
        return student, None, "Hệ thống phát hiện mã sinh viên đang đồng thời thuộc nhiều lớp phù hợp. Vui lòng liên hệ giảng viên chủ nhiệm để kiểm tra dữ liệu lớp."
    return student, enrollments[0].school_class, None



def eligible_classes_for_qr(qr: QrCode):
    """Chỉ cho người quét chọn lớp đang áp dụng cho môn học/mô đun của QR."""
    query = SchoolClass.query.filter(SchoolClass.active.is_(True))
    if qr.course_id:
        query = query.join(CourseClass, CourseClass.school_class_id == SchoolClass.id).filter(
            CourseClass.course_id == qr.course_id,
            CourseClass.active.is_(True),
        )
    return query.distinct().order_by(SchoolClass.name, SchoolClass.id).all()


def current_students_in_class(school_class: SchoolClass, on_date: date | None = None):
    """Danh sách SV đang học, không lấy SV đã bảo lưu/nghỉ/chuyển lớp."""
    on_date = on_date or utcnow().date()
    rows = (Student.query.join(ClassEnrollment, ClassEnrollment.student_id == Student.id)
            .filter(ClassEnrollment.school_class_id == school_class.id,
                    ClassEnrollment.started_on <= on_date,
                    or_(ClassEnrollment.ended_on.is_(None), ClassEnrollment.ended_on > on_date))
            .distinct().all())
    return sorted(rows, key=lambda item: ((item.full_name or "").casefold(), item.student_code))


def validate_student_selection_for_qr(qr: QrCode, school_class_id: int, student_id: int, class_name: str):
    """Đối chiếu lựa chọn với lớp, liên kết môn học và sĩ số tại thời điểm hiện tại."""
    school_class = db.session.get(SchoolClass, school_class_id)
    if not school_class or school_class.id not in {c.id for c in eligible_classes_for_qr(qr)}:
        return None, None, "Lớp không thuộc môn học/mô đun của mã QR này. Vui lòng chọn lại lớp."
    normalized = re.sub(r"\s+", " ", (class_name or "").strip()).casefold()
    if normalized != re.sub(r"\s+", " ", school_class.name.strip()).casefold():
        return None, None, "Tên lớp chưa khớp với lớp đã chọn. Vui lòng chọn lớp từ danh sách gợi ý."
    student = db.session.get(Student, student_id)
    if not student or student.id not in {row.id for row in current_students_in_class(school_class)}:
        return None, None, "Sinh viên không còn trong danh sách đang học của lớp đã chọn. Vui lòng kiểm tra lại."
    return student, school_class, None


def role_label(role: str) -> str:
    return "Quản trị viên" if role == "admin" else "Giảng viên"



def migrate_legacy_classes():
    """Chuyển dữ liệu lớp cũ theo môn học/ mô đun sang danh mục lớp dùng chung, không xóa dữ liệu cũ."""
    try:
        for old in ClassGroup.query.order_by(ClassGroup.id).all():
            name = (old.name or "").strip()
            if not name:
                continue
            school_class = SchoolClass.query.filter(func.lower(SchoolClass.name) == name.lower()).first()
            if not school_class:
                code = (old.code or f"L{old.id:06d}").strip()
                if SchoolClass.query.filter_by(code=code).first():
                    code = f"LOP{old.id:06d}"
                school_class = SchoolClass(code=code, name=name, active=bool(getattr(old, "active", True)))
                db.session.add(school_class)
                db.session.flush()
            link = CourseClass.query.filter_by(course_id=old.course_id, school_class_id=school_class.id).first()
            if not link:
                db.session.add(CourseClass(course_id=old.course_id, school_class_id=school_class.id, active=bool(getattr(old, "active", True))))
        db.session.commit()
    except Exception:
        db.session.rollback()


def create_app():
    app = Flask(__name__)
    app.config.from_object(Config)
    app.wsgi_app = ProxyFix(app.wsgi_app, x_for=1, x_proto=1, x_host=1)

    db.init_app(app)
    csrf.init_app(app)
    oauth.init_app(app)

    if app.config["GOOGLE_CLIENT_ID"] and app.config["GOOGLE_CLIENT_SECRET"]:
        oauth.register(
            name="google",
            client_id=app.config["GOOGLE_CLIENT_ID"],
            client_secret=app.config["GOOGLE_CLIENT_SECRET"],
            server_metadata_url="https://accounts.google.com/.well-known/openid-configuration",
            client_kwargs={"scope": app.config["GOOGLE_SCOPES"]},
        )

    if app.config["GOOGLE_CLIENT_ID"] and app.config["GOOGLE_CLIENT_SECRET"]:
        oauth.register(
            name="google_additional", client_id=app.config["GOOGLE_CLIENT_ID"],
            client_secret=app.config["GOOGLE_CLIENT_SECRET"],
            server_metadata_url="https://accounts.google.com/.well-known/openid-configuration",
            client_kwargs={"scope": app.config["GOOGLE_SCOPES"]},
        )

    register_routes(app)

    with app.app_context():
        if app.config["AUTO_CREATE_DB"]:
            db.create_all()
            ensure_schema_upgrades()
            migrate_legacy_classes()

    return app



def ensure_schema_upgrades():
    """Nâng cấp nhẹ schema khi deploy mà không xóa PostgreSQL hiện có.

    db.create_all() không tự thêm cột vào bảng đã tồn tại. v2.6 tiếp tục giữ nguyên
    dữ liệu cũ, bổ sung khoa/bộ môn (nếu thiếu) và trạng thái hoạt động cho lớp để
    có thể "xóa" lớp khỏi giao diện mà không làm mất dữ liệu lịch sử.
    """
    inspector = inspect(db.engine)
    tables = set(inspector.get_table_names())
    with db.engine.begin() as conn:
        if "users" in tables:
            user_columns = {c["name"] for c in inspector.get_columns("users")}
            if "lecturer_code" not in user_columns:
                conn.execute(text("ALTER TABLE users ADD COLUMN lecturer_code VARCHAR(32)"))
                try:
                    conn.execute(text("CREATE UNIQUE INDEX IF NOT EXISTS ix_users_lecturer_code ON users (lecturer_code)"))
                except Exception:
                    pass
            if "professional_group" not in user_columns:
                conn.execute(text("ALTER TABLE users ADD COLUMN professional_group VARCHAR(16)"))
            if "profile_completed" not in user_columns:
                conn.execute(text("ALTER TABLE users ADD COLUMN profile_completed BOOLEAN NOT NULL DEFAULT FALSE"))
        if "youtube_channels" in tables:
            channel_columns = {c["name"] for c in inspector.get_columns("youtube_channels")}
            if "connected_account_id" not in channel_columns:
                conn.execute(text("ALTER TABLE youtube_channels ADD COLUMN connected_account_id INTEGER"))
        if "school_classes" in tables:
            class_columns = {c["name"] for c in inspector.get_columns("school_classes")}
            if "declared_size" not in class_columns:
                conn.execute(text("ALTER TABLE school_classes ADD COLUMN declared_size INTEGER"))
        if "courses" in tables:
            columns = {c["name"] for c in inspector.get_columns("courses")}
            if "faculty_name" not in columns:
                conn.execute(text("ALTER TABLE courses ADD COLUMN faculty_name VARCHAR(255)"))
            if "department_name" not in columns:
                conn.execute(text("ALTER TABLE courses ADD COLUMN department_name VARCHAR(255)"))
            conn.execute(text("""
                UPDATE courses
                SET faculty_name = unit_name
                WHERE (faculty_name IS NULL OR faculty_name = '')
                  AND unit_name IS NOT NULL AND unit_name <> ''
            """))
        if "videos" in tables:
            video_columns = {c["name"] for c in inspector.get_columns("videos")}
            if "uploader_channel_id" not in video_columns:
                # Không tự suy đoán rằng video cũ thuộc kênh đang gắn với bản ghi.
                # Lần đồng bộ tiếp theo sẽ kiểm chứng bằng YouTube API.
                conn.execute(text("ALTER TABLE videos ADD COLUMN uploader_channel_id VARCHAR(64)"))
                conn.execute(text("CREATE INDEX IF NOT EXISTS ix_videos_uploader_channel_id ON videos (uploader_channel_id)"))
        if "playlists" in tables:
            playlist_columns = {c["name"] for c in inspector.get_columns("playlists")}
            if "is_available" not in playlist_columns:
                conn.execute(text("ALTER TABLE playlists ADD COLUMN is_available BOOLEAN NOT NULL DEFAULT TRUE"))
        if "qr_codes" in tables:
            qr_columns = {c["name"] for c in inspector.get_columns("qr_codes")}
            if "target_type" not in qr_columns:
                conn.execute(text("ALTER TABLE qr_codes ADD COLUMN target_type VARCHAR(16) NOT NULL DEFAULT 'video'"))
            if "playlist_id" not in qr_columns:
                conn.execute(text("ALTER TABLE qr_codes ADD COLUMN playlist_id INTEGER"))
            conn.execute(text("UPDATE qr_codes SET target_type = 'video' WHERE target_type IS NULL OR target_type = ''"))
        if "course_playlists" in tables:
            cp_columns = {c["name"] for c in inspector.get_columns("course_playlists")}
            if "active" not in cp_columns:
                conn.execute(text("ALTER TABLE course_playlists ADD COLUMN active BOOLEAN NOT NULL DEFAULT TRUE"))
        if "course_videos" in tables:
            cv_columns = {c["name"] for c in inspector.get_columns("course_videos")}
            if "manual_added" not in cv_columns:
                # Dữ liệu cũ được coi là gán thủ công để không vô tình bị loại bỏ khi nâng cấp.
                conn.execute(text("ALTER TABLE course_videos ADD COLUMN manual_added BOOLEAN NOT NULL DEFAULT TRUE"))
        if "class_groups" in tables:
            class_columns = {c["name"] for c in inspector.get_columns("class_groups")}
            if "active" not in class_columns:
                conn.execute(text("ALTER TABLE class_groups ADD COLUMN active BOOLEAN NOT NULL DEFAULT TRUE"))
        if "error_logs" in tables:
            error_columns = {c["name"] for c in inspector.get_columns("error_logs")}
            if "action" not in error_columns:
                conn.execute(text("ALTER TABLE error_logs ADD COLUMN action VARCHAR(128)"))
            if "course_module" not in error_columns:
                conn.execute(text("ALTER TABLE error_logs ADD COLUMN course_module VARCHAR(255)"))
            if "actor_user_id" not in error_columns:
                conn.execute(text("ALTER TABLE error_logs ADD COLUMN actor_user_id INTEGER"))
            if "actor_student_code" not in error_columns:
                conn.execute(text("ALTER TABLE error_logs ADD COLUMN actor_student_code VARCHAR(64)"))
            if "actor_class_code" not in error_columns:
                conn.execute(text("ALTER TABLE error_logs ADD COLUMN actor_class_code VARCHAR(64)"))
        # Chỉ mục phục vụ thống kê phân trang và ghép lượt quét với phiên xem.
        # Không thay đổi dữ liệu cũ; các chỉ mục chỉ được tạo nếu chưa có.
        if "scan_logs" in tables:
            conn.execute(text("CREATE INDEX IF NOT EXISTS ix_scan_logs_match_viewer ON scan_logs (qr_id, participant_code, class_code, scanned_at)"))
        if "watch_sessions" in tables:
            conn.execute(text("CREATE INDEX IF NOT EXISTS ix_watch_sessions_match_viewer ON watch_sessions (qr_id, participant_code, class_code, started_at)"))
        if "functional_test_runs" in tables:
            test_columns = {c["name"] for c in inspector.get_columns("functional_test_runs")}
            if "source" not in test_columns:
                conn.execute(text("ALTER TABLE functional_test_runs ADD COLUMN source VARCHAR(16) NOT NULL DEFAULT 'manual'"))
            if "participant_code" not in test_columns:
                conn.execute(text("ALTER TABLE functional_test_runs ADD COLUMN participant_code VARCHAR(128)"))
            if "participant_class_code" not in test_columns:
                conn.execute(text("ALTER TABLE functional_test_runs ADD COLUMN participant_class_code VARCHAR(64)"))
            conn.execute(text("CREATE INDEX IF NOT EXISTS ix_functional_test_runs_participant_code ON functional_test_runs (participant_code)"))
            # Cho phép ghi lỗi OAuth trước khi xác định được user. PostgreSQL hỗ trợ trực tiếp.
            if db.engine.dialect.name == "postgresql":
                conn.execute(text("ALTER TABLE functional_test_runs ALTER COLUMN owner_id DROP NOT NULL"))

def record_functional_test(case_code, case_name, passed, owner_id=None, note="", source="automatic", participant_code=None, participant_class_code=None):
    """Ghi bằng chứng kiểm thử phát sinh từ hoạt động thật của hệ thống.

    Hàm chỉ thêm bản ghi vào session hiện tại; caller chủ động commit cùng giao dịch nghiệp vụ.
    """
    db.session.add(FunctionalTestRun(
        owner_id=owner_id,
        case_code=(case_code or "AUTO")[:64],
        case_name=(case_name or "Kiểm thử tự động")[:255],
        passed=bool(passed),
        source=(source or "automatic")[:16],
        note=(note or "")[:4000],
        participant_code=(participant_code or None),
        participant_class_code=(participant_class_code or None),
    ))


def record_webapp_experiment(owner_id, task_key, elapsed_seconds, operations_count=1, success=True, note="", processing_seconds=None):
    """Ghi phép đo Web App tự động; dữ liệu đối chứng thủ công vẫn do quản trị viên nhập."""
    db.session.add(ExperimentRun(
        owner_id=owner_id,
        task_key=(task_key or "WEBAPP")[:128],
        mode="webapp",
        operations_count=max(0, int(operations_count or 0)),
        elapsed_seconds=max(0.0, float(elapsed_seconds or 0)),
        system_processing_seconds=(float(processing_seconds) if processing_seconds is not None else None),
        success=bool(success),
        note=(note or "")[:4000],
    ))


def research_setting(key, default=""):
    row = ResearchSetting.query.filter_by(key=key).first()
    return row.value if row and row.value else default


def save_research_setting(key, value):
    row = ResearchSetting.query.filter_by(key=key).first()
    if not row:
        row = ResearchSetting(key=key)
        db.session.add(row)
    row.value = (value or "").strip()


def login_required(fn):
    @wraps(fn)
    def wrapper(*args, **kwargs):
        if not session.get("user_id"):
            return redirect(url_for("login", next=request.path))
        return fn(*args, **kwargs)
    return wrapper


def current_user():
    uid = session.get("user_id")
    return db.session.get(User, uid) if uid else None


def allocate_lecturer_code(user: User) -> str:
    """Cấp mã GV tuần tự GV0001, GV0002... độc lập với ID tài khoản.

    Sau khi quản trị viên xóa toàn bộ dữ liệu, danh sách mã GV trống nên lần cấp
    tiếp theo tự bắt đầu lại từ GV0001 mà không cần đăng xuất/đăng nhập lại.
    """
    if user.lecturer_code:
        return user.lecturer_code
    # Khóa riêng cho cấp mã trên PostgreSQL để tránh hai tài khoản cùng nhận một mã.
    if db.engine.dialect.name == "postgresql":
        db.session.execute(text("SELECT pg_advisory_xact_lock(48191270)"))
    used = []
    for value, in db.session.query(User.lecturer_code).filter(User.lecturer_code.isnot(None)).all():
        m = re.fullmatch(r"GV(\d{4,})", (value or "").strip().upper())
        if m:
            used.append(int(m.group(1)))
    n = 1
    used_set = set(used)
    while n in used_set:
        n += 1
    code = f"GV{n:04d}"
    user.lecturer_code = code
    return code


def safe_next(value):
    if not value:
        return None
    parsed = urlparse(value)
    if parsed.scheme or parsed.netloc:
        return None
    if not value.startswith("/"):
        return None
    return value


def base_url():
    return current_app.config.get("BASE_URL") or request.url_root.rstrip("/")


def qr_public_url(qr):
    return f"{base_url()}/q/{qr.slug}"


def active_video_for_qr(qr):
    target = qr.current_target
    return target.video if target else None


def owner_filter(model):
    user = current_user()
    if user.role == "admin":
        return model.query
    if hasattr(model, "owner_id"):
        return model.query.filter_by(owner_id=user.id)
    return model.query


def safe_filename(value: str, fallback: str = "ma_qr") -> str:
    """Tạo tên tệp an toàn nhưng GIỮ NGUYÊN tiếng Việt có dấu.

    Chỉ loại các ký tự không hợp lệ trong tên tệp Windows/macOS và ký tự điều khiển;
    không phiên âm/khử dấu. Nhờ đó tên tệp QR tải về đúng tên video YouTube.
    """
    value = unicodedata.normalize("NFC", (value or "").strip())
    value = re.sub(r'[<>:"/\\|?*\x00-\x1F]', "_", value)
    value = re.sub(r"\s+", " ", value).strip(" .")
    return (value[:180] or fallback)




def assign_auto_course_code(course: Course) -> str:
    """Sinh mã môn học/ mô đun tự động, ổn định và không yêu cầu GV nhập.

    Dùng khóa chính của CSDL để mã không trùng ngay cả khi nhiều giảng viên thao tác
    đồng thời. Mã chỉ phục vụ nhận diện nội bộ; người dùng chủ yếu làm việc bằng tên môn học/ mô đun.
    """
    if not course.id:
        db.session.flush()
    course.code = f"HP{course.id:06d}"
    return course.code


def assign_auto_class_code(class_group: ClassGroup) -> str:
    """Sinh mã lớp tự động; GV chỉ cần nhập tên lớp."""
    if not class_group.id:
        db.session.flush()
    class_group.code = f"L{class_group.id:06d}"
    return class_group.code


def assign_auto_school_class_code(school_class: SchoolClass) -> str:
    """Sinh mã lớp dùng chung tự động; quản trị viên chỉ nhập tên lớp và giao GVCN."""
    if not school_class.id:
        db.session.flush()
    school_class.code = f"L{school_class.id:06d}"
    return school_class.code


def video_is_usable(video: Video | None) -> bool:
    return bool(video and video.is_available and video.embeddable is not False and video.privacy_status != "private")


def video_is_from_connected_channel(video: Video | None) -> bool:
    """Quyền sở hữu phải được xác nhận qua YouTube API, không qua playlist."""
    return bool(video and video.channel and video.uploader_channel_id
                and video.uploader_channel_id == video.channel.channel_id)


def course_owned_by_user(course: Course, user: User) -> bool:
    return bool(course and user and course.owner_id == user.id)


def managed_qr_scope(user_id: int):
    """Phân quyền theo GV hiện phụ trách môn, KHÔNG theo tác giả QR ban đầu.

    QR không gắn môn vẫn do người tạo quản lý. Điều kiện dùng thống nhất cho
    tải PNG/ZIP, sửa/xóa QR, thống kê quét/xem và Dashboard.
    """
    owned_courses = db.session.query(Course.id).filter(Course.owner_id == user_id)
    return or_(
        and_(QrCode.course_id.is_(None), QrCode.owner_id == user_id),
        and_(QrCode.course_id.isnot(None), QrCode.course_id.in_(owned_courses)),
    )


def qr_managed_by_user(qr: QrCode, user: User) -> bool:
    if not qr or not user:
        return False
    if qr.course_id is not None:
        return bool(qr.course and qr.course.owner_id == user.id)
    return qr.owner_id == user.id


def visible_video_query(user: User):
    """Chỉ trả về video đã xác minh do kênh kết nối của chính GV đăng tải.

    Việc một video xuất hiện trong playlist không chứng minh GV sở hữu video.
    Những video nhập sai từ bản cũ được giữ trong CSDL để không mất lịch sử QR,
    nhưng không được đưa vào Kho video hoặc danh sách tạo QR mới.
    """
    return (Video.query.join(YouTubeChannel)
            .filter(YouTubeChannel.user_id == user.id,
                    Video.uploader_channel_id == YouTubeChannel.channel_id))


def visible_playlist_query(user: User):
    """Playlist có thể gán vào môn học/ mô đun của tài khoản đang đăng nhập.

    Chỉ lấy playlist do chính kênh YouTube của người dùng sở hữu và còn tồn tại.
    Playlist Uploads hệ thống của YouTube không đưa vào danh sách gán môn học/ mô đun.
    """
    q = Playlist.query.join(YouTubeChannel).filter(YouTubeChannel.user_id == user.id)
    q = q.filter(Playlist.is_available.is_(True))
    q = q.filter((YouTubeChannel.uploads_playlist_id.is_(None)) | (Playlist.youtube_playlist_id != YouTubeChannel.uploads_playlist_id))
    return q


def videos_library_return(anchor: str, raw_query: str = "") -> str:
    """Redirect only within the video library, retaining approved filter parameters.

    A stable row anchor stops the browser jumping back to the beginning after POST.
    """
    allowed = {"q", "status", "assignment", "playlist_id", "p_q", "p_assignment", "p_page", "v_page"}
    query = urlencode([(key, value[:150]) for key, value in parse_qsl((raw_query or "")[:1200])
                       if key in allowed and value
                       and not ((anchor.startswith("playlist-") and key == "p_assignment") or
                                (anchor.startswith("video-") and key == "assignment"))], doseq=True)
    target = url_for("videos_page") + (("?" + query) if query else "")
    return target + "#" + anchor


def ensure_course_playlist_link(course: Course, playlist: Playlist, user: User):
    link = CoursePlaylist.query.filter_by(course_id=course.id, playlist_id=playlist.id).first()
    if link:
        was_reactivated = not link.active
        link.active = True
        if not link.added_by:
            link.added_by = user.id
        return link, was_reactivated
    link = CoursePlaylist(course_id=course.id, playlist_id=playlist.id, added_by=user.id, active=True)
    db.session.add(link)
    db.session.flush()
    return link, True


def create_qr_for_course_playlist(course: Course, playlist: Playlist, user: User, link: CoursePlaylist | None = None):
    link = link or CoursePlaylist.query.filter_by(course_id=course.id, playlist_id=playlist.id, active=True).first()
    if link and link.qr_id:
        existing = db.session.get(QrCode, link.qr_id)
        if existing and existing.active:
            existing.title = (playlist.title or "Playlist")[:255]
            existing.target_type = "playlist"
            existing.playlist_id = playlist.id
            return existing, False
    existing = (QrCode.query
                .filter_by(course_id=course.id, playlist_id=playlist.id, active=True)
                .filter(QrCode.target_type == "playlist")
                .order_by(QrCode.created_at.desc())
                .first())
    if existing:
        if link and link.qr_id != existing.id:
            link.qr_id = existing.id
        existing.title = (playlist.title or "Playlist")[:255]
        return existing, False
    if playlist.channel.user_id != user.id:
        raise ValueError("Playlist thuộc tài khoản YouTube của giảng viên trước. Mã QR cũ vẫn dùng được, nhưng chỉ được tạo mã mới từ playlist do giảng viên hiện tại kết nối.")
    t0 = time.perf_counter()
    slug = secrets.token_urlsafe(8).replace("-", "").replace("_", "")[:12]
    qr = QrCode(
        owner_id=user.id, course_id=course.id, slug=slug,
        title=(playlist.title or "Playlist")[:255], target_type="playlist", playlist_id=playlist.id,
    )
    db.session.add(qr)
    db.session.flush()
    if link:
        link.qr_id = qr.id
    db.session.flush()
    qr.processing_seconds = round(time.perf_counter() - t0, 4)
    return qr, True


def active_qr_for_course_video(course_id: int, video_id: int, owner_id: int):
    # owner_id giữ lại trong chữ ký tương thích; QR cũ do GV trước tạo vẫn dùng.
    return (QrCode.query
            .join(QrTargetHistory, QrTargetHistory.qr_id == QrCode.id)
            .filter(QrCode.course_id == course_id, QrCode.active.is_(True),
                    QrTargetHistory.video_id == video_id, QrTargetHistory.active_to.is_(None))
            .order_by(QrCode.created_at.desc())
            .first())


def ensure_course_video_link(course: Course, video: Video, user: User, source_playlist_id=None, manual=True):
    """Bảo đảm video thuộc môn học/ mô đun.

    manual=True khi GV chọn video trực tiếp; manual=False khi video được kế thừa
    tự động từ một playlist đã gán cho môn học/ mô đun. Cờ này giúp bỏ playlist mà không
    làm mất những video GV vẫn muốn giữ riêng trong môn học/ mô đun.
    """
    link = CourseVideo.query.filter_by(course_id=course.id, video_id=video.id).first()
    if link:
        if manual:
            link.manual_added = True
        elif not link.source_playlist_id and source_playlist_id:
            link.source_playlist_id = source_playlist_id
        return link, False
    link = CourseVideo(
        course_id=course.id,
        video_id=video.id,
        added_by=user.id,
        source_playlist_id=source_playlist_id,
        manual_added=bool(manual),
    )
    db.session.add(link)
    db.session.flush()
    return link, True


def reconcile_course_videos_from_playlists(course: Course, user: User):
    """Đồng bộ danh sách video của môn học/ mô đun theo các playlist đang được gán.

    - Video mới xuất hiện trong bất kỳ playlist đã gán nào sẽ tự động được thêm.
    - Khi bỏ playlist hoặc video bị gỡ khỏi playlist, liên kết tự động sẽ bị bỏ nếu
      video không còn thuộc playlist nào khác của môn học/ mô đun và chưa từng được GV gán riêng.
    - QR video đã phát hành không bị xóa; chỉ liên kết video-môn học/ mô đun được điều chỉnh.
    """
    active_links = CoursePlaylist.query.filter_by(course_id=course.id, active=True).all()
    playlist_ids = [x.playlist_id for x in active_links if x.playlist_id]
    desired = {}  # video_id -> một playlist nguồn đại diện
    if playlist_ids:
        rows = (db.session.query(PlaylistVideo.video_id, PlaylistVideo.playlist_id)
                .filter(PlaylistVideo.playlist_id.in_(playlist_ids)).all())
        for video_id, playlist_id in rows:
            desired.setdefault(video_id, playlist_id)

    existing = {x.video_id: x for x in CourseVideo.query.filter_by(course_id=course.id).all()}
    added = removed = 0
    for video_id, playlist_id in desired.items():
        if video_id not in existing:
            video = db.session.get(Video, video_id)
            if video:
                ensure_course_video_link(course, video, user, source_playlist_id=playlist_id, manual=False)
                added += 1
        else:
            link = existing[video_id]
            if not link.source_playlist_id:
                link.source_playlist_id = playlist_id

    for video_id, link in list(existing.items()):
        if video_id not in desired and not bool(getattr(link, "manual_added", True)):
            db.session.delete(link)
            removed += 1

    db.session.flush()
    return added, removed


def set_playlist_courses(playlist: Playlist, selected_course_ids: set[int], user: User):
    """Đặt lại toàn bộ các môn học/ mô đun áp dụng cho một playlist theo lựa chọn hiện tại."""
    owned_courses = Course.query.filter_by(owner_id=user.id, active=True).all()
    owned_by_id = {c.id: c for c in owned_courses}
    selected_course_ids = {cid for cid in selected_course_ids if cid in owned_by_id}
    existing = {x.course_id: x for x in CoursePlaylist.query.filter_by(playlist_id=playlist.id).all()
                if x.course and x.course.owner_id == user.id}
    affected = set()
    activated = deactivated = 0

    for course_id, course in owned_by_id.items():
        link = existing.get(course_id)
        should_be_active = course_id in selected_course_ids
        if should_be_active:
            if link:
                if not link.active:
                    link.active = True
                    activated += 1
                if not link.added_by:
                    link.added_by = user.id
            else:
                db.session.add(CoursePlaylist(course_id=course_id, playlist_id=playlist.id, added_by=user.id, active=True))
                activated += 1
            affected.add(course_id)
        elif link and link.active:
            link.active = False
            deactivated += 1
            affected.add(course_id)

    db.session.flush()
    auto_added = auto_removed = 0
    for course_id in affected:
        course = owned_by_id[course_id]
        a, r = reconcile_course_videos_from_playlists(course, user)
        auto_added += a
        auto_removed += r
    return activated, deactivated, auto_added, auto_removed


def create_qr_for_course_video(course: Course, video: Video, user: User, link: CourseVideo | None = None):
    link = link or CourseVideo.query.filter_by(course_id=course.id, video_id=video.id).first()
    if link and link.qr_id:
        existing = db.session.get(QrCode, link.qr_id)
        if existing and existing.active:
            # Luôn đồng bộ tên mã QR theo đúng tên video YouTube, giữ nguyên tiếng Việt có dấu.
            existing.title = (video.title or "Video")[:255]
            return existing, False
    existing = active_qr_for_course_video(course.id, video.id, user.id)
    if existing:
        if link and link.qr_id != existing.id:
            link.qr_id = existing.id
        existing.title = (video.title or "Video")[:255]
        return existing, False
    if video.channel.user_id != user.id:
        raise ValueError("Video thuộc kênh của giảng viên trước. QR đã phát hành vẫn hoạt động; mã QR mới chỉ được tạo từ kênh giảng viên đang phụ trách.")
    t0 = time.perf_counter()
    slug = secrets.token_urlsafe(8).replace("-", "").replace("_", "")[:12]
    qr = QrCode(owner_id=user.id, course_id=course.id, slug=slug, title=(video.title or "Video")[:255])
    db.session.add(qr)
    db.session.flush()
    db.session.add(QrTargetHistory(qr_id=qr.id, video_id=video.id, changed_by=user.id))
    if link:
        link.qr_id = qr.id
    db.session.flush()
    qr.processing_seconds = round(time.perf_counter() - t0, 4)
    return qr, True


def generate_branded_qr(qr: QrCode):
    """Tao QR thuong hieu theo mau chinh thuc cua nha truong.

    Mau chot: ma den - nen trang, muc sua loi H, logo Truong Cao dang
    Son La dat giua tren mot dia trang nho. URL van la URL QR dong,
    vi vay co the doi video dich ma khong doi hinh QR da phat hanh.
    """
    builder = qrcode.QRCode(
        version=None,
        error_correction=ERROR_CORRECT_H,
        box_size=int(current_app.config.get("QR_BOX_SIZE", 14)),
        border=int(current_app.config.get("QR_BORDER", 4)),
    )
    builder.add_data(qr_public_url(qr))
    builder.make(fit=True)
    img = builder.make_image(
        fill_color=current_app.config.get("QR_FILL_COLOR", "#000000"),
        back_color=current_app.config.get("QR_BACKGROUND_COLOR", "#FFFFFF"),
    ).convert("RGBA")

    logo_path = os.path.join(
        current_app.static_folder,
        current_app.config.get("BRAND_LOGO", "branding/logo_truong.png"),
    )
    if os.path.exists(logo_path):
        logo = Image.open(logo_path).convert("RGBA")
        alpha_bbox = logo.getchannel("A").getbbox()
        if alpha_bbox:
            logo = logo.crop(alpha_bbox)

        logo_ratio = float(current_app.config.get("QR_LOGO_RATIO", 0.23))
        logo_side = max(48, int(img.size[0] * logo_ratio))
        logo.thumbnail((logo_side, logo_side), Image.Resampling.LANCZOS)

        # Vong trang bao quanh logo giup mau trung tam ro rang va tang do on dinh khi quet.
        pad_ratio = float(current_app.config.get("QR_LOGO_WHITE_PAD_RATIO", 0.012))
        pad = max(6, int(img.size[0] * pad_ratio))
        badge_side = max(logo.width, logo.height) + pad * 2
        badge = Image.new("RGBA", (badge_side, badge_side), (255, 255, 255, 0))
        from PIL import ImageDraw
        draw = ImageDraw.Draw(badge)
        draw.ellipse((0, 0, badge_side - 1, badge_side - 1), fill=(255, 255, 255, 255))
        lx = (badge_side - logo.width) // 2
        ly = (badge_side - logo.height) // 2
        badge.alpha_composite(logo, (lx, ly))

        x = (img.width - badge.width) // 2
        y = (img.height - badge.height) // 2
        img.alpha_composite(badge, (x, y))

    return img


def qr_download_name(qr: QrCode) -> str:
    # Tên tệp QR = đúng tên video hoặc playlist, giữ nguyên tiếng Việt có dấu.
    if qr.is_playlist:
        source_title = qr.playlist.title if qr.playlist and qr.playlist.title else qr.title
    else:
        target = qr.current_target
        source_title = target.video.title if target and target.video and target.video.title else qr.title
    return f"{safe_filename(source_title, qr.slug)}.png"


def register_routes(app):
    @app.errorhandler(403)
    def handle_forbidden(error):
        if request.path.startswith("/api/"):
            return jsonify({"error": "khong_co_quyen_truy_cap"}), 403
        if current_user():
            flash("Bạn không có quyền truy cập chức năng này.", "warning")
            return redirect(url_for("index"))
        return redirect(url_for("login", next=request.path))

    @app.errorhandler(404)
    def handle_not_found(error):
        if request.path.startswith("/api/"):
            return jsonify({"error": "khong_tim_thay_du_lieu"}), 404
        return render_template("http_error.html", title="Không tìm thấy trang", message="Trang hoặc dữ liệu bạn yêu cầu không tồn tại."), 404

    @app.errorhandler(500)
    def handle_unexpected_error(error):
        """Ghi nhật ký lỗi thật kèm người thao tác khi đã được định danh.

        Ngoại lệ không có phiên đăng nhập (ví dụ OAuth trước khi hoàn thành) để trống
        thông tin người dùng thay vì gán nhầm cho một GV/SV khác.
        """
        try:
            db.session.rollback()
            uid = session.get("user_id")
            user = db.session.get(User, uid) if uid else None
            is_student = session.get("participant_role") == "SV" and not user
            code = session.get("participant_code") if is_student else None
            class_code = session.get("participant_class_code") if is_student else None
            db.session.add(ErrorLog(
                module="viewer" if request.path.startswith("/q/") or request.path.startswith("/identify") else "system",
                action="Truy cập học liệu" if is_student else "Thao tác hệ thống",
                course_module="Hệ thống", message="Lỗi xử lý yêu cầu",
                details=str(getattr(error, "original_exception", None) or error)[:2000],
                actor_user_id=(user.id if user else None),
                actor_student_code=(str(code)[:64] if code else None),
                actor_class_code=(str(class_code)[:64] if class_code else None),
            ))
            db.session.commit()
        except Exception:
            db.session.rollback()
        if request.path.startswith("/api/"):
            return jsonify({"error": "loi_he_thong"}), 500
        return render_template("http_error.html", title="Có lỗi hệ thống",
                               message="Hệ thống gặp sự cố. Vui lòng thử lại sau."), 500

    @app.context_processor
    def inject_globals():
        endpoint = request.endpoint or ""
        section_map = {
            "youtube_page": "youtube", "youtube_sync": "youtube", "youtube_websub_subscribe": "youtube", "connect_another_youtube":"youtube", "disconnect_youtube_account":"youtube",
            "videos_page": "videos", "assign_videos_to_course": "videos", "assign_playlists_from_library": "videos",
            "courses_page": "courses", "add_class": "courses", "set_course_school_classes": "courses", "update_course_info": "courses",
            "assign_playlists_to_course": "courses", "delete_course_playlist": "courses",
            "create_course_qrs_bulk": "courses", "create_course_playlist_qrs_bulk": "courses", "download_course_qrs": "courses",
            "qr_page": "qr", "create_playlist_qr": "qr", "qr_retarget": "qr", "qr_png": "qr", "qr_download_all": "qr", "delete_all_qrs": "qr",
            "analytics": "analytics", "export_scans": "analytics", "export_watch": "analytics",
            "classes_page": "class_admin", "my_classes_page": "homeroom", "teaching_classes_page":"teaching_classes", "create_school_class": "class_admin", "assign_homeroom_teacher": "class_admin",
            "upload_class_students": "homeroom", "add_class_student": "homeroom", "change_student_status": "homeroom", "delete_class_student": "homeroom",
            "lecturers_page": "lecturers", "assign_lecturer_code_admin": "lecturers", "update_lecturer_professional_group": "lecturers", "delete_lecturer_account": "lecturers",
            "research_page": "research",
            "error_logs_page": "errors",
            "data_management": "data_admin",
            "update_display_name": "", "edit_student_name":"homeroom", "add_school_class_to_course":"courses", "rename_shared_catalog":"lecturers", "reset_lecturer_data": "lecturers", "profile_setup": "", "sync_wait": "",
            "backup_zip": "data_admin",
            "restore_backup": "data_admin",
        }
        return {
            "me": current_user(),
            "active_section": section_map.get(endpoint, ""),
            "heartbeat_seconds": current_app.config["HEARTBEAT_SECONDS"],
            "completion_threshold": current_app.config["COMPLETION_THRESHOLD"],
            "app_name": current_app.config["APP_NAME"],
            "app_tagline": current_app.config["APP_TAGLINE"],
            "school_name": current_app.config["SCHOOL_NAME"],
            "authors": current_app.config["AUTHORS"],
            "brand_logo": current_app.config["BRAND_LOGO"],
            "support_email": sorted(current_app.config["ADMIN_EMAILS"])[0] if current_app.config["ADMIN_EMAILS"] else "",
            "has_homeroom_classes": bool(current_user() and SchoolClass.query.filter_by(homeroom_teacher_id=current_user().id, active=True).first()),
        }

    @app.before_request
    def require_complete_profile():
        """Yêu cầu kê khai thông tin một lần trước khi sử dụng chức năng nghiệp vụ."""
        if not session.get("user_id"):
            return None
        if (request.endpoint in {"static", "login", "auth_callback", "logout", "profile_setup",
                                 "sync_wait", "sync_run", "connect_youtube_callback", "privacy_policy", "health"}
                or (request.endpoint or "").startswith("websub_")):
            return None
        user = current_user()
        if user and not user.profile_completed:
            return redirect(url_for("profile_setup"))
        return None

    @app.get("/health")
    def health():
        try:
            db.session.execute(text("SELECT 1"))
            return jsonify({"status": "ok", "database": "ok"})
        except Exception as exc:
            return jsonify({"status": "error", "detail": str(exc)}), 500

    @app.get("/privacy")
    def privacy_policy():
        return render_template("privacy.html")

    @app.get("/")
    def index():
        user = current_user()
        if not user:
            return render_template("landing.html")
        whole_system = user.role == "admin" and request.args.get("scope") == "all"
        videos = (db.session.query(Video.id).join(YouTubeChannel)
                  .filter(Video.uploader_channel_id == YouTubeChannel.channel_id).count()
                  if whole_system else visible_video_query(user).count())
        qrs = (QrCode.query.filter_by(active=True).count() if whole_system else
               QrCode.query.filter(managed_qr_scope(user.id), QrCode.active.is_(True)).count())
        courses = (Course.query.filter_by(active=True).count() if whole_system else
                   Course.query.filter_by(owner_id=user.id, active=True).count())
        scan_q = ScanLog.query.join(QrCode, ScanLog.qr_id == QrCode.id).filter(ScanLog.role == "SV")
        watch_q = WatchSession.query.join(QrCode, WatchSession.qr_id == QrCode.id).filter(WatchSession.role == "SV")
        if not whole_system:
            scan_q = scan_q.filter(managed_qr_scope(user.id))
            watch_q = watch_q.filter(managed_qr_scope(user.id))
        scans = scan_q.count()
        watch = watch_q.count()
        completed = watch_q.filter(WatchSession.completed.is_(True)).count()
        completion_rate = round(100 * completed / watch, 1) if watch else 0
        recent_scans = scan_q.order_by(ScanLog.scanned_at.desc()).limit(10).all()
        lecturer_count = User.query.filter(User.last_login_at.isnot(None)).count()
        class_count = (SchoolClass.query.filter_by(active=True).count() if whole_system else
                       SchoolClass.query.filter_by(active=True, homeroom_teacher_id=user.id).count())
        return render_template(
            "dashboard.html", videos=videos, courses=courses, qrs=qrs, scans=scans, watch=watch,
            completion_rate=completion_rate, recent_scans=recent_scans, lecturer_count=lecturer_count,
            class_count=class_count, whole_system=whole_system,
        )

    @app.get("/auth/login")
    def login():
        if not current_app.config["GOOGLE_CLIENT_ID"]:
            flash("Chưa cấu hình GOOGLE_CLIENT_ID / GOOGLE_CLIENT_SECRET.", "danger")
            return redirect(url_for("index"))
        redirect_uri = f"{base_url()}/auth/callback"
        session["oauth_retry_attempt"] = 1 if request.args.get("retry") == "1" else 0
        session["after_login"] = safe_next(request.args.get("next")) or url_for("index")
        return oauth.google.authorize_redirect(
            redirect_uri,
            access_type="offline",
            prompt="consent",
            include_granted_scopes="true",
        )

    @app.get("/auth/callback")
    def auth_callback():
        try:
            token = oauth.google.authorize_access_token()
            access_token = token.get("access_token")
            if not access_token:
                raise RuntimeError("Google không trả về access_token")
            profile = requests.get(
                "https://openidconnect.googleapis.com/v1/userinfo",
                headers={"Authorization": f"Bearer {access_token}"},
                timeout=20,
            )
            profile.raise_for_status()
            info = profile.json()
            email = (info.get("email") or "").lower().strip()
            if not email:
                raise RuntimeError("Không đọc được email từ Google")
            user = User.query.filter_by(email=email).first()
            if not user:
                role = "admin" if email in current_app.config["ADMIN_EMAILS"] else "lecturer"
                user = User(email=email, name=info.get("name") or email, google_sub=info.get("sub"), role=role)
                db.session.add(user)
                db.session.flush()
            else:
                # Họ tên hiển thị có thể được người dùng hiệu chỉnh trong Web App;
                # không ghi đè lại bằng tên hồ sơ Google ở những lần đăng nhập sau.
                if not (user.name or "").strip():
                    user.name = info.get("name") or email
                user.google_sub = info.get("sub") or user.google_sub
                # Quyền quản trị chỉ dựa trên ADMIN_EMAILS hiện hành. Mọi tài khoản
                # Google khác luôn là giảng viên, kể cả trước đây từng được cấu hình admin.
                user.role = "admin" if email in current_app.config["ADMIN_EMAILS"] else "lecturer"
            user.last_login_at = utcnow()

            account = OAuthAccount.query.filter_by(user_id=user.id).first()
            existing_refresh = None
            if account:
                try:
                    from security import decrypt_json
                    existing_refresh = decrypt_json(account.token_encrypted).get("refresh_token")
                except Exception:
                    existing_refresh = None
            if existing_refresh and not token.get("refresh_token"):
                token["refresh_token"] = existing_refresh
            if not account:
                account = OAuthAccount(user_id=user.id, provider="google", token_encrypted="")
                db.session.add(account)
            account.token_encrypted = encrypt_json(token)
            token_scope = token.get("scope") or current_app.config["GOOGLE_SCOPES"]
            account.scopes = " ".join(token_scope) if isinstance(token_scope, (list, tuple)) else str(token_scope)
            record_functional_test(
                "AUTH_LOGIN", "Đăng nhập Google", True, owner_id=user.id,
                note=f"Đăng nhập thành công với vai trò {role_label(user.role)}."
            )
            db.session.commit()
            session["user_id"] = user.id
            session.pop("oauth_retry_attempt", None)
            session.permanent = True
            # Mã GV được cấp sau khi xác nhận thông tin hồ sơ (cả Quản trị viên).
            # Tự đồng bộ trên màn chờ mỗi lần đăng nhập thành công.
            flash(f"Đăng nhập thành công với vai trò {role_label(user.role)}.", "success")
            session.pop("after_login", None)
            if not user.profile_completed:
                return redirect(url_for("profile_setup"))
            if not user.lecturer_code:
                allocate_lecturer_code(user)
                db.session.commit()
            return redirect(url_for("sync_wait"))
        except Exception as exc:
            db.session.rollback()
            # OAuth state là biện pháp phòng chống CSRF: tuyệt đối KHÔNG bỏ qua
            # kiểm tra state hoặc tự chấp nhận callback không khớp.
            error_text = str(exc)
            state_failed = "mismatching_state" in error_text.lower() or (
                "csrf warning" in error_text.lower() and "state" in error_text.lower())
            db.session.add(ErrorLog(module="oauth", action="Đăng nhập hệ thống", course_module="Hệ thống", message="Đăng nhập Google thất bại", details=str(exc)))
            record_functional_test(
                "AUTH_LOGIN", "Đăng nhập Google", False, owner_id=None,
                note=f"OAuth thất bại: {exc}"
            )
            db.session.commit()
            if state_failed:
                # Tab cũ, nhấn đăng nhập nhiều lần hoặc cookie phiên thay đổi
                # làm state không còn khớp. Thử lại đúng MỘT lần, sau đó dừng.
                if not session.get("oauth_retry_attempt", 0) and "after_login" in session:
                    flash("Phiên xác thực Google đã thay đổi. Hệ thống sẽ khởi tạo lại đăng nhập an toàn một lần.", "warning")
                    return redirect(url_for("login", retry=1))
                session.pop("oauth_retry_attempt", None)
                flash("Phiên đăng nhập Google không khớp. Vui lòng đóng các tab đăng nhập cũ, bật cookie và thử lại bằng một tab duy nhất.", "danger")
            else:
                flash("Đăng nhập Google không thành công. Vui lòng thử lại; nếu tiếp diễn, báo quản trị viên kiểm tra Nhật ký lỗi.", "danger")
            return redirect(url_for("index"))

    @app.route("/profile/setup", methods=["GET", "POST"])
    @login_required
    def profile_setup():
        user = current_user()
        if user.profile_completed:
            return redirect(url_for("index"))
        if request.method == "POST":
            fullname = re.sub(r"\s+", " ", (request.form.get("full_name") or "").strip())
            group = (request.form.get("professional_group") or "").strip()
            homeroom = request.form.get("is_homeroom") == "yes"
            class_name = re.sub(r"\s+", " ", (request.form.get("homeroom_class_name") or "").strip())
            class_id = request.form.get("homeroom_class_id", type=int)
            declared_raw = (request.form.get("declared_size") or "").strip()
            if len(fullname) < 2 or len(fullname) > 255 or group not in {"it", "other"}:
                flash("Vui lòng nhập họ tên đầy đủ và chọn nhóm chuyên môn.", "warning")
                return render_template("profile_setup.html", available_classes=SchoolClass.query.filter_by(active=True).order_by(SchoolClass.name).all())
            existing = db.session.get(SchoolClass, class_id) if homeroom and class_id else None
            if homeroom:
                if (not existing and (not class_name or len(class_name) > 255)) or not declared_raw.isdigit() or not (1 <= int(declared_raw) <= 1000):
                    flash("Vui lòng chọn hoặc nhập tên lớp và sĩ số khai báo hợp lệ.", "warning")
                    return render_template("profile_setup.html", available_classes=SchoolClass.query.filter_by(active=True).order_by(SchoolClass.name).all())
                if existing and existing.homeroom_teacher_id not in (None, user.id):
                    flash("Lớp đã có GVCN. Quản trị viên cần bàn giao lớp trước khi bạn nhận chủ nhiệm.", "warning")
                    return render_template("profile_setup.html", available_classes=SchoolClass.query.filter_by(active=True).order_by(SchoolClass.name).all())
            try:
                user.name = fullname
                user.professional_group = group
                if homeroom:
                    if existing:
                        school_class = existing
                    else:
                        school_class, _ = create_or_find_school_class(class_name, user, int(declared_raw))
                    if school_class.homeroom_teacher_id not in (None, user.id):
                        raise ValueError("Lớp đã có giảng viên chủ nhiệm; không thể tự thay thế.")
                    if school_class.homeroom_teacher_id is None:
                        db.session.add(HomeroomChange(school_class_id=school_class.id,
                                                      new_teacher_id=user.id, changed_by=user.id,
                                                      effective_on=utcnow().date()))
                    school_class.homeroom_teacher_id = user.id
                    school_class.declared_size = int(declared_raw)
                allocate_lecturer_code(user)
                user.profile_completed = True
                db.session.commit()
                flash("Đã lưu thông tin giảng viên và cấp mã định danh tự động.", "success")
                return redirect(url_for("sync_wait"))
            except Exception as exc:
                db.session.rollback()
                flash(f"Không thể hoàn tất khai báo: {exc}", "danger")
        return render_template("profile_setup.html", available_classes=SchoolClass.query.filter_by(active=True).order_by(SchoolClass.name).all())

    @app.get("/auth/sync-wait")
    @login_required
    def sync_wait():
        if not current_user().profile_completed:
            return redirect(url_for("profile_setup"))
        return render_template("sync_wait.html")

    @app.post("/auth/sync-run")
    @login_required
    def sync_run():
        user = current_user()
        if not user.profile_completed:
            return jsonify({"ok": False, "message": "Bạn chưa hoàn tất thông tin giảng viên."}), 400
        try:
            channels, sync_errors = sync_all_connected_channels(user)
            video_count = sum(len(c.videos) for c in channels)
            record_functional_test("YT_SYNC", "Đồng bộ kênh/video YouTube", True, owner_id=user.id,
                                   note=f"Tự đồng bộ {len(channels)} kênh và {video_count} video sau đăng nhập.")
            db.session.commit()
            return jsonify({"ok": bool(channels), "message": f"Đã đồng bộ {len(channels)} kênh YouTube và {video_count} video." + (" Một số kênh cần kiểm tra: " + "; ".join(sync_errors) if sync_errors else "")})
        except Exception as exc:
            db.session.rollback()
            db.session.add(ErrorLog(module="youtube", action="Tự động đồng bộ sau đăng nhập", course_module="Hệ thống", actor_user_id=user.id,
                                    message="Không thể tự động đồng bộ kênh YouTube", details=str(exc)[:10000]))
            record_functional_test("YT_SYNC", "Đồng bộ kênh/video YouTube", False, owner_id=user.id, note=str(exc)[:1000])
            db.session.commit()
            return jsonify({"ok": False, "message": "Chưa đồng bộ được kênh YouTube. Bạn vẫn có thể vào hệ thống và thử lại sau."})

    @app.post("/profile/display-name")
    @login_required
    def update_display_name():
        user = current_user()
        display_name = re.sub(r"\s+", " ", (request.form.get("display_name") or "").strip())
        if len(display_name) < 2:
            flash("Vui lòng nhập họ và tên đầy đủ để hiển thị trên hệ thống.", "warning")
            return redirect(url_for("index"))
        user.name = display_name[:255]
        db.session.commit()
        flash("Đã cập nhật họ tên hiển thị.", "success")
        return redirect(url_for("index"))

    @app.get("/auth/logout")
    def logout():
        session.clear()
        return redirect(url_for("index"))

    @app.get("/youtube")
    @login_required
    def youtube_page():
        user = current_user()
        channels = YouTubeChannel.query.filter_by(user_id=user.id).all()
        channel_ids = [c.id for c in channels]
        hub_events = HubEvent.query.filter(HubEvent.channel_id.in_(channel_ids)).order_by(HubEvent.received_at.desc()).limit(20).all() if channel_ids else []
        return render_template("youtube.html", channels=channels, hub_events=hub_events, extra_accounts=ConnectedGoogleAccount.query.filter_by(user_id=user.id).all())

    @app.post("/youtube/sync")
    @login_required
    def youtube_sync():
        user = current_user()
        t0 = time.perf_counter()
        try:
            channels, sync_errors = sync_all_connected_channels(user)
            # sync_my_channel đã đồng bộ uploads, các playlist do GV tạo và trạng thái
            # video hiện có; không gọi đối soát lần hai để tránh tiêu tốn hạn mức API.
            checked = sum(len(c.videos) for c in channels)
            elapsed = time.perf_counter() - t0
            record_functional_test(
                "YT_SYNC", "Đồng bộ kênh/video YouTube", True, owner_id=user.id,
                note=f"Đồng bộ {len(channels)} kênh; đối soát {checked} video."
            )
            db.session.commit()
            flash(f"Đã đồng bộ {len(channels)} kênh YouTube, kiểm tra {checked} video." + (" Lỗi kênh khác: " + "; ".join(sync_errors) if sync_errors else ""), "success" if channels else "warning")
        except Exception as exc:
            db.session.rollback()
            elapsed = time.perf_counter() - t0
            db.session.add(ErrorLog(
                module="youtube", action="Đồng bộ kênh/video", course_module="Hệ thống", actor_user_id=user.id,
                message="Đồng bộ kênh/video YouTube thất bại", details=str(exc)
            ))
            record_functional_test(
                "YT_SYNC", "Đồng bộ kênh/video YouTube", False, owner_id=user.id, note=str(exc)
            )
            db.session.commit()
            flash(f"Đồng bộ thất bại: {exc}", "danger")
        return redirect(url_for("youtube_page"))

    @app.get("/youtube/connect-another")
    @login_required
    def connect_another_youtube():
        """Kết nối YouTube của tài khoản Google khác, vẫn giữ phiên GV hiện tại."""
        session["extra_youtube_owner_id"] = current_user().id
        return oauth.google_additional.authorize_redirect(
            f"{base_url()}/youtube/connect-callback",
            access_type="offline", prompt="select_account consent",
            include_granted_scopes="true",
        )

    @app.get("/youtube/connect-callback")
    @login_required
    def connect_youtube_callback():
        user = current_user()
        if session.pop("extra_youtube_owner_id", None) != user.id:
            flash("Phiên kết nối YouTube đã hết hạn. Vui lòng thử lại.", "warning")
            return redirect(url_for("youtube_page"))
        try:
            token = oauth.google_additional.authorize_access_token()
            r = requests.get("https://openidconnect.googleapis.com/v1/userinfo",
                             headers={"Authorization": f"Bearer {token['access_token']}"}, timeout=20)
            r.raise_for_status()
            info = r.json()
            sub = info.get("sub")
            email = (info.get("email") or "").strip().lower()
            if not sub or not email:
                raise ValueError("Google không cung cấp danh tính tài khoản kết nối.")
            primary_user = User.query.filter_by(google_sub=sub).first()
            if primary_user and primary_user.id != user.id:
                raise ValueError("Tài khoản Google này đã là tài khoản đăng nhập chính của giảng viên khác.")
            existing = ConnectedGoogleAccount.query.filter_by(google_sub=sub).first()
            if existing and existing.user_id != user.id:
                raise ValueError("Tài khoản YouTube này đã liên kết với giảng viên khác.")
            if primary_user and primary_user.id == user.id:
                flash("Đây là tài khoản Google đang dùng để đăng nhập; kênh đã được quản lý ở kết nối chính.", "info")
                return redirect(url_for("youtube_page"))
            if not existing:
                existing = ConnectedGoogleAccount(user_id=user.id, google_sub=sub, email=email, token_encrypted="")
                db.session.add(existing)
            old_refresh = None
            if existing.token_encrypted:
                try:
                    old_refresh = decrypt_json(existing.token_encrypted).get("refresh_token")
                except Exception:
                    pass
            if old_refresh and not token.get("refresh_token"):
                token["refresh_token"] = old_refresh
            existing.token_encrypted = encrypt_json(token)
            existing.scopes = token.get("scope") or current_app.config["GOOGLE_SCOPES"]
            db.session.commit()
            flash(f"Đã kết nối tài khoản YouTube bổ sung {email}. Bấm Đồng bộ lại để tải học liệu.", "success")
        except Exception as exc:
            db.session.rollback()
            flash(f"Không thể kết nối thêm kênh YouTube: {exc}", "danger")
        return redirect(url_for("youtube_page"))

    @app.post("/youtube/disconnect/<int:connection_id>")
    @login_required
    def disconnect_youtube_account(connection_id):
        user = current_user()
        conn = db.session.get(ConnectedGoogleAccount, connection_id) or abort(404)
        if conn.user_id != user.id:
            abort(403)
        # Để bảo toàn học liệu, QR và thống kê, chỉ ngắt token; các video đã lưu vẫn giữ.
        for channel in YouTubeChannel.query.filter_by(user_id=user.id, connected_account_id=conn.id).all():
            channel.connected_account_id = None
        db.session.delete(conn)
        db.session.commit()
        flash("Đã ngắt liên kết tài khoản YouTube bổ sung. Dữ liệu video và QR cũ vẫn được giữ.", "success")
        return redirect(url_for("youtube_page"))

    @app.post("/youtube/websub/<int:channel_pk>/subscribe")
    @login_required
    def youtube_websub_subscribe(channel_pk):
        channel = db.session.get(YouTubeChannel, channel_pk) or abort(404)
        user = current_user()
        if user.role != "admin" and channel.user_id != user.id:
            abort(403)
        try:
            subscribe_channel(channel)
            flash("Đã kích hoạt cập nhật tự động cho kênh YouTube. Hệ thống đang chờ xác nhận kết nối.", "success")
        except Exception as exc:
            flash("Không thể kích hoạt cập nhật tự động từ YouTube. Vui lòng kiểm tra nhật ký lỗi.", "danger")
        return redirect(url_for("youtube_page"))

    @app.route("/websub/callback/<channel_youtube_id>/<secret>", methods=["GET", "POST"])
    @csrf.exempt
    def websub_callback(channel_youtube_id, secret):
        channel = YouTubeChannel.query.filter_by(channel_id=channel_youtube_id).first_or_404()
        if request.method == "GET":
            challenge = verify_callback(channel, secret, request.args)
            if challenge is None:
                abort(404)
            return Response(challenge, mimetype="text/plain")

        sub = HubSubscription.query.filter_by(channel_id=channel.id, callback_secret=secret).first()
        if not sub:
            abort(404)
        evt = parse_notification(channel, request.data)
        if evt.youtube_video_id:
            try:
                sync_video_ids(channel.user, channel, [evt.youtube_video_id])
                evt.processed = True
                db.session.commit()
            except Exception as exc:
                db.session.add(ErrorLog(module="websub", action="Đồng bộ video", course_module="Hệ thống", message="Nhận thông báo cập nhật nhưng đồng bộ video lỗi", details=str(exc)))
                db.session.commit()
        return "", 204

    @app.get("/videos")
    @login_required
    def videos_page():
        user = current_user()

        # --- Bộ lọc video ---
        q = visible_video_query(user)
        status = request.args.get("status", "").strip()
        assignment = request.args.get("assignment", "").strip()
        search = request.args.get("q", "").strip()
        playlist_filter_id = request.args.get("playlist_id", type=int)

        if status == "private":
            q = q.filter(Video.privacy_status == "private")
        elif status == "noembed":
            q = q.filter(Video.embeddable.is_(False))
        elif status == "unavailable":
            q = q.filter(Video.is_available.is_(False))
        elif status == "usable":
            q = q.filter(Video.is_available.is_(True), Video.privacy_status != "private", Video.embeddable.isnot(False))
        if search:
            q = q.filter(Video.title.ilike(f"%{search}%"))

        # Danh sách playlist của chính giảng viên để lọc video và gán môn học/ mô đun.
        available_playlists = visible_playlist_query(user).order_by(Playlist.title).all()
        allowed_playlist_ids = {p.id for p in available_playlists}
        if playlist_filter_id:
            if playlist_filter_id in allowed_playlist_ids:
                playlist_video_ids = db.session.query(PlaylistVideo.video_id).filter(PlaylistVideo.playlist_id == playlist_filter_id)
                q = q.filter(Video.id.in_(playlist_video_ids))
            else:
                # Không cho dùng ID playlist ngoài phạm vi tài khoản để lọc dữ liệu.
                q = q.filter(Video.id == -1)

        owned_course_ids = db.session.query(Course.id).filter(Course.owner_id == user.id, Course.active.is_(True))
        assigned_video_ids = db.session.query(CourseVideo.video_id).filter(CourseVideo.course_id.in_(owned_course_ids))
        if assignment == "unassigned":
            q = q.filter(~Video.id.in_(assigned_video_ids))
        elif assignment == "assigned":
            q = q.filter(Video.id.in_(assigned_video_ids))
        video_total = q.count()
        video_pages = max(1, (video_total + 4) // 5)
        video_page = min(max(1, request.args.get("v_page", 1, type=int)), video_pages)
        videos = q.order_by(Video.published_at.desc().nullslast(), Video.id.desc()).offset((video_page-1)*5).limit(5).all()

        # --- Bộ lọc playlist ---
        playlist_search = request.args.get("p_q", "").strip()
        playlist_assignment = request.args.get("p_assignment", "").strip()
        assigned_playlist_ids = (
            db.session.query(CoursePlaylist.playlist_id)
            .filter(CoursePlaylist.course_id.in_(owned_course_ids), CoursePlaylist.active.is_(True))
        )
        # Vẫn liệt kê các playlist đã mất trên YouTube nếu còn được gán môn,
        # để giáo viên nhìn thấy tình trạng và ngừng QR nhưng không xóa lịch sử.
        pq = (Playlist.query.join(YouTubeChannel).filter(
            YouTubeChannel.user_id == user.id,
            or_(Playlist.is_available.is_(True), Playlist.id.in_(assigned_playlist_ids)),
            or_(YouTubeChannel.uploads_playlist_id.is_(None),
                Playlist.youtube_playlist_id != YouTubeChannel.uploads_playlist_id)))
        if playlist_search:
            pq = pq.filter(Playlist.title.ilike(f"%{playlist_search}%"))
        if playlist_assignment == "unassigned":
            pq = pq.filter(~Playlist.id.in_(assigned_playlist_ids))
        elif playlist_assignment == "assigned":
            pq = pq.filter(Playlist.id.in_(assigned_playlist_ids))
        # Playlist đã bị xóa/ẩn trên YouTube vẫn hiện để GV xử lý QR cũ,
        # nhưng không cho tạo QR mới từ nguồn không khả dụng.
        playlist_total = pq.count()
        playlist_pages = max(1, (playlist_total + 4) // 5)
        playlist_page = min(max(1, request.args.get("p_page", 1, type=int)), playlist_pages)
        playlists = pq.order_by(Playlist.is_available.desc(), Playlist.title).offset((playlist_page-1)*5).limit(5).all()

        courses = Course.query.filter_by(owner_id=user.id, active=True).order_by(Course.name).all()
        return render_template(
            "videos.html",
            videos=videos,
            status=status,
            assignment=assignment,
            search=search,
            playlist_filter_id=playlist_filter_id,
            playlists=playlists,
            available_playlists=available_playlists,
            playlist_search=playlist_search,
            playlist_assignment=playlist_assignment,
            video_page=video_page, video_pages=video_pages, video_total=video_total,
            playlist_page=playlist_page, playlist_pages=playlist_pages, playlist_total=playlist_total,
            library_args=request.args.to_dict(flat=True),
            retarget_videos=visible_video_query(user).filter(
                Video.is_available.is_(True), Video.privacy_status != "private",
                Video.embeddable.isnot(False)).order_by(Video.title).all(),
            courses=courses,
        )

    @app.post("/videos/assign-course")
    @login_required
    def assign_videos_to_course():
        user = current_user()
        course_id = request.form.get("course_id", type=int)
        video_ids = request.form.getlist("video_ids")
        course = db.session.get(Course, course_id) if course_id else None
        if not course or not course.active or not course_owned_by_user(course, user):
            abort(403)
        if not video_ids:
            flash("Vui lòng chọn ít nhất một video để gán vào môn học/ mô đun.", "warning")
            return redirect(url_for("videos_page"))
        allowed_ids = {v.id for v in visible_video_query(user).filter(Video.id.in_([int(x) for x in video_ids if str(x).isdigit()])).all()}
        t0 = time.perf_counter()
        created = 0
        skipped = 0
        for video_id in allowed_ids:
            video = db.session.get(Video, video_id)
            _, was_created = ensure_course_video_link(course, video, user)
            if was_created:
                created += 1
            else:
                skipped += 1
        elapsed = time.perf_counter() - t0
        record_functional_test(
            "ASSIGN_COURSE", "Gán video vào môn học/ mô đun", True, owner_id=user.id,
            note=f"Môn học/ mô đun: {course.name}; gán mới {created}; bỏ qua {skipped}."
        )
        db.session.commit()
        flash(f"Đã gán {created} video vào môn học/ mô đun “{course.name}”. Bỏ qua {skipped} video đã được gán trước đó.", "success")
        return redirect(videos_library_return("videos", request.form.get("return_query", "")))

    @app.post("/playlists/assign-course")
    @login_required
    def assign_playlists_from_library():
        """Gán playlist ngay tại Kho video học liệu để giảm thao tác cho giảng viên."""
        user = current_user()
        course_id = request.form.get("course_id", type=int)
        playlist_ids = [int(x) for x in request.form.getlist("playlist_ids") if str(x).isdigit()]
        course = db.session.get(Course, course_id) if course_id else None
        if not course or not course.active or not course_owned_by_user(course, user):
            abort(403)
        if not playlist_ids:
            flash("Vui lòng chọn ít nhất một playlist để gán vào môn học/ mô đun.", "warning")
            return redirect(url_for("videos_page"))

        allowed = {p.id: p for p in visible_playlist_query(user).filter(Playlist.id.in_(playlist_ids)).all()}
        created = 0
        skipped = 0
        for playlist_id in playlist_ids:
            playlist = allowed.get(playlist_id)
            if not playlist:
                continue
            _, was_created = ensure_course_playlist_link(course, playlist, user)
            if was_created:
                created += 1
            else:
                skipped += 1
        auto_added, auto_removed = reconcile_course_videos_from_playlists(course, user)
        record_functional_test(
            "ASSIGN_PLAYLIST", "Gán playlist vào môn học/ mô đun", bool(created or skipped), owner_id=user.id,
            note=f"Kho học liệu → {course.name}; playlist mới/khôi phục {created}; đã có {skipped}; video tự gán {auto_added}; video tự bỏ {auto_removed}."
        )
        db.session.commit()
        flash(
            f"Đã cập nhật playlist cho môn học/ mô đun “{course.name}”. Hệ thống tự gán thêm {auto_added} video thuộc playlist; không cần chọn lại từng video.",
            "success",
        )
        return redirect(url_for("videos_page") + "#playlists")

    @app.post("/playlists/<int:playlist_id>/set-courses")
    @login_required
    def set_playlist_course_assignments(playlist_id):
        """Lưu toàn bộ môn học/ mô đun áp dụng cho một playlist trong một thao tác.

        Danh sách môn học/ mô đun gửi lên là trạng thái cuối cùng: môn học/ mô đun được chọn sẽ
        nhận playlist và toàn bộ video của playlist; môn học/ mô đun bỏ chọn sẽ bỏ playlist
        và các video chỉ được kế thừa từ playlist đó.
        """
        user = current_user()
        playlist = db.session.get(Playlist, playlist_id) or abort(404)
        if playlist.channel.user_id != user.id:
            abort(403)
        selected_course_ids = {int(x) for x in request.form.getlist("course_ids") if str(x).isdigit()}
        activated, deactivated, auto_added, auto_removed = set_playlist_courses(playlist, selected_course_ids, user)
        record_functional_test(
            "SET_PLAYLIST_COURSES", "Thiết lập môn học/ mô đun cho playlist", True, owner_id=user.id,
            note=(f"Playlist: {playlist.title}; thêm/khôi phục môn học/ mô đun {activated}; bỏ môn học/ mô đun {deactivated}; "
                  f"video tự gán {auto_added}; video tự bỏ {auto_removed}."),
        )
        db.session.commit()
        flash(
            f"Đã lưu môn học/ mô đun cho playlist “{playlist.title}”. Video trong playlist được tự động đồng bộ vào các môn học/ mô đun đã chọn.",
            "success",
        )
        return redirect(videos_library_return(
            f"playlist-{playlist.id}", request.form.get("return_query", "")
        ))

    @app.post("/playlists/<int:playlist_id>/qr")
    @login_required
    def create_playlist_qr_from_library(playlist_id):
        user = current_user()
        playlist = visible_playlist_query(user).filter(Playlist.id == playlist_id).first_or_404()
        course_id = request.form.get("course_id", type=int)
        course = db.session.get(Course, course_id) if course_id else None
        if not course or not course.active or not course_owned_by_user(course, user):
            abort(403)
        link = CoursePlaylist.query.filter_by(course_id=course.id, playlist_id=playlist.id, active=True).first()
        if not link:
            flash("Cần lưu môn học/mô đun áp dụng cho playlist trước khi tạo mã QR.", "warning")
            return redirect(videos_library_return(f"playlist-{playlist.id}", request.form.get("return_query", "")))
        qr, was_created = create_qr_for_course_playlist(course, playlist, user, link)
        record_functional_test(
            "QR_PLAYLIST_SINGLE", "Tạo mã QR cho một playlist", True, owner_id=user.id,
            note=f"Kho học liệu: {playlist.title}; môn học/mô đun: {course.name}; tạo mới: {was_created}."
        )
        db.session.commit()
        flash("Đã tạo mã QR cho playlist." if was_created else
              "Playlist đã có mã QR ở môn học/mô đun này; không tạo mã trùng.",
              "success" if was_created else "info")
        return redirect(videos_library_return(f"playlist-{playlist.id}", request.form.get("return_query", "")))

    @app.post("/videos/<int:video_id>/qr")
    @login_required
    def create_video_qr_from_library(video_id):
        user = current_user()
        video = visible_video_query(user).filter(Video.id == video_id).first_or_404()
        course_id = request.form.get("course_id", type=int)
        course = db.session.get(Course, course_id) if course_id else None
        if not course or not course.active or not course_owned_by_user(course, user):
            abort(403)
        if not video_is_usable(video):
            flash("Video chưa đủ điều kiện phát và tạo mã QR. Hãy kiểm tra lại video.", "warning")
            return redirect(videos_library_return(f"video-{video.id}", request.form.get("return_query", "")))
        started = time.perf_counter()
        link, _ = ensure_course_video_link(course, video, user, manual=True)
        qr, was_created = create_qr_for_course_video(course, video, user, link)
        elapsed = time.perf_counter() - started
        record_functional_test(
            "QR_SINGLE", "Tạo một mã QR", True, owner_id=user.id,
            note=f"Kho học liệu: {video.title}; môn học/mô đun: {course.name}; tạo mới: {was_created}."
        )
        record_webapp_experiment(
            user.id, "TAO_QR_RIENG", elapsed, operations_count=1, success=True,
            note=f"Video: {video.title}; tạo mới: {was_created}",
            processing_seconds=(qr.processing_seconds if was_created else 0.0),
        )
        db.session.commit()
        flash("Đã tạo mã QR cho video." if was_created else
              "Video đã có mã QR ở môn học/mô đun này; không tạo mã trùng.",
              "success" if was_created else "info")
        return redirect(videos_library_return(f"video-{video.id}", request.form.get("return_query", "")))

    @app.route("/courses", methods=["GET", "POST"])
    @login_required
    def courses_page():
        user = current_user()
        if request.method == "POST":
            name = request.form.get("course_name", "").strip()
            faculty = request.form.get("faculty_name", "").strip()
            department = request.form.get("department_name", "").strip()
            if not name or not faculty or not department:
                flash("Vui lòng nhập tên môn học/ mô đun, tên khoa và bộ môn quản lý môn học/ mô đun.", "warning")
            else:
                try:
                    normalized_name = re.sub(r"\s+", " ", name).casefold()
                    same_course = next(
                        (c for c in Course.query.filter_by(owner_id=user.id, active=True).all()
                         if re.sub(r"\s+", " ", c.name.strip()).casefold() == normalized_name),
                        None
                    )
                    if same_course:
                        # Một môn học dùng được cho nhiều lớp; không tạo bản ghi trùng.
                        normalize = lambda value: re.sub(r"\s+", " ", (value or "").strip()).casefold()
                        if (normalize(same_course.faculty_name or same_course.unit_name) != normalize(faculty)
                                or normalize(same_course.department_name) != normalize(department)):
                            flash(f"Môn học/mô đun ‘{same_course.name}’ đã tồn tại, nhưng tên khoa/bộ môn khác thông tin vừa nhập. "
                                  "Hãy kiểm tra và chỉnh sửa môn hiện có trước khi thêm lớp; hệ thống không tạo bản sao.", "warning")
                            return redirect(url_for("courses_page") + f"#course-{same_course.id}")
                        selected_ids = {int(x) for x in request.form.getlist("school_class_ids") if x.isdigit()}
                        new_class_name = (request.form.get("new_class_name") or "").strip()
                        if new_class_name:
                            school_class, _ = create_or_find_school_class(new_class_name, user)
                            selected_ids.add(school_class.id)
                        valid_ids = {cl.id for cl in SchoolClass.query.filter(SchoolClass.active.is_(True),
                                      SchoolClass.id.in_(selected_ids)).all()} if selected_ids else set()
                        for class_id in valid_ids:
                            link = CourseClass.query.filter_by(course_id=same_course.id, school_class_id=class_id).first()
                            if link:
                                link.active = True
                            else:
                                db.session.add(CourseClass(course_id=same_course.id, school_class_id=class_id,
                                                           added_by=user.id, active=True))
                        db.session.commit()
                        flash(f"Môn học/mô đun ‘{same_course.name}’ đã tồn tại, hệ thống không tạo trùng. "
                              f"Đã cập nhật lớp áp dụng nếu có chọn lớp mới.", "info")
                        return redirect(url_for("courses_page") + f"#course-{same_course.id}")
                    # GV không phải tự đặt mã. Tạo mã tạm để lấy ID, sau đó đổi sang mã tự động ổn định.
                    course = Course(
                        owner_id=user.id,
                        code=f"TMP-{secrets.token_hex(8)}",
                        name=name,
                        unit_name=faculty,  # giữ tương thích dữ liệu/backup cũ
                        faculty_name=faculty,
                        department_name=department,
                    )
                    db.session.add(course)
                    db.session.flush()
                    assign_auto_course_code(course)
                    new_class_name = (request.form.get("new_class_name") or "").strip()
                    added_class_id = None
                    if new_class_name:
                        new_class, was_created = create_or_find_school_class(new_class_name, user)
                        added_class_id = new_class.id
                        if not was_created:
                            flash("Lớp đã tồn tại nên hệ thống sử dụng lớp chung, không tạo lớp trùng.", "info")
                    selected_ids = {int(x) for x in request.form.getlist("school_class_ids") if x.isdigit()}
                    if added_class_id:
                        selected_ids.add(added_class_id)
                    allowed_ids = {cl.id for cl in SchoolClass.query.filter(SchoolClass.active.is_(True), SchoolClass.id.in_(selected_ids)).all()} if selected_ids else set()
                    for class_id in allowed_ids:
                        db.session.add(CourseClass(course_id=course.id, school_class_id=class_id, added_by=user.id, active=True))
                    db.session.commit()
                    flash(f"Đã thêm môn học/ mô đun “{course.name}”. Mã môn học/ mô đun được hệ thống tự tạo.", "success")
                except Exception as exc:
                    db.session.rollback()
                    flash(f"Không thêm được môn học/ mô đun: {exc}", "danger")
            return redirect(url_for("courses_page"))
        all_owned_courses = Course.query.filter_by(owner_id=user.id).order_by(Course.created_at.desc()).all()
        courses = [c for c in all_owned_courses if c.active]
        # Không gợi ý tên môn học/ mô đun cũ. Tên khoa và bộ môn từng dùng vẫn được giữ
        # trong gợi ý để GV chọn lại nhanh, kể cả khi một môn học/ mô đun cũ đã được ẩn.
        faculty_options = sorted({((c.faculty_name or c.unit_name) or "").strip() for c in all_owned_courses if ((c.faculty_name or c.unit_name) or "").strip()}, key=str.casefold)
        department_options = sorted({(c.department_name or "").strip() for c in all_owned_courses if (c.department_name or "").strip()}, key=str.casefold)
        available_playlists = visible_playlist_query(user).order_by(Playlist.title).all()
        available_school_classes = SchoolClass.query.filter_by(active=True).order_by(SchoolClass.name).all()
        # Quản trị viên có chức năng bàn giao, không chia sẻ quyền này cho GV thường.
        transfer_courses = (Course.query.filter_by(active=True).order_by(Course.name, Course.id).all()
                            if user.role == "admin" else [])
        transfer_lecturers = (User.query.filter(User.role.in_(["admin", "lecturer"]))
                              .order_by(User.name, User.email).all() if user.role == "admin" else [])
        transfer_history = (CourseTransfer.query.order_by(CourseTransfer.created_at.desc(), CourseTransfer.id.desc())
                            .limit(15).all() if user.role == "admin" else [])
        return render_template(
            "courses.html",
            courses=courses,
            transfer_courses=transfer_courses, transfer_lecturers=transfer_lecturers,
            transfer_history=transfer_history,
            faculty_options=faculty_options,
            department_options=department_options,
            available_playlists=available_playlists,
            available_school_classes=available_school_classes,
        )

    @app.post("/admin/courses/transfer")
    @login_required
    def transfer_course_owner():
        """Chuyển GV phụ trách, không đổi Course.id/QR.slug/tác giả QR và các log.

        Không nhận ngày tương lai: giao dịch chuyển quyền có hiệu lực ngay, lưu ngày
        thực hiện để truy vết. Mọi liên kết lớp, video, playlist vẫn nguyên vẹn.
        """
        admin = current_user()
        if admin.role != "admin":
            abort(403)
        course_id = request.form.get("course_id", type=int)
        new_owner_id = request.form.get("new_owner_id", type=int)
        if request.form.get("confirm_transfer") != "yes":
            flash("Cần xác nhận bàn giao và giữ nguyên mã QR đã phát hành.", "warning")
            return redirect(url_for("courses_page") + "#admin-course-transfer")
        if not course_id or not new_owner_id:
            flash("Vui lòng chọn môn học/mô đun và giảng viên tiếp nhận.", "warning")
            return redirect(url_for("courses_page") + "#admin-course-transfer")
        try:
            course = Course.query.filter_by(id=course_id, active=True).with_for_update().first()
            target = db.session.get(User, new_owner_id)
            if not course or not target or target.role not in {"admin", "lecturer"}:
                flash("Môn học hoặc giảng viên tiếp nhận không hợp lệ.", "warning")
                return redirect(url_for("courses_page") + "#admin-course-transfer")
            if course.owner_id == target.id:
                flash("Giảng viên được chọn đang phụ trách môn học/mô đun này.", "info")
                return redirect(url_for("courses_page") + "#admin-course-transfer")
            normalized = lambda value: re.sub(r"\s+", " ", (value or "").strip()).casefold()
            duplicate = any(normalized(other.name) == normalized(course.name)
                            for other in Course.query.filter_by(owner_id=target.id, active=True).all())
            if duplicate:
                flash("Giảng viên tiếp nhận đã có môn học/mô đun cùng tên. Cần kiểm tra và xử lý trùng môn trước; "
                      "hệ thống không tự gộp để tránh làm mất liên kết QR.", "warning")
                return redirect(url_for("courses_page") + "#admin-course-transfer")
            previous_owner_id = course.owner_id
            course.owner_id = target.id
            db.session.add(CourseTransfer(
                course_id=course.id, previous_owner_id=previous_owner_id,
                new_owner_id=target.id, changed_by=admin.id, effective_on=utcnow().date()))
            record_functional_test(
                "COURSE_TRANSFER", "Bàn giao môn học/mô đun", True, owner_id=admin.id,
                note=f"Môn: {course.name} ({course.code}); GV cũ ID {previous_owner_id}; GV mới: {target.name or target.email}; giữ QR và lịch sử.")
            db.session.commit()
            flash(f"Đã bàn giao môn “{course.name}” cho {target.name or target.email}. "
                  "Mã QR đã phát hành, địa chỉ QR, lớp, video, playlist và lịch sử quét/xem được giữ nguyên. "
                  "Giảng viên mới quản lý QR thuộc môn; tài khoản YouTube gốc không bị chuyển quyền.", "success")
        except Exception as exc:
            db.session.rollback()
            flash(f"Không bàn giao được môn học/mô đun: {exc}", "danger")
        return redirect(url_for("courses_page") + "#admin-course-transfer")

    @app.post("/courses/<int:course_id>/classes/new")
    @login_required
    def add_school_class_to_course(course_id):
        user = current_user()
        course = db.session.get(Course, course_id) or abort(404)
        if not course.active or course.owner_id != user.id:
            abort(403)
        name = (request.form.get("new_class_name") or "").strip()
        try:
            school_class, created = create_or_find_school_class(name, user)
            link = CourseClass.query.filter_by(course_id=course.id, school_class_id=school_class.id).first()
            if link:
                link.active = True
            else:
                db.session.add(CourseClass(course_id=course.id, school_class_id=school_class.id, added_by=user.id))
            db.session.commit()
            flash(("Đã thêm lớp mới" if created else "Đã chọn lớp có sẵn") + f" “{school_class.name}” cho môn học.", "success")
        except Exception as exc:
            db.session.rollback()
            flash(f"Không thêm được lớp: {exc}", "danger")
        return redirect(url_for("courses_page"))

    @app.post("/courses/<int:course_id>/info")
    @login_required
    def update_course_info(course_id):
        user = current_user()
        course = db.session.get(Course, course_id) or abort(404)
        if not course.active or not course_owned_by_user(course, user):
            abort(403)
        name = re.sub(r"\s+", " ", request.form.get("course_name", "").strip())
        faculty = re.sub(r"\s+", " ", request.form.get("faculty_name", "").strip())
        department = re.sub(r"\s+", " ", request.form.get("department_name", "").strip())
        if not name or not faculty or not department or max(len(name), len(faculty), len(department)) > 255:
            flash("Vui lòng nhập đầy đủ tên môn học/mô đun, khoa và bộ môn hợp lệ.", "warning")
            return redirect(url_for("courses_page"))
        normalized_name = re.sub(r"\s+", " ", name).casefold()
        if any(c.id != course.id and re.sub(r"\s+", " ", c.name.strip()).casefold() == normalized_name
               for c in Course.query.filter_by(owner_id=user.id, active=True).all()):
            flash("Đã tồn tại môn học/mô đun cùng tên. Hãy dùng một môn cho nhiều lớp, không đổi thành tên trùng.", "warning")
            return redirect(url_for("courses_page") + f"#course-{course.id}")
        for field, before, after in (("name", course.name, name),
                                     ("faculty_name", course.faculty_name, faculty),
                                     ("department_name", course.department_name, department)):
            audit_edit(user, "course", course.id, field, before, after)
        course.name = name
        course.faculty_name = faculty
        course.department_name = department
        course.unit_name = faculty
        db.session.commit()
        flash("Đã lưu thông tin môn học/mô đun. Mã môn, lớp, QR và thống kê không thay đổi.", "success")
        return redirect(url_for("courses_page"))

    @app.post("/courses/<int:course_id>/playlists")
    @login_required
    def assign_playlists_to_course(course_id):
        user = current_user()
        course = db.session.get(Course, course_id) or abort(404)
        if not course.active or not course_owned_by_user(course, user):
            abort(403)
        playlist_ids = [int(x) for x in request.form.getlist("playlist_ids") if str(x).isdigit()]
        if not playlist_ids:
            flash("Vui lòng chọn ít nhất một playlist để gán vào môn học/ mô đun.", "warning")
            return redirect(url_for("courses_page"))
        allowed = {p.id: p for p in visible_playlist_query(user).filter(Playlist.id.in_(playlist_ids)).all()}
        created = 0
        skipped = 0
        for playlist_id in playlist_ids:
            playlist = allowed.get(playlist_id)
            if not playlist:
                continue
            _, was_created = ensure_course_playlist_link(course, playlist, user)
            if was_created:
                created += 1
            else:
                skipped += 1
        auto_added, auto_removed = reconcile_course_videos_from_playlists(course, user)
        record_functional_test(
            "ASSIGN_PLAYLIST", "Gán playlist vào môn học/ mô đun", bool(created or skipped), owner_id=user.id,
            note=f"Môn học/ mô đun: {course.name}; playlist mới/khôi phục {created}; đã có {skipped}; video tự gán {auto_added}; video tự bỏ {auto_removed}."
        )
        db.session.commit()
        flash(f"Đã gán playlist vào môn học/ mô đun “{course.name}” và tự động cập nhật các video thuộc playlist.", "success")
        return redirect(url_for("courses_page"))

    @app.post("/course-playlists/<int:link_id>/delete")
    @login_required
    def delete_course_playlist(link_id):
        user = current_user()
        link = db.session.get(CoursePlaylist, link_id) or abort(404)
        if not course_owned_by_user(link.course, user):
            abort(403)
        link.active = False
        auto_added, auto_removed = reconcile_course_videos_from_playlists(link.course, user)
        db.session.commit()
        flash(
            f"Đã bỏ playlist “{link.playlist.title}” khỏi môn học/ mô đun và tự động bỏ {auto_removed} video chỉ được kế thừa từ playlist này. Video đã gán riêng hoặc còn thuộc playlist khác vẫn được giữ.",
            "success",
        )
        return redirect(url_for("courses_page"))

    @app.post("/courses/<int:course_id>/school-classes")
    @login_required
    def set_course_school_classes(course_id):
        user = current_user()
        course = db.session.get(Course, course_id) or abort(404)
        if not course.active or not course_owned_by_user(course, user):
            abort(403)
        selected_ids = {int(x) for x in request.form.getlist("school_class_ids") if str(x).isdigit()}
        valid_ids = {c.id for c in SchoolClass.query.filter(SchoolClass.id.in_(selected_ids), SchoolClass.active.is_(True)).all()} if selected_ids else set()
        existing = {link.school_class_id: link for link in CourseClass.query.filter_by(course_id=course.id).all()}
        for class_id in valid_ids:
            link = existing.get(class_id)
            if link:
                link.active = True
                link.added_by = link.added_by or user.id
            else:
                db.session.add(CourseClass(course_id=course.id, school_class_id=class_id, added_by=user.id, active=True))
        for class_id, link in existing.items():
            if class_id not in valid_ids:
                link.active = False
        db.session.commit()
        flash(f"Đã cập nhật lớp áp dụng cho môn học/ mô đun “{course.name}”.", "success")
        return redirect(url_for("courses_page"))

    @app.post("/courses/<int:course_id>/classes")
    @login_required
    def add_class(course_id):
        course = db.session.get(Course, course_id) or abort(404)
        user = current_user()
        if user.role != "admin" and course.owner_id != user.id:
            abort(403)
        if not course.active:
            abort(404)
        name = request.form.get("class_name", "").strip()
        if not name:
            flash("Vui lòng nhập tên lớp.", "warning")
            return redirect(url_for("courses_page"))
        try:
            class_group = ClassGroup(
                course_id=course.id,
                code=f"TMP-{secrets.token_hex(8)}",
                name=name,
            )
            db.session.add(class_group)
            db.session.flush()
            assign_auto_class_code(class_group)
            db.session.commit()
            flash(f"Đã thêm lớp “{class_group.name}”. Mã lớp được hệ thống tự tạo.", "success")
        except Exception as exc:
            db.session.rollback()
            flash(f"Không thêm được lớp: {exc}", "danger")
        return redirect(url_for("courses_page"))

    @app.post("/courses/<int:course_id>/delete")
    @login_required
    def delete_course(course_id):
        """Ẩn môn học/ mô đun khỏi danh sách sử dụng nhưng giữ dữ liệu lịch sử.

        Không xóa cứng vì ScanLog/WatchSession/QR là minh chứng nghiên cứu. Các mã QR
        đã phát hành vẫn được giữ nguyên và có thể được xóa riêng nếu giảng viên muốn.
        """
        user = current_user()
        course = db.session.get(Course, course_id) or abort(404)
        if not course_owned_by_user(course, user):
            abort(403)
        course.active = False
        for cls in course.classes:
            if hasattr(cls, "active"):
                cls.active = False
        for link in course.playlist_links:
            link.active = False
        for link in course.school_class_links:
            link.active = False
        db.session.commit()
        flash(
            f"Đã xóa môn học/ mô đun “{course.name}” khỏi danh sách sử dụng. Dữ liệu lịch sử và mã QR đã phát hành được giữ lại để không mất minh chứng.",
            "success",
        )
        return redirect(url_for("courses_page"))

    @app.post("/classes/<int:class_id>/delete")
    @login_required
    def delete_class(class_id):
        user = current_user()
        class_group = db.session.get(ClassGroup, class_id) or abort(404)
        if not course_owned_by_user(class_group.course, user):
            abort(403)
        class_group.active = False
        db.session.commit()
        flash(
            f"Đã xóa lớp “{class_group.name}” khỏi danh sách. Dữ liệu quét/xem đã ghi trước đó vẫn được giữ nguyên.",
            "success",
        )
        return redirect(url_for("courses_page"))

    @app.post("/courses/<int:course_id>/qr/bulk")
    @login_required
    def create_course_qrs_bulk(course_id):
        user = current_user()
        course = db.session.get(Course, course_id) or abort(404)
        if not course.active or not course_owned_by_user(course, user):
            abort(403)
        links = CourseVideo.query.filter_by(course_id=course.id).order_by(CourseVideo.id).all()
        t0 = time.perf_counter()
        created = 0
        existing = 0
        skipped = 0
        total_processing = 0.0
        for link in links:
            video = link.video
            if not video_is_usable(video) or not video_is_from_connected_channel(video):
                # Không sinh QR mới từ bản ghi video không đúng nguồn của phiên bản cũ.
                skipped += 1
                continue
            # GV mới vẫn sử dụng QR cũ từ kênh GV trước; không tự tạo QR mới
            # cho nguồn YouTube chưa được GV mới kết nối.
            existing = (db.session.get(QrCode, link.qr_id) if link.qr_id else
                        active_qr_for_course_video(course.id, video.id, user.id))
            if video.channel.user_id != user.id and not (existing and existing.active):
                skipped += 1
                continue
            qr, was_created = create_qr_for_course_video(course, video, user, link)
            if was_created:
                created += 1
                total_processing += qr.processing_seconds or 0.0
            else:
                existing += 1
        elapsed = time.perf_counter() - t0
        passed = bool(created or existing or not links)
        record_functional_test(
            "QR_BULK", "Tạo mã QR hàng loạt", passed, owner_id=user.id,
            note=f"Môn học/ mô đun: {course.name}; mới {created}; đã có {existing}; bỏ qua {skipped}."
        )
        record_webapp_experiment(
            user.id, "TAO_QR_HANG_LOAT", elapsed, operations_count=1, success=passed,
            note=f"{len(links)} video; mới {created}; đã có {existing}; bỏ qua {skipped}",
            processing_seconds=(total_processing / created if created else None),
        )
        db.session.commit()
        flash(
            f"Tạo mã QR hoàn tất: {created} mã mới, {existing} video đã có mã, {skipped} video chưa đủ điều kiện sử dụng.",
            "success" if created or existing else "warning",
        )
        # Khi tạo mã từ trang Quản lý mã QR động, giữ người dùng ở chính trang đó.
        # Khi tạo từ trang Môn học/ mô đun, quay lại trang Môn học/ mô đun như trước.
        return_to = request.form.get("return_to", "courses").strip().lower()
        return redirect(url_for("qr_page" if return_to == "qr" else "courses_page"))

    @app.post("/courses/<int:course_id>/playlist-qr/bulk")
    @login_required
    def create_course_playlist_qrs_bulk(course_id):
        user = current_user()
        course = db.session.get(Course, course_id) or abort(404)
        if not course.active or not course_owned_by_user(course, user):
            abort(403)
        links = CoursePlaylist.query.filter_by(course_id=course.id, active=True).order_by(CoursePlaylist.id).all()
        t0 = time.perf_counter()
        created = 0
        existing = 0
        skipped = 0
        total_processing = 0.0
        for link in links:
            playlist = link.playlist
            if not playlist or not playlist.is_available:
                skipped += 1
                continue
            existing = (db.session.get(QrCode, link.qr_id) if link.qr_id else
                        QrCode.query.filter_by(course_id=course.id, playlist_id=playlist.id, active=True).first())
            if playlist.channel.user_id != user.id and not (existing and existing.active):
                skipped += 1
                continue
            qr, was_created = create_qr_for_course_playlist(course, playlist, user, link)
            if was_created:
                created += 1
                total_processing += qr.processing_seconds or 0.0
            else:
                existing += 1
        elapsed = time.perf_counter() - t0
        passed = bool(created or existing or not links)
        record_functional_test(
            "QR_PLAYLIST_BULK", "Tạo mã QR hàng loạt cho playlist", passed, owner_id=user.id,
            note=f"Môn học/ mô đun: {course.name}; mới {created}; đã có {existing}; bỏ qua {skipped}."
        )
        # QR cho nhiều playlist thuộc cùng nhóm nhiệm vụ đối chứng tạo QR hàng loạt.
        record_webapp_experiment(
            user.id, "TAO_QR_HANG_LOAT", elapsed, operations_count=1, success=passed,
            note=f"Playlist môn học/ mô đun {course.name}: tổng {len(links)}; mới {created}; đã có {existing}; bỏ qua {skipped}",
            processing_seconds=(total_processing / created if created else None),
        )
        db.session.commit()
        flash(
            f"Tạo QR playlist hoàn tất: {created} mã mới, {existing} playlist đã có mã, {skipped} playlist không còn khả dụng.",
            "success" if created or existing else "warning",
        )
        return_to = request.form.get("return_to", "courses").strip().lower()
        return redirect(url_for("qr_page" if return_to == "qr" else "courses_page"))

    @app.get("/courses/<int:course_id>/qr.zip")
    @login_required
    def download_course_qrs(course_id):
        user = current_user()
        course = db.session.get(Course, course_id) or abort(404)
        if not course_owned_by_user(course, user):
            abort(403)
        # Bao gồm mọi QR đang hoạt động thuộc môn học/mô đun (cả QR đã tách
        # khỏi liên kết học liệu sau chỉnh sửa playlist). Không bỏ sót QR đã phát hành.
        qrs = (QrCode.query.filter_by(course_id=course.id, active=True)
               .order_by(QrCode.created_at, QrCode.id).all())
        bio = io.BytesIO()
        used_names = set()
        with zipfile.ZipFile(bio, "w", zipfile.ZIP_DEFLATED) as zf:
            for qr in qrs:
                img = generate_branded_qr(qr)
                png = io.BytesIO()
                img.save(png, format="PNG")
                name = qr_download_name(qr)
                if name in used_names:
                    stem, ext = os.path.splitext(name)
                    name = f"{stem}_{qr.slug}{ext}"
                used_names.add(name)
                zf.writestr(name, png.getvalue())
        bio.seek(0)
        filename = f"Ma_QR_{safe_filename(course.code)}.zip"
        return send_file(bio, mimetype="application/zip", as_attachment=True, download_name=filename)


    @app.route("/qr", methods=["GET", "POST"])
    @login_required
    def qr_page():
        user = current_user()
        if request.method == "POST":
            action_started = time.perf_counter()
            video_id = request.form.get("video_id", type=int)
            course_id = request.form.get("course_id", type=int)
            video = visible_video_query(user).filter(Video.id == video_id).first() if video_id else None
            course = db.session.get(Course, course_id) if course_id else None
            if not video or not video_is_usable(video):
                flash("Vui lòng chọn một video có thể sử dụng.", "warning")
                return redirect(url_for("qr_page"))
            if course and not course_owned_by_user(course, user):
                abort(403)
            # Tên mã QR luôn đúng bằng tên video YouTube, giữ nguyên tiếng Việt có dấu.
            title = (video.title or "Video")[:255]
            if course:
                link, _ = ensure_course_video_link(course, video, user)
                qr, was_created = create_qr_for_course_video(course, video, user, link)
                qr.title = title
                elapsed = time.perf_counter() - action_started
                record_functional_test(
                    "QR_SINGLE", "Tạo một mã QR", True, owner_id=user.id,
                    note=f"Video: {video.title}; môn học/ mô đun: {course.name}; tạo mới: {was_created}."
                )
                record_webapp_experiment(
                    user.id, "TAO_QR_RIENG", elapsed, operations_count=1, success=True,
                    note=f"Video: {video.title}; tạo mới: {was_created}",
                    processing_seconds=(qr.processing_seconds if was_created else 0.0),
                )
                db.session.commit()
                if was_created:
                    flash("Đã tạo mã QR động thành công.", "success")
                else:
                    flash("Video này đã có mã QR trong môn học/ mô đun; hệ thống không tạo trùng.", "info")
            else:
                t0 = time.perf_counter()
                slug = secrets.token_urlsafe(8).replace("-", "").replace("_", "")[:12]
                qr = QrCode(owner_id=user.id, course_id=None, slug=slug, title=title, target_type="video")
                db.session.add(qr)
                db.session.flush()
                db.session.add(QrTargetHistory(qr_id=qr.id, video_id=video.id, changed_by=user.id))
                qr.processing_seconds = round(time.perf_counter() - t0, 4)
                elapsed = time.perf_counter() - action_started
                record_functional_test(
                    "QR_SINGLE", "Tạo một mã QR", True, owner_id=user.id, note=f"Video: {video.title}; không gắn môn học/ mô đun."
                )
                record_webapp_experiment(
                    user.id, "TAO_QR_RIENG", elapsed, operations_count=1, success=True,
                    note=f"Video: {video.title}; không gắn môn học/ mô đun", processing_seconds=qr.processing_seconds
                )
                db.session.commit()
                flash("Đã tạo mã QR động thành công.", "success")
            return redirect(url_for("qr_page"))

        selected_course_id = request.args.get("course_id", type=int)
        courses = Course.query.filter_by(owner_id=user.id, active=True).order_by(Course.name).all()
        selected_course = None
        if selected_course_id is not None and selected_course_id < 0:
            abort(404)
        if selected_course_id and selected_course_id > 0:
            selected_course = next((c for c in courses if c.id == selected_course_id), None)
            if not selected_course:
                abort(403)
        all_qrs = (QrCode.query.filter(managed_qr_scope(user.id), QrCode.active.is_(True))
                   .order_by(QrCode.created_at.desc()).all())
        if selected_course_id is None:
            qrs = all_qrs
        elif selected_course_id == 0:
            qrs = [q for q in all_qrs if q.course_id is None]
        else:
            qrs = [q for q in all_qrs if q.course_id == selected_course_id]
        # Đồng bộ tên QR theo đúng tiêu đề nguồn hiện tại: video hoặc playlist.
        title_changed = False
        for existing_qr in all_qrs:
            if existing_qr.is_playlist:
                if existing_qr.playlist and existing_qr.playlist.title:
                    exact_title = existing_qr.playlist.title[:255]
                else:
                    exact_title = existing_qr.title
            else:
                target = existing_qr.current_target
                exact_title = target.video.title[:255] if target and target.video and target.video.title else existing_qr.title
            if existing_qr.title != exact_title:
                existing_qr.title = exact_title
                title_changed = True
        if title_changed:
            db.session.commit()

        qr_needs_review = set()
        for q in qrs:
            if q.course_id and not q.is_playlist:
                target = active_video_for_qr(q)
                if target and (not video_is_from_connected_channel(target) or
                               not CourseVideo.query.filter_by(course_id=q.course_id, video_id=target.id).first()):
                    qr_needs_review.add(q.id)
        # Danh sách video thay thế vẫn cần cho nút Đổi video đích trong Kho mã QR.
        videos = visible_video_query(user).filter(
            Video.is_available.is_(True), Video.privacy_status != "private",
            Video.embeddable.isnot(False)
        ).order_by(Video.title).all()
        return render_template(
            "qr.html", qr_needs_review=qr_needs_review, qrs=qrs, videos=videos, courses=courses,
            selected_course=selected_course, selected_course_id=selected_course_id,
            total_qrs=len(all_qrs),
        )

    @app.post("/qr/playlist")
    @login_required
    def create_playlist_qr():
        user = current_user()
        link_id = request.form.get("course_playlist_id", type=int)
        link = db.session.get(CoursePlaylist, link_id) if link_id else None
        if not link or not link.active or not link.course.active or not course_owned_by_user(link.course, user):
            abort(403)
        if not link.playlist or not link.playlist.is_available:
            flash("Playlist hiện không còn khả dụng trên YouTube. Hãy đồng bộ lại kênh hoặc chọn playlist khác.", "warning")
            return redirect(url_for("qr_page"))
        if link.playlist.channel.user_id != user.id and not (link.qr and link.qr.active):
            flash("Playlist này thuộc kênh của giảng viên trước, không thể tạo QR mới khi chưa kết nối nguồn. Các QR cũ vẫn hoạt động.", "warning")
            return redirect(url_for("qr_page"))
        qr, was_created = create_qr_for_course_playlist(link.course, link.playlist, user, link)
        record_functional_test(
            "QR_PLAYLIST_SINGLE", "Tạo mã QR cho một playlist", True, owner_id=user.id,
            note=f"Môn học/ mô đun: {link.course.name}; playlist: {link.playlist.title}; tạo mới: {was_created}."
        )
        db.session.commit()
        flash("Đã tạo mã QR playlist thành công." if was_created else "Playlist này đã có mã QR trong môn học/ mô đun; hệ thống không tạo trùng.", "success" if was_created else "info")
        return redirect(url_for("qr_page"))


    def qr_action_destination(qr, fallback="qr"):
        """Redirect back to originating library row, never to a supplied external URL."""
        if request.form.get("return_to") == "videos":
            row = request.form.get("library_row", "")
            if re.fullmatch(r"(?:video|playlist)-[0-9]+", row):
                return videos_library_return(row, request.form.get("return_query", ""))
        return url_for("qr_page")

    @app.post("/qr/<int:qr_id>/retarget")
    @login_required
    def qr_retarget(qr_id):
        qr = db.session.get(QrCode, qr_id) or abort(404)
        user = current_user()
        if not qr_managed_by_user(qr, user):
            abort(403)
        if not qr.active:
            abort(404)
        if qr.is_playlist:
            flash("Mã QR playlist tự cập nhật theo nội dung playlist sau mỗi lần đồng bộ YouTube; không đổi video đích theo cách này.", "info")
            return redirect(qr_action_destination(qr))
        video_id = request.form.get("video_id", type=int)
        video = visible_video_query(user).filter(Video.id == video_id).first() if video_id else None
        if not video:
            abort(403)
        if not video_is_usable(video):
            flash("Video mới chưa đủ điều kiện sử dụng.", "warning")
            return redirect(qr_action_destination(qr))
        linked = CourseVideo.query.filter_by(qr_id=qr.id).first()
        if qr.course_id and not linked:
            course = db.session.get(Course, qr.course_id)
            if not course or course.owner_id != user.id:
                abort(403)
            target_link = CourseVideo.query.filter_by(course_id=course.id, video_id=video.id).first()
            if not target_link:
                target_link = CourseVideo(course_id=course.id, video_id=video.id,
                                          added_by=user.id, manual_added=True)
                db.session.add(target_link)
                db.session.flush()
            if target_link.qr_id and target_link.qr_id != qr.id:
                flash("Video đích đã có QR riêng ở môn học này.", "warning")
                return redirect(qr_action_destination(qr))
            target_link.manual_added = True
            target_link.qr_id = qr.id
            linked = target_link
        if linked and linked.video_id != video.id:
            duplicate = CourseVideo.query.filter_by(course_id=linked.course_id, video_id=video.id).first()
            if duplicate and duplicate.id != linked.id:
                if duplicate.qr_id and duplicate.qr_id != qr.id:
                    flash("Video mới đã có một mã QR khác trong cùng môn học/ mô đun; không thể tạo liên kết trùng.", "warning")
                    return redirect(qr_action_destination(qr))
                db.session.delete(duplicate)
                db.session.flush()
            linked.video_id = video.id
            linked.manual_added = True  # GV đã chọn đích thủ công; không tự xóa khi playlist thay đổi.
        now = utcnow()
        old = qr.current_target
        if old:
            old.active_to = now
        db.session.add(QrTargetHistory(qr_id=qr.id, video_id=video.id, active_from=now, changed_by=user.id))
        qr.title = (video.title or qr.title)[:255]
        record_functional_test(
            "QR_RETARGET", "Đổi video đích của mã QR", True, owner_id=user.id,
            note=f"QR {qr.slug}; video mới: {video.title}."
        )
        db.session.commit()
        flash("Đã đổi video đích. Mã QR và đường dẫn không thay đổi.", "success")
        return redirect(qr_action_destination(qr))

    @app.post("/qr/<int:qr_id>/delete")
    @login_required
    def delete_qr(qr_id):
        """Ngừng sử dụng một QR nhưng giữ log/lịch sử để phục vụ nghiên cứu."""
        qr = db.session.get(QrCode, qr_id) or abort(404)
        user = current_user()
        if not qr_managed_by_user(qr, user):
            abort(403)
        if not qr.active:
            flash("Mã QR này đã được xóa khỏi danh sách sử dụng.", "info")
            return redirect(qr_action_destination(qr))
        qr.active = False
        # Gỡ liên kết hiện hành để video trong môn học/ mô đun có thể tạo một QR mới nếu cần.
        for link in CourseVideo.query.filter_by(qr_id=qr.id).all():
            link.qr_id = None
        for link in CoursePlaylist.query.filter_by(qr_id=qr.id).all():
            link.qr_id = None
        db.session.commit()
        flash(
            f"Đã xóa mã QR “{qr.title}” khỏi danh sách sử dụng. Lịch sử quét/xem trước đây vẫn được giữ nguyên.",
            "success",
        )
        return redirect(qr_action_destination(qr))

    @app.post("/qr/delete-all")
    @login_required
    def delete_all_qrs():
        user = current_user()
        qrs = QrCode.query.filter(managed_qr_scope(user.id), QrCode.active.is_(True)).all()
        if not qrs:
            flash("Không có mã QR nào để xóa.", "info")
            return redirect(url_for("qr_page"))
        qr_ids = [q.id for q in qrs]
        for q in qrs:
            q.active = False
        if qr_ids:
            for link in CourseVideo.query.filter(CourseVideo.qr_id.in_(qr_ids)).all():
                link.qr_id = None
            for link in CoursePlaylist.query.filter(CoursePlaylist.qr_id.in_(qr_ids)).all():
                link.qr_id = None
        db.session.commit()
        flash(f"Đã xóa {len(qrs)} mã QR khỏi danh sách sử dụng. Lịch sử quét/xem trước đây vẫn được giữ lại.", "success")
        return redirect(url_for("qr_page"))

    @app.get("/qr/<int:qr_id>/png")
    @login_required
    def qr_png(qr_id):
        qr = db.session.get(QrCode, qr_id) or abort(404)
        user = current_user()
        if not qr_managed_by_user(qr, user):
            abort(403)
        if not qr.active:
            abort(404)
        img = generate_branded_qr(qr)
        bio = io.BytesIO()
        img.save(bio, format="PNG")
        bio.seek(0)
        return send_file(bio, mimetype="image/png", as_attachment=True, download_name=qr_download_name(qr))

    @app.get("/qr/download-all.zip")
    @login_required
    def qr_download_all():
        """Xuất ZIP đúng phạm vi đang lọc; khi xem tất cả, chia thư mục theo MH/MĐ.

        Tên thư mục có mã môn để tách được những môn có tên giống nhau.
        Chỉ QR được quản lý theo phân quyền hiện hành mới có trong ZIP.
        """
        user = current_user()
        selected = request.args.get("course_id", type=int)
        if request.args.get("course_id") not in (None, "") and selected is None:
            abort(400)
        if selected is not None and selected < 0:
            abort(404)
        course = None
        if selected and selected > 0:
            course = db.session.get(Course, selected) or abort(404)
            if not course.active or not course_owned_by_user(course, user):
                abort(403)
        qry = QrCode.query.filter(managed_qr_scope(user.id), QrCode.active.is_(True))
        if selected is not None:
            qry = qry.filter(QrCode.course_id == selected if selected > 0 else QrCode.course_id.is_(None))
        qrs = qry.order_by(QrCode.course_id, QrCode.created_at, QrCode.id).all()
        bio = io.BytesIO()
        used_names = set()
        with zipfile.ZipFile(bio, "w", zipfile.ZIP_DEFLATED) as zf:
            for q in qrs:
                # Có thể lọc tất cả: các môn được nhóm theo mã môn, QR chưa gán tách riêng.
                folder = (safe_filename(f"{q.course.code}_{q.course.name}")
                          if q.course else "Chua_gan_mon_hoc")
                img = generate_branded_qr(q)
                png = io.BytesIO()
                img.save(png, format="PNG")
                name = f"{folder}/{qr_download_name(q)}"
                if name in used_names:
                    stem, ext = os.path.splitext(name)
                    name = f"{stem}_{q.slug}{ext}"
                used_names.add(name)
                zf.writestr(name, png.getvalue())
        bio.seek(0)
        filename = (f"Ma_QR_{safe_filename(course.code)}.zip" if course else
                    "Ma_QR_chua_gan_mon.zip" if selected == 0 else "Ma_QR_theo_mon_hoc.zip")
        return send_file(bio, mimetype="application/zip", as_attachment=True, download_name=filename)

    @app.get("/q/<slug>")
    def public_qr(slug):
        qr = QrCode.query.filter_by(slug=slug).first_or_404()
        if not qr.active:
            return render_template("viewer_error.html", qr=qr,
                message="Mã QR này đã ngừng sử dụng. Vui lòng liên hệ giảng viên để nhận học liệu mới."), 410
        signed_in_user = current_user()

        # "Mở thử" dành cho GV/Quản trị viên đã đăng nhập: bỏ qua màn hình nhận diện
        # và không ghi lượt quét/thời lượng xem vào dữ liệu thực nghiệm.
        preview_mode = bool(
            request.args.get("preview") == "1"
            and signed_in_user
            and (signed_in_user.role == "admin" or qr_managed_by_user(qr, signed_in_user))
        )
        tracking_enabled = not preview_mode

        if preview_mode:
            pcode = signed_in_user.lecturer_code or "XEM-THU"
            role = "GV"
            class_code = None
            viewer_identity = "Chế độ xem thử"
        elif signed_in_user and signed_in_user.lecturer_code:
            # Quản trị viên đồng thời là GV. Lượt quét của GV không đưa vào thống kê SV.
            if not signed_in_user.lecturer_code:
                return render_template("viewer_error.html", qr=qr,
                                       message="Vui lòng hoàn tất thông tin giảng viên để được cấp mã định danh tự động."), 403
            pcode = signed_in_user.lecturer_code
            role = "GV"
            class_code = None
            viewer_identity = pcode
            session["participant_code"] = pcode
            session["participant_role"] = "GV"
            session["participant_class_code"] = ""
        elif signed_in_user and signed_in_user.role == "admin":
            # Quản trị viên kiểm tra QR nhưng không tạo dữ liệu người học giả.
            pcode = "QUAN-TRI"
            role = "GV"
            class_code = None
            viewer_identity = "Quản trị viên"
            tracking_enabled = False
        else:
            # Mỗi lượt mở QR mới cần SV tự chọn lớp và tên để hạn chế dùng nhầm
            # phiên người xem trên thiết bị dùng chung. Bước chuyển về từ biểu mẫu
            # nhận diện được đánh dấu một lần trong phiên, không dùng cho QR khác.
            session_code = session.get("participant_code") if session.get("participant_role") == "SV" else None
            just_identified = (session.pop("identified_for_qr", None) == slug)
            resolved_class = None
            if session_code and just_identified:
                school_class = SchoolClass.query.filter_by(code=session.get("participant_class_code"), active=True).first()
                student = Student.query.filter_by(student_code=normalize_student_code(session_code)).first()
                if school_class and student:
                    _, resolved_class, resolve_error = validate_student_selection_for_qr(
                        qr, school_class.id, student.id, school_class.name)
                    if resolve_error:
                        session_code = None
                else:
                    session_code = None
            else:
                session_code = None
            if not session_code:
                session["identify_next"] = url_for("public_qr", slug=slug)
                session["identify_qr_slug"] = slug
                message = session.pop("identify_error", None)
                classes = eligible_classes_for_qr(qr)
                class_choices = [{"id": cls.id, "name": cls.name} for cls in classes]
                return render_template("identify.html", qr=qr, message=message,
                                       class_choices=class_choices)
            pcode = normalize_student_code(session_code)
            role = "SV"
            class_code = resolved_class.code
            viewer_identity = pcode
            session["participant_class_code"] = class_code
            session["participant_qr_slug"] = slug

        if qr.is_playlist:
            playlist = qr.playlist
            links = PlaylistVideo.query.filter_by(playlist_id=playlist.id).order_by(PlaylistVideo.position.asc().nullslast(), PlaylistVideo.id).all() if playlist else []
            videos = [link.video for link in links if video_is_usable(link.video)]
            unavailable_count = max(0, len(links) - len(videos))
            success = bool(playlist and playlist.is_available and videos)
            if tracking_enabled:
                cutoff = utcnow() - timedelta(seconds=30)
                existing = ScanLog.query.filter(
                    ScanLog.participant_code == pcode,
                    ScanLog.qr_id == qr.id,
                    ScanLog.scanned_at >= cutoff,
                ).first()
                if not existing:
                    db.session.add(ScanLog(
                        participant_code=pcode, role=role, class_code=class_code,
                        course_id=qr.course_id, qr_id=qr.id, video_id=None, success=success,
                    ))
                    record_functional_test(
                        "QR_PLAYLIST_SCAN", "Quét mã QR playlist và mở học liệu", success, owner_id=(qr.course.owner_id if qr.course else qr.owner_id),
                        note=(f"{playlist.title if playlist else qr.title} – "
                              f"{qr.course.name if qr.course else 'Chưa xác định môn học/mô đun'}"),
                        participant_code=pcode if role == "SV" else None,
                        participant_class_code=class_code if role == "SV" else None,
                    )
                    db.session.commit()
            if not playlist:
                return render_template("viewer_error.html", qr=qr, message="Mã QR playlist chưa có playlist đích. Vui lòng liên hệ giảng viên phụ trách."), 404
            if not playlist.is_available:
                return render_template("viewer_error.html", qr=qr, message="Playlist đích đã bị xóa hoặc không còn khả dụng trên YouTube. Mã QR vẫn được giữ nguyên để giảng viên xử lý trong hệ thống."), 409
            if not videos:
                return render_template("viewer_error.html", qr=qr, message="Playlist hiện chưa có video nào đủ điều kiện phát trong trang học liệu. Vui lòng liên hệ giảng viên phụ trách."), 409
            return render_template(
                "playlist_viewer.html", qr=qr, playlist=playlist, videos=videos,
                unavailable_count=unavailable_count, preview_mode=preview_mode,
                tracking_enabled=tracking_enabled, viewer_identity=viewer_identity,
            )

        video = active_video_for_qr(qr)
        success = bool(video and video.is_available and video.embeddable is not False and video.privacy_status != "private")
        if tracking_enabled:
            cutoff = utcnow() - timedelta(seconds=30)
            existing = ScanLog.query.filter(
                ScanLog.participant_code == pcode,
                ScanLog.qr_id == qr.id,
                ScanLog.scanned_at >= cutoff,
            ).first()
            if not existing:
                db.session.add(ScanLog(
                    participant_code=pcode, role=role, class_code=class_code,
                    course_id=qr.course_id,
                    qr_id=qr.id,
                    video_id=video.id if video else None,
                    success=success,
                ))
                record_functional_test(
                    "QR_SCAN", "Quét mã QR và mở học liệu", success, owner_id=(qr.course.owner_id if qr.course else qr.owner_id),
                    note=(f"{video.title if video else qr.title} – "
                          f"{qr.course.name if qr.course else 'Chưa xác định môn học/mô đun'}"),
                    participant_code=pcode if role == "SV" else None,
                    participant_class_code=class_code if role == "SV" else None,
                )
                db.session.commit()

        if not video:
            return render_template("viewer_error.html", qr=qr, message="Mã QR chưa được gán video đích. Vui lòng liên hệ giảng viên phụ trách."), 404
        if not success:
            if not video.is_available:
                message = "Video đích đã bị xóa hoặc không còn tồn tại trên YouTube. Mã QR vẫn được giữ nguyên để giảng viên có thể thay video mới mà không phải phát hành lại mã QR."
            elif video.privacy_status == "private":
                message = "Video đích hiện đang ở chế độ riêng tư. Mã QR vẫn được giữ nguyên; giảng viên cần đổi quyền video hoặc cập nhật video thay thế."
            elif video.embeddable is False:
                message = "Video đích hiện không cho phép phát nhúng. Mã QR vẫn được giữ nguyên; giảng viên có thể cập nhật một video thay thế phù hợp."
            else:
                message = f"Video hiện không thể phát trong trang học liệu ({video.status_label}). Mã QR vẫn được giữ nguyên để có thể cập nhật video đích."
            return render_template("viewer_error.html", qr=qr, message=message), 409
        return render_template(
            "viewer.html", qr=qr, video=video, preview_mode=preview_mode,
            tracking_enabled=tracking_enabled, viewer_identity=viewer_identity,
        )


    @app.get("/q/<slug>/students")
    def qr_class_students(slug):
        """Chỉ trả họ tên, mã SV đang học của một lớp được gán cho QR hợp lệ."""
        qr = QrCode.query.filter_by(slug=slug, active=True).first_or_404()
        class_id = request.args.get("class_id", type=int)
        allowed_ids = {row.id for row in eligible_classes_for_qr(qr)}
        if not class_id or class_id not in allowed_ids:
            response = jsonify({"students": [], "error": "Lớp không thuộc học liệu được giao."})
            response.status_code = 400
        else:
            school_class = db.session.get(SchoolClass, class_id)
            students = current_students_in_class(school_class)
            response = jsonify({"students": [
                {"id": student.id, "name": student.full_name,
                 "short_code": (re.sub(r"\D", "", student.student_code) or student.student_code)[-3:]}
                for student in students
            ]})
        response.headers["Cache-Control"] = "no-store"
        return response

    @app.post("/identify")
    def identify():
        nxt = session.get("identify_next") or url_for("index")
        slug = session.get("identify_qr_slug")
        qr = QrCode.query.filter_by(slug=slug, active=True).first() if slug else None
        if not qr:
            session["identify_error"] = "Không xác định được mã QR đang truy cập. Vui lòng quét lại mã QR."
            return redirect(nxt)
        class_id = request.form.get("school_class_id", type=int)
        student_id = request.form.get("student_id", type=int)
        class_name = request.form.get("class_name", "")
        if not class_id or not student_id:
            session["identify_error"] = "Vui lòng chọn tên lớp và họ tên sinh viên trong danh sách."
            return redirect(nxt)
        student, school_class, error = validate_student_selection_for_qr(
            qr, class_id, student_id, class_name)
        if error:
            session["identify_error"] = error
            return redirect(nxt)
        session["participant_code"] = student.student_code
        session["participant_role"] = "SV"
        session["participant_class_code"] = school_class.code
        session["participant_qr_slug"] = slug
        session["identified_for_qr"] = slug
        session.pop("identify_qr_slug", None)
        session.pop("identify_error", None)
        session.pop("identify_next", None)
        return redirect(nxt)

    @app.post("/participant/reset")
    def participant_reset():
        for k in ("participant_code", "participant_role", "participant_class_code",
                  "participant_qr_slug", "identified_for_qr"):
            session.pop(k, None)
        return redirect(request.referrer or url_for("index"))

    @app.post("/api/watch/start")
    def watch_start():
        data = request.get_json(silent=True) or {}
        slug = data.get("slug")
        qr = QrCode.query.filter_by(slug=slug, active=True).first_or_404()
        if not session.get("participant_code"):
            return jsonify({"error": "missing_participant_or_video"}), 400
        if session.get("participant_role") == "SV" and session.get("participant_qr_slug") != slug:
            return jsonify({"error": "Vui lòng xác nhận lại lớp và tên sinh viên cho mã QR này."}), 403

        if qr.is_playlist:
            video_id = data.get("video_id")
            try:
                video_id = int(video_id)
            except (TypeError, ValueError):
                return jsonify({"error": "missing_playlist_video"}), 400
            link = (PlaylistVideo.query
                    .filter_by(playlist_id=qr.playlist_id, video_id=video_id)
                    .first())
            video = link.video if link else None
            if not video_is_usable(video):
                return jsonify({"error": "invalid_playlist_video"}), 400
        else:
            video = active_video_for_qr(qr)
            if not video:
                return jsonify({"error": "missing_participant_or_video"}), 400

        ws = WatchSession(
            session_key=secrets.token_urlsafe(24),
            participant_code=session["participant_code"],
            role=session.get("participant_role", "SV"),
            class_code=session.get("participant_class_code"),
            course_id=qr.course_id,
            qr_id=qr.id,
            video_id=video.id,
            duration_seconds=float(data.get("duration") or 0),
        )
        db.session.add(ws)
        db.session.commit()
        return jsonify({"session_key": ws.session_key})


    @app.post("/api/watch/event")
    def watch_event():
        data = request.get_json(silent=True) or {}
        key = data.get("session_key")
        ws = WatchSession.query.filter_by(session_key=key).first_or_404()
        if session.get("participant_code") != ws.participant_code:
            abort(403)
        event_type = (data.get("event_type") or "heartbeat").lower()[:24]
        state = (data.get("state") or "unknown").lower()[:16]
        position = max(0.0, float(data.get("position") or 0))
        duration = max(0.0, float(data.get("duration") or ws.duration_seconds or 0))
        now = utcnow()
        elapsed = max(0.0, (now - ws.last_event_at).total_seconds()) if ws.last_event_at else 0.0
        position_delta = position - (ws.last_position or 0.0)
        counted = 0.0

        # Chỉ cộng khi phiên trước đang phát, thời gian thực hợp lý và không có bước nhảy tua lớn.
        if ws.last_state == "playing" and elapsed <= max(15, current_app.config["HEARTBEAT_SECONDS"] * 3):
            if -1.0 <= position_delta <= elapsed + 3.0:
                counted = max(0.0, min(elapsed, position_delta + 0.5))
                ws.estimated_watch_seconds += counted

        ws.max_position = max(ws.max_position or 0.0, position)
        ws.duration_seconds = max(ws.duration_seconds or 0.0, duration)
        ws.last_event_at = now
        ws.last_position = position
        ws.last_state = state

        if ws.duration_seconds > 0:
            ws.completion_percent = min(100.0, 100.0 * ws.estimated_watch_seconds / ws.duration_seconds)
        if event_type == "ended" or state == "ended":
            ws.completed = True
            ws.ended_at = now
            ws.completion_percent = max(ws.completion_percent, 100.0)
        elif ws.duration_seconds > 0 and ws.estimated_watch_seconds / ws.duration_seconds >= current_app.config["COMPLETION_THRESHOLD"]:
            ws.completed = True

        db.session.add(WatchEvent(
            watch_session_id=ws.id,
            event_type=event_type,
            player_state=state,
            position=position,
            duration=duration,
            counted_seconds=counted,
        ))
        db.session.commit()
        return jsonify({
            "ok": True,
            "estimated_watch_seconds": round(ws.estimated_watch_seconds, 1),
            "max_position": round(ws.max_position, 1),
            "completion_percent": round(ws.completion_percent, 1),
            "completed": ws.completed,
        })

    @app.get("/analytics")
    @login_required
    def analytics():
        """Hoạt động GV/Quản trị viên và một bảng tổng hợp lượt quét - xem của SV.

        Không nhân bản dữ liệu trong PostgreSQL; đối chiếu phiên xem với lượt quét
        gần nhất trước thời điểm bắt đầu phiên của cùng SV, lớp và QR.
        """
        user = current_user()
        course_q = Course.query.filter(Course.active.is_(True))
        qr_q = QrCode.query.filter(QrCode.active.is_(True), QrCode.course_id.isnot(None))
        if user.role != "admin":
            course_q = course_q.filter(Course.owner_id == user.id)
            qr_q = qr_q.filter(managed_qr_scope(user.id))
        course_options = course_q.order_by(Course.name).all()
        qr_options = qr_q.order_by(QrCode.course_id, QrCode.title).all()
        if user.role == "admin":
            class_options = SchoolClass.query.filter(SchoolClass.active.is_(True)).order_by(SchoolClass.name).all()
        else:
            class_options = (SchoolClass.query.join(CourseClass, CourseClass.school_class_id == SchoolClass.id)
                .join(Course, CourseClass.course_id == Course.id)
                .filter(Course.owner_id == user.id, Course.active.is_(True),
                        CourseClass.active.is_(True), SchoolClass.active.is_(True))
                .distinct().order_by(SchoolClass.name).all())
        class_name_map = {c.code: c.name for c in SchoolClass.query.all()}
        for legacy in ClassGroup.query.all():
            class_name_map.setdefault(legacy.code, legacy.name)

        # Bộ lọc theo tên học liệu, lớp, môn học và thời gian; mã QR chỉ dùng nội bộ.
        course_id = request.args.get("course_id", type=int)
        class_code = (request.args.get("class_code") or "").strip()
        qr_id = request.args.get("qr_id", type=int)
        status = (request.args.get("status") or "all").strip()
        if status not in {"all", "not_scanned", "not_completed", "unwatched", "completed", "failed"}:
            status = "all"
        selected_course = next((c for c in course_options if c.id == course_id), None)
        selected_qr = next((q for q in qr_options if q.id == qr_id), None)
        selected_class = next((c for c in class_options if c.code == class_code), None)
        media_search = (request.args.get("media_search") or "").strip()[:120]
        sv_period = (request.args.get("sv_period") or "all").strip()
        if sv_period not in {"all", "today", "7", "30", "custom"}:
            sv_period = "30"
        sv_from, sv_to = (request.args.get("sv_from") or "").strip(), (request.args.get("sv_to") or "").strip()
        sv_start_at, sv_end_at = None, None
        if sv_period == "today":
            sv_start_at = datetime.combine(utcnow().date(), datetime.min.time())
        elif sv_period in {"7", "30"}:
            sv_start_at = utcnow() - timedelta(days=int(sv_period))
        elif sv_period == "custom":
            try:
                sv_start_at = datetime.combine(date.fromisoformat(sv_from), datetime.min.time()) if sv_from else None
                sv_end_at = datetime.combine(date.fromisoformat(sv_to) + timedelta(days=1), datetime.min.time()) if sv_to else None
                if sv_start_at and sv_end_at and sv_end_at <= sv_start_at:
                    raise ValueError("Ngày kết thúc phải từ ngày bắt đầu trở đi")
            except ValueError:
                flash("Khoảng ngày tra cứu của sinh viên không hợp lệ; vui lòng chọn lại.", "warning")
                sv_period, sv_start_at, sv_end_at = "all", None, None
        sv_page = max(1, request.args.get("sv_page", default=1, type=int))
        if class_code and not selected_class:
            # Không cho dùng mã lớp tùy ý ngoài quyền người dùng.
            class_code = "__unauthorized__"
        # Truy vấn tại máy chủ: chỉ tải 10 dòng đang xem, không đưa toàn bộ
        # nhật ký toàn trường vào bộ nhớ hay vào trang HTML.
        test_period = (request.args.get("sys_period") or "30").strip()
        test_feature = (request.args.get("sys_feature") or "").strip()
        test_result = (request.args.get("sys_result") or "all").strip()
        sys_actor_id = request.args.get("sys_actor", type=int)
        if user.role != "admin" and sys_actor_id is not None and sys_actor_id != user.id:
            abort(403)
        test_search = ""  # Không còn tra cứu theo học liệu ở mục hoạt động hệ thống.
        actor_query = User.query.filter(User.id.in_(
            db.session.query(FunctionalTestRun.owner_id).filter(
                FunctionalTestRun.source == "automatic",
                FunctionalTestRun.participant_code.is_(None),
                ~FunctionalTestRun.case_code.in_(["QR_SCAN", "QR_PLAYLIST_SCAN"]),
                FunctionalTestRun.owner_id.isnot(None)
            )))
        if user.role != "admin":
            actor_query = actor_query.filter(User.id == user.id)
        system_actors = actor_query.order_by(User.name, User.email).all()
        test_from = (request.args.get("sys_from") or "").strip()
        test_to = (request.args.get("sys_to") or "").strip()
        test_page = max(1, request.args.get("sys_page", 1, type=int))
        test_query = FunctionalTestRun.query.filter_by(source="automatic").filter(
            FunctionalTestRun.participant_code.is_(None),
            ~FunctionalTestRun.case_code.in_(["QR_SCAN", "QR_PLAYLIST_SCAN"]),
        )
        if user.role != "admin":
            test_query = test_query.filter(FunctionalTestRun.owner_id == user.id)
        now = utcnow()
        if test_period == "today":
            test_query = test_query.filter(FunctionalTestRun.created_at >= datetime.combine(now.date(), datetime.min.time()))
        elif test_period in {"7", "30"}:
            test_query = test_query.filter(FunctionalTestRun.created_at >= now - timedelta(days=int(test_period)))
        elif test_period == "custom":
            try:
                if test_from:
                    test_query = test_query.filter(FunctionalTestRun.created_at >= datetime.combine(date.fromisoformat(test_from), datetime.min.time()))
                if test_to:
                    test_query = test_query.filter(FunctionalTestRun.created_at < datetime.combine(date.fromisoformat(test_to) + timedelta(days=1), datetime.min.time()))
            except ValueError:
                flash("Ngày lọc không hợp lệ; vui lòng chọn lại.", "warning")
        elif test_period != "all":
            test_period = "30"
            test_query = test_query.filter(FunctionalTestRun.created_at >= now - timedelta(days=30))
        if test_feature:
            test_query = test_query.filter(FunctionalTestRun.case_name == test_feature)
        if test_result == "pass":
            test_query = test_query.filter(FunctionalTestRun.passed.is_(True))
        elif test_result == "fail":
            test_query = test_query.filter(FunctionalTestRun.passed.is_(False))
        if sys_actor_id is not None:
            test_query = test_query.filter(FunctionalTestRun.owner_id == sys_actor_id)
        tests_total = test_query.count()
        tests_pass = test_query.filter(FunctionalTestRun.passed.is_(True)).count()
        tests_fail = tests_total - tests_pass
        test_rate = round(100 * tests_pass / tests_total, 1) if tests_total else None
        test_pages = max(1, (tests_total + 9) // 10)
        test_page = min(test_page, test_pages)
        tests = test_query.order_by(FunctionalTestRun.created_at.desc(), FunctionalTestRun.id.desc()).offset((test_page - 1) * 10).limit(10).all()
        test_page_numbers = list(range(max(1, test_page-2), min(test_pages, test_page+2)+1))
        test_feature_options = [row[0] for row in db.session.query(FunctionalTestRun.case_name)
                                .filter(FunctionalTestRun.source == "automatic", FunctionalTestRun.participant_code.is_(None),
                                        ~FunctionalTestRun.case_code.in_(["QR_SCAN", "QR_PLAYLIST_SCAN"]))
                                .distinct().order_by(FunctionalTestRun.case_name).all()]
        test_rows = [{"test": t, "actor": (f"{t.owner.name or 'Giảng viên'} ({t.owner.email})" if t.owner else "Chưa xác định"),
                      "evidence": t.note or "—"} for t in tests]
        test_filter_args = {"sys_period": test_period, "sys_feature": test_feature,
                            "sys_result": test_result, "sys_actor": sys_actor_id or "",
                            "sys_from": test_from, "sys_to": test_to}

        # Không tính bản ghi quét mã hoặc xem video của SV vào hoạt động hệ thống.
        # Hai loại sự kiện này chỉ xuất hiện trong mục 2.
        
        pending_rows, pending_notice, participant_stats = [], "", None
        if status in {"not_scanned", "not_completed"}:
            selected_class = next((c for c in class_options if c.code == class_code), None)
            if not selected_class or not selected_qr:
                pending_notice = "Vui lòng chọn lớp và tên học liệu để xác định chính xác danh sách sinh viên chưa tham gia."
            elif not selected_qr.course_id:
                pending_notice = "Học liệu chưa thuộc môn học/mô đun nên chưa xác định được danh sách lớp áp dụng."
            elif selected_course and selected_qr.course_id != selected_course.id:
                pending_notice = "Học liệu không thuộc môn học/mô đun đang chọn."
            elif not CourseClass.query.filter_by(course_id=selected_qr.course_id,
                                                  school_class_id=selected_class.id, active=True).first():
                pending_notice = "Lớp chưa được gán cho môn học/mô đun của học liệu đã chọn."
            else:
                today = utcnow().date()
                roster = (db.session.query(Student).join(ClassEnrollment, ClassEnrollment.student_id == Student.id)
                          .filter(ClassEnrollment.school_class_id == selected_class.id,
                                  ClassEnrollment.started_on <= today,
                                  or_(ClassEnrollment.ended_on.is_(None), ClassEnrollment.ended_on > today))
                          .distinct().order_by(Student.student_code).all())
                roster_codes = [s.student_code for s in roster]
                valid_scan_codes = set()
                completed_pairs = set()
                if roster_codes:
                    roster_scans = db.session.query(ScanLog.participant_code).filter(
                        ScanLog.role == "SV", ScanLog.qr_id == selected_qr.id,
                        ScanLog.class_code == selected_class.code,
                        ScanLog.success.is_(True), ScanLog.participant_code.in_(roster_codes))
                    roster_watches = db.session.query(WatchSession.participant_code, WatchSession.video_id).filter(
                        WatchSession.role == "SV", WatchSession.qr_id == selected_qr.id,
                        WatchSession.class_code == selected_class.code,
                        WatchSession.completed.is_(True), WatchSession.participant_code.in_(roster_codes))
                    if sv_start_at:
                        roster_scans = roster_scans.filter(ScanLog.scanned_at >= sv_start_at)
                        roster_watches = roster_watches.filter(WatchSession.started_at >= sv_start_at)
                    if sv_end_at:
                        roster_scans = roster_scans.filter(ScanLog.scanned_at < sv_end_at)
                        roster_watches = roster_watches.filter(WatchSession.started_at < sv_end_at)
                    valid_scan_codes = {r[0] for r in roster_scans.distinct().all()}
                    completed_pairs = {(r[0], r[1]) for r in roster_watches.distinct().all()}
                target_video_names = {}
                if selected_qr.is_playlist:
                    for pv in PlaylistVideo.query.filter_by(playlist_id=selected_qr.playlist_id).all():
                        if video_is_usable(pv.video):
                            target_video_names[pv.video_id] = pv.video.title
                    target_video_ids = list(target_video_names)
                else:
                    video = active_video_for_qr(selected_qr)
                    target_video_ids = [video.id] if video and video_is_usable(video) else []
                    if target_video_ids:
                        target_video_names[video.id] = video.title
                complete_count = sum(1 for code in roster_codes
                                     if code in valid_scan_codes and target_video_ids
                                     and all((code, vid) in completed_pairs for vid in target_video_ids))
                participant_stats = {"roster": len(roster), "scanned": len(valid_scan_codes),
                                     "not_scanned": len(roster) - len(valid_scan_codes),
                                     "completed": complete_count,
                                     "not_completed": max(0, len(valid_scan_codes) - complete_count),
                                     "video_total": len(target_video_ids)}
                if status == "not_completed" and not target_video_ids:
                    pending_notice = "Học liệu hiện không có video khả dụng; chưa thể xác định tình trạng hoàn thành."
                else:
                    for student in roster:
                        scanned = student.student_code in valid_scan_codes
                        count_completed = sum((student.student_code, vid) in completed_pairs for vid in target_video_ids)
                        if status == "not_scanned" and not scanned:
                            pending_rows.append({"student": student, "class_name": selected_class.name,
                                                 "course_name": selected_qr.course.name if selected_qr.course else "—",
                                                 "qr_title": selected_qr.title, "status": "Chưa quét thành công"})
                        elif status == "not_completed" and scanned and count_completed < len(target_video_ids):
                            missing = [target_video_names[vid] for vid in target_video_ids
                                       if (student.student_code, vid) not in completed_pairs]
                            pending_rows.append({"student": student, "class_name": selected_class.name,
                                                 "course_name": selected_qr.course.name if selected_qr.course else "—",
                                                 "qr_title": selected_qr.title,
                                                 "missing_videos": missing,
                                                 "status": f"Đã hoàn thành {count_completed}/{len(target_video_ids)} video"})

        if status in {"not_scanned", "not_completed"} and media_search:
            term = media_search.casefold()
            pending_rows = [r for r in pending_rows if
                            term in (r["qr_title"] or "").casefold() or
                            any(term in video_name.casefold() for video_name in r.get("missing_videos", []))]
        sv_rows = []
        sv_total = 0
        sv_pages = 1
        if status in {"not_scanned", "not_completed"}:
            sv_total = len(pending_rows)
            sv_pages = max(1, (sv_total + 9) // 10)
            sv_page = min(sv_page, sv_pages)
            for r in pending_rows[(sv_page - 1)*10:sv_page*10]:
                code = r["student"].student_code
                sv_rows.append({
                    "time": None, "student": f"{r['student'].full_name} ({student_code_tail(code)})",
                    "class_name": r["class_name"], "course_name": r["course_name"],
                    "media_name": ", ".join(r.get("missing_videos") or []) or r["qr_title"],
                    "result": "Chưa quét" if status == "not_scanned" else "Đã mở học liệu",
                    "watch_time": "00:00", "watch_status": r["status"],
                    "success": status != "not_scanned",
                })
        else:
            # Mỗi phiên xem ứng với một video thực tế. Một QR playlist có thể phát
            # nhiều video nhưng chỉ có một bản ghi quét, không nhân lượt quét.
            # Chỉ hiển thị lượt quét riêng khi chưa có phiên xem gắn với lượt đó.
            scan_ref = aliased(ScanLog)
            closest_scan_id = (select(scan_ref.id)
                .where(scan_ref.role == "SV", scan_ref.qr_id == WatchSession.qr_id,
                       scan_ref.participant_code == WatchSession.participant_code,
                       or_(scan_ref.class_code == WatchSession.class_code,
                           and_(scan_ref.class_code.is_(None), WatchSession.class_code.is_(None))),
                       scan_ref.scanned_at <= WatchSession.started_at,
                       scan_ref.scanned_at >= WatchSession.started_at - timedelta(days=1))
                .order_by(scan_ref.scanned_at.desc(), scan_ref.id.desc()).limit(1)
                .correlate(WatchSession).scalar_subquery())
            scan_with_video = (select(ScanLog.id.label("scan_id"), WatchSession.id.label("watch_id"),
                                      ScanLog.scanned_at.label("stamp"))
                .select_from(ScanLog)
                .outerjoin(WatchSession, and_(WatchSession.role == "SV", closest_scan_id == ScanLog.id))
                .join(QrCode, ScanLog.qr_id == QrCode.id)
                .where(ScanLog.role == "SV"))
            orphan_video = (select(literal(None).label("scan_id"), WatchSession.id.label("watch_id"),
                                   WatchSession.started_at.label("stamp"))
                .select_from(WatchSession)
                .outerjoin(ScanLog, ScanLog.id == closest_scan_id)
                .join(QrCode, WatchSession.qr_id == QrCode.id)
                .where(WatchSession.role == "SV", ScanLog.id.is_(None)))
            if user.role != "admin":
                scan_with_video = scan_with_video.where(managed_qr_scope(user.id))
                orphan_video = orphan_video.where(managed_qr_scope(user.id))
            events = union_all(scan_with_video, orphan_video).subquery("sv_events")
            class_expr = func.coalesce(ScanLog.class_code, WatchSession.class_code)
            course_expr = func.coalesce(ScanLog.course_id, WatchSession.course_id)
            qr_expr = func.coalesce(ScanLog.qr_id, WatchSession.qr_id)
            video_expr = func.coalesce(WatchSession.video_id, ScanLog.video_id)
            sv_query = (db.session.query(events.c.scan_id, events.c.watch_id, events.c.stamp,
                                         ScanLog, WatchSession, QrCode, Video)
                .select_from(events).outerjoin(ScanLog, ScanLog.id == events.c.scan_id)
                .outerjoin(WatchSession, WatchSession.id == events.c.watch_id)
                .join(QrCode, QrCode.id == qr_expr)
                .outerjoin(Video, Video.id == video_expr))
            if class_code:
                sv_query = sv_query.filter(class_expr == class_code)
            if course_id:
                sv_query = sv_query.filter(course_expr == course_id if selected_course else course_expr == -1)
            if qr_id:
                sv_query = sv_query.filter(qr_expr == qr_id if selected_qr else qr_expr == -1)
            if media_search:
                like = f"%{media_search}%"
                sv_query = sv_query.filter(or_(Video.title.ilike(like), QrCode.title.ilike(like)))
            if sv_start_at:
                sv_query = sv_query.filter(events.c.stamp >= sv_start_at)
            if sv_end_at:
                sv_query = sv_query.filter(events.c.stamp < sv_end_at)
            if status == "completed":
                sv_query = sv_query.filter(WatchSession.completed.is_(True))
            elif status == "unwatched":
                sv_query = sv_query.filter(WatchSession.id.is_(None), ScanLog.success.is_(True))
            elif status == "failed":
                sv_query = sv_query.filter(ScanLog.success.is_(False))
            sv_total = sv_query.order_by(None).count()
            sv_pages = max(1, (sv_total + 9) // 10)
            sv_page = min(sv_page, sv_pages)
            selection = (sv_query.order_by(events.c.stamp.desc(), events.c.scan_id.desc(), events.c.watch_id.desc())
                         .offset((sv_page - 1)*10).limit(10).all())
            codes = {s.participant_code if s else w.participant_code for _, _, _, s, w, _, _ in selection}
            students = {x.student_code: x.full_name for x in Student.query.filter(Student.student_code.in_(codes)).all()} if codes else {}
            for scan_id, watch_id, stamp, scan, watch, qr, video in selection:
                code = (scan.participant_code if scan else "") or (watch.participant_code if watch else "")
                class_ref = (scan.class_code if scan else "") or (watch.class_code if watch else "")
                course_ref = (scan.course if scan else None) or (watch.course if watch else None)
                played_video = video.title if video else (qr.title if qr else "Không xác định")
                if not watch and qr and qr.is_playlist:
                    played_video = f"{qr.title} (chưa phát video)"
                seconds = max(0, round(watch.estimated_watch_seconds)) if watch else 0
                watch_time = f"{seconds//3600:02d}:{(seconds//60)%60:02d}:{seconds%60:02d}" if seconds >= 3600 else f"{seconds//60:02d}:{seconds%60:02d}"
                success = scan.success if scan else None
                sv_rows.append({
                    "time": stamp, "student": f"{students.get(code, 'Chưa có hồ sơ SV')} ({student_code_tail(code)})",
                    "class_name": class_name_map.get(class_ref, class_ref or "—"),
                    "course_name": course_ref.name if course_ref else "—",
                    "media_name": played_video,
                    "result": ("Thành công" if success else "Không thành công") if success is not None else "Chưa có dữ liệu quét",
                    "watch_time": watch_time,
                    "watch_status": ("Hoàn thành" if watch.completed else "Chưa hoàn thành") if watch else ("Chưa xem" if success else "Không mở được"),
                    "success": success,
                })
        sv_page_numbers = list(range(max(1, sv_page-2), min(sv_pages, sv_page+2)+1))
        # Hai bộ lọc có tham số độc lập để chuyển trang hoặc lọc mục nào
        # thì mục kia không mất lựa chọn hiện tại.
        all_args = request.args.to_dict(flat=True)
        sys_args = {k:v for k,v in all_args.items() if k.startswith("sys_")}
        sv_args = {k:v for k,v in all_args.items() if not k.startswith("sys_")}
        sys_nav = {**all_args, **test_filter_args}
        sv_nav = dict(all_args)
        return render_template("analytics.html",
            course_options=course_options, class_options=class_options,
            qr_options=qr_options, status=status, pending_notice=pending_notice,
            participant_stats=participant_stats, sv_rows=sv_rows, sv_total=sv_total,
            sv_page=sv_page, sv_pages=sv_pages, sv_page_numbers=sv_page_numbers,
            sv_nav=sv_nav, sys_nav=sys_nav, sys_args=sys_args, sv_args=sv_args,
            sv_period=sv_period, sv_from=sv_from, sv_to=sv_to, media_search=media_search,
            test_rows=test_rows, tests_total=tests_total, tests_pass=tests_pass,
            tests_fail=tests_fail, test_rate=test_rate, test_page=test_page,
            test_pages=test_pages, test_page_numbers=test_page_numbers,
            test_period=test_period, test_feature=test_feature,
            test_feature_options=test_feature_options, test_result=test_result,
            system_actors=system_actors, sys_actor_id=sys_actor_id,
            test_search=test_search, test_from=test_from, test_to=test_to,
            test_filter_args=test_filter_args)

    @app.get("/exports/scans.csv")
    @login_required
    def export_scans():
        out = io.StringIO()
        w = csv.writer(out)
        w.writerow(["Mã lượt quét", "Mã sinh viên", "Mã lớp", "Môn học/ mô đun", "Mã QR", "Mã video YouTube", "Thời gian quét", "Thành công"])
        user = current_user()
        q = ScanLog.query.join(QrCode, ScanLog.qr_id == QrCode.id).filter(ScanLog.role == "SV")
        if user.role != "admin":
            q = q.filter(managed_qr_scope(user.id))
        for s in q.order_by(ScanLog.scanned_at).all():
            w.writerow([s.id, s.participant_code, s.class_code, s.course.code if s.course else "", s.qr.slug if s.qr else "", s.video.youtube_video_id if s.video else "", s.scanned_at.isoformat(), int(s.success)])
        return Response(out.getvalue(), mimetype="text/csv", headers={"Content-Disposition": "attachment; filename=scan_logs.csv"})

    @app.get("/exports/watch.csv")
    @login_required
    def export_watch():
        out = io.StringIO()
        w = csv.writer(out)
        w.writerow(["Mã phiên xem", "Mã sinh viên", "Mã lớp", "Mã QR", "Mã video YouTube", "Thời gian bắt đầu", "Thời lượng xem ước tính", "Vị trí xa nhất", "Thời lượng video", "Tỷ lệ hoàn thành", "Đã hoàn thành"])
        user = current_user()
        q = WatchSession.query.join(QrCode, WatchSession.qr_id == QrCode.id).filter(WatchSession.role == "SV")
        if user.role != "admin":
            q = q.filter(managed_qr_scope(user.id))
        for s in q.order_by(WatchSession.started_at).all():
            w.writerow([s.session_key, s.participant_code, s.class_code, s.qr.slug if s.qr else "", s.video.youtube_video_id if s.video else "", s.started_at.isoformat(), round(s.estimated_watch_seconds, 2), round(s.max_position, 2), round(s.duration_seconds, 2), round(s.completion_percent, 2), int(s.completed)])
        return Response(out.getvalue(), mimetype="text/csv", headers={"Content-Disposition": "attachment; filename=watch_sessions.csv"})


    @app.get("/classes/teaching")
    @login_required
    def teaching_classes_page():
        """GV bộ môn được tải danh sách ban đầu của lớp đang dạy, khi chưa có SV."""
        user = current_user()
        taught = (SchoolClass.query.join(CourseClass, CourseClass.school_class_id == SchoolClass.id)
                 .join(Course, Course.id == CourseClass.course_id)
                 .filter(Course.owner_id == user.id, Course.active.is_(True),
                         CourseClass.active.is_(True), SchoolClass.active.is_(True))
                 .distinct().order_by(SchoolClass.name).all())
        class_rows = []
        for cls in taught:
            taught_courses = (Course.query.join(CourseClass, CourseClass.course_id == Course.id)
                .filter(CourseClass.school_class_id == cls.id, CourseClass.active.is_(True),
                        Course.owner_id == user.id, Course.active.is_(True))
                .order_by(Course.name).all())
            class_rows.append({"class": cls, "current_size": class_size_at(cls),
                               "course_names": [c.name for c in taught_courses],
                               "student_rows": [],
                               "can_seed": (ClassEnrollment.query.filter_by(school_class_id=cls.id).first() is None
                                            and can_seed_school_class(user, cls))})
        return render_template("classes.html", class_rows=class_rows, lecturers=[],
                               status_labels=STUDENT_STATUS_LABELS, today=utcnow().date(),
                               own_classes=True, teaching_only=True)

    @app.get("/my-classes")
    @login_required
    def my_classes_page():
        """Giao diện GVCN cho cả Quản trị viên đang kiêm giảng viên."""
        user = current_user()
        classes = SchoolClass.query.filter_by(homeroom_teacher_id=user.id, active=True).order_by(SchoolClass.name).all()
        if not classes:
            flash("Tài khoản của bạn chưa khai báo lớp chủ nhiệm.", "info")
            return redirect(url_for("index"))
        today = utcnow().date()
        rows = []
        for cls in classes:
            latest = {}
            for e in sorted(cls.enrollments, key=lambda x: (x.started_on, x.id)):
                latest[e.student_id] = e
            student_rows = []
            for e in latest.values():
                key, reference = student_current_status_in_class(e.student, cls, today)
                student_rows.append({"student": e.student, "status_key": key,
                                     "status_label": STUDENT_STATUS_LABELS.get(key, "Khác"), "reference": reference})
            student_rows.sort(key=lambda x: ((x["student"].full_name or "").casefold(), x["student"].student_code))
            rows.append({"class": cls, "current_size": class_size_at(cls, today), "student_rows": student_rows})
        return render_template("classes.html", class_rows=rows, lecturers=[],
                               status_labels=STUDENT_STATUS_LABELS, today=today, own_classes=True)

    @app.get("/classes")
    @login_required
    def classes_page():
        user = current_user()
        today = utcnow().date()
        class_rows = []

        if user.role == "admin":
            # Quản trị viên chỉ quản lý danh mục lớp và phân công GVCN.
            # Không tải/hiển thị chi tiết danh sách sinh viên như giao diện GVCN.
            classes = SchoolClass.query.order_by(SchoolClass.name).all()
            for school_class in classes:
                class_rows.append({
                    "class": school_class,
                    # Sĩ số hiện tại chỉ tính SV đang có quan hệ lớp còn hiệu lực,
                    # tương ứng trạng thái Đang học tại ngày hiện tại.
                    "current_size": class_size_at(school_class, today),
                })
            lecturers = User.query.order_by(User.name, User.email).all()
        else:
            classes = SchoolClass.query.filter_by(homeroom_teacher_id=user.id, active=True).order_by(SchoolClass.name).all()
            if not classes:
                flash("Tài khoản của bạn hiện chưa được giao chủ nhiệm lớp nào.", "info")
                return redirect(url_for("index"))
            for school_class in classes:
                latest_by_student = {}
                for e in sorted(school_class.enrollments, key=lambda x: (x.started_on, x.id)):
                    latest_by_student[e.student_id] = e
                student_rows = []
                for student_id, latest in latest_by_student.items():
                    student = latest.student
                    status_key, ref = student_current_status_in_class(student, school_class, today)
                    student_rows.append({
                        "student": student,
                        "status_key": status_key,
                        "status_label": STUDENT_STATUS_LABELS.get(status_key, "Khác"),
                        "reference": ref,
                    })
                student_rows.sort(key=lambda x: ((x["student"].full_name or "").casefold(), x["student"].student_code))
                class_rows.append({
                    "class": school_class,
                    "current_size": class_size_at(school_class, today),
                    "student_rows": student_rows,
                })
            lecturers = []

        return render_template(
            "classes.html", class_rows=class_rows, lecturers=lecturers,
            status_labels=STUDENT_STATUS_LABELS, today=today,
        )

    @app.post("/classes/<int:class_id>/rename")
    @login_required
    def rename_school_class(class_id):
        user = current_user()
        school_class = db.session.get(SchoolClass, class_id) or abort(404)
        if not can_manage_school_class(user, school_class):
            abort(403)
        new_name = re.sub(r"\s+", " ", (request.form.get("class_name") or "").strip())
        if not new_name or len(new_name) > 255:
            flash("Tên lớp không hợp lệ.", "warning")
        elif (other := shared_class_by_name(new_name)) and other.id != school_class.id:
            flash("Tên lớp này đã có trong hệ thống. Hãy chọn lớp có sẵn để tránh trùng dữ liệu.", "warning")
        else:
            audit_edit(user, "school_class", school_class.id, "name", school_class.name, new_name)
            school_class.name = new_name
            db.session.commit()
            flash("Đã cập nhật tên lớp; mã lớp, danh sách SV và thống kê giữ nguyên.", "success")
        return redirect(url_for("my_classes_page") if school_class.homeroom_teacher_id == user.id else url_for("classes_page"))

    @app.post("/admin/classes/create")
    @login_required
    def create_school_class():
        user = current_user()
        if user.role != "admin":
            flash("Bạn không có quyền thực hiện thao tác này.", "warning")
            return redirect(url_for("index"))
        name = (request.form.get("class_name") or "").strip()
        homeroom_id = request.form.get("homeroom_teacher_id", type=int)
        if not name:
            flash("Vui lòng nhập tên lớp.", "warning")
            return redirect(url_for("classes_page"))
        if shared_class_by_name(name):
            flash("Lớp đã có trên hệ thống; không tạo trùng.", "warning")
            return redirect(url_for("classes_page"))
        teacher = db.session.get(User, homeroom_id) if homeroom_id else None
        if teacher and teacher.role not in {"lecturer", "admin"}:
            teacher = None
        school_class = SchoolClass(code=f"TMP-{secrets.token_hex(8)}", name=name[:255], homeroom_teacher_id=teacher.id if teacher else None)
        db.session.add(school_class)
        db.session.flush()
        assign_auto_school_class_code(school_class)
        db.session.commit()
        flash(f"Đã tạo lớp “{name}”. Mã lớp {school_class.code} được hệ thống tự động tạo.", "success")
        return redirect(url_for("classes_page"))

    @app.post("/admin/classes/<int:class_id>/homeroom")
    @login_required
    def assign_homeroom_teacher(class_id):
        user = current_user()
        if user.role != "admin":
            flash("Bạn không có quyền thực hiện thao tác này.", "warning")
            return redirect(url_for("index"))
        school_class = db.session.get(SchoolClass, class_id) or abort(404)
        teacher_id = request.form.get("homeroom_teacher_id", type=int)
        teacher = db.session.get(User, teacher_id) if teacher_id else None
        if teacher and teacher.role not in {"lecturer", "admin"}:
            flash("Tài khoản được chọn không phải là giảng viên.", "warning")
            return redirect(url_for("classes_page"))
        effective_text = (request.form.get("effective_date") or "").strip()
        try:
            effective = date.fromisoformat(effective_text) if effective_text else utcnow().date()
        except ValueError:
            flash("Ngày chuyển giao không hợp lệ.", "warning")
            return redirect(url_for("classes_page"))
        if effective > utcnow().date():
            flash("Vui lòng chọn ngày có hiệu lực không muộn hơn hôm nay để chuyển giao ngay.", "warning")
            return redirect(url_for("classes_page"))
        if school_class.homeroom_teacher_id != (teacher.id if teacher else None):
            db.session.add(HomeroomChange(school_class_id=school_class.id,
                         previous_teacher_id=school_class.homeroom_teacher_id,
                         new_teacher_id=teacher.id if teacher else None,
                         effective_on=effective, changed_by=user.id))
        school_class.homeroom_teacher_id = teacher.id if teacher else None
        db.session.commit()
        flash(f"Đã cập nhật giảng viên chủ nhiệm cho lớp “{school_class.name}”.", "success")
        return redirect(url_for("classes_page"))

    @app.get("/classes/student-template.xlsx")
    @login_required
    def student_excel_template():
        from openpyxl import Workbook
        wb = Workbook()
        ws = wb.active
        ws.title = "Danh sách sinh viên"
        ws.append(["Mã sinh viên", "Họ và tên"])
        ws.append(["MSSV001", "Nguyễn Văn A"])
        bio = io.BytesIO()
        wb.save(bio)
        bio.seek(0)
        return send_file(bio, as_attachment=True, download_name="Mau_danh_sach_sinh_vien.xlsx", mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")

    def read_students_from_excel(upload):
        from openpyxl import load_workbook
        wb = load_workbook(upload, read_only=True, data_only=True)
        ws = wb.active
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            raise ValueError("Tệp Excel không có dữ liệu.")
        def clean_header(v):
            s = unicodedata.normalize("NFD", str(v or "").strip().lower())
            s = "".join(ch for ch in s if unicodedata.category(ch) != "Mn")
            return re.sub(r"[^a-z0-9]+", " ", s).strip()
        headers = [clean_header(v) for v in rows[0]]
        code_idx = next((i for i,h in enumerate(headers) if h in {"ma sinh vien","ma sv","mssv","student code"}), None)
        name_idx = next((i for i,h in enumerate(headers) if h in {"ho va ten","ho ten","ten sinh vien","full name"}), None)
        if code_idx is None or name_idx is None:
            raise ValueError("Tệp Excel cần có hai cột: “Mã sinh viên” và “Họ và tên”.")
        result = []
        seen = set()
        for row in rows[1:]:
            code = normalize_student_code(row[code_idx] if code_idx < len(row) else "")
            name = str(row[name_idx] or "").strip() if name_idx < len(row) else ""
            if not code or not name or code in seen:
                continue
            seen.add(code)
            result.append((code, name[:255]))
        return result

    def add_student_to_school_class(school_class, code, name, effective_date):
        student = Student.query.filter_by(student_code=code).first()
        if not student:
            student = Student(student_code=code, full_name=name)
            db.session.add(student)
            db.session.flush()
        # Không tự ghi đè tên đã được GVCN hiệu chỉnh bởi lần tải Excel khác.
        for e in active_enrollments_for_student(student, effective_date):
            if e.school_class_id == school_class.id:
                return "already"
            return f"other:{e.school_class.name}"
        db.session.add(ClassEnrollment(school_class_id=school_class.id, student_id=student.id, started_on=effective_date))
        return "added"

    @app.post("/classes/<int:class_id>/students/upload")
    @login_required
    def upload_class_students(class_id):
        user = current_user()
        school_class = db.session.get(SchoolClass, class_id) or abort(404)
        if not can_seed_school_class(user, school_class):
            flash("Chỉ GVCN được cập nhật danh sách đã có. GV dạy lớp chỉ được nhập danh sách ban đầu nếu lớp chưa có SV.", "warning")
            return redirect(url_for("my_classes_page") if SchoolClass.query.filter_by(homeroom_teacher_id=user.id).first() else url_for("teaching_classes_page"))
        upload = request.files.get("student_file")
        effective_text = (request.form.get("effective_date") or "").strip()
        if not upload or not upload.filename:
            flash("Vui lòng chọn tệp Excel danh sách sinh viên.", "warning")
            return redirect(class_management_url(user, school_class))
        try:
            effective_date = date.fromisoformat(effective_text) if effective_text else utcnow().date()
            rows = read_students_from_excel(upload)
            if not rows:
                raise ValueError("Tệp Excel không có sinh viên hợp lệ. Vui lòng kiểm tra cột Mã sinh viên và Họ và tên.")
            added = already = conflict = 0
            conflicts = []
            for code, name in rows:
                result = add_student_to_school_class(school_class, code, name, effective_date)
                if result == "added": added += 1
                elif result == "already": already += 1
                else:
                    conflict += 1
                    conflicts.append(f"{code} đang thuộc {result.split(':',1)[1]}")
            db.session.commit()
            msg = f"Đã bổ sung {added} sinh viên vào lớp; {already} sinh viên đã có trong lớp"
            if conflict:
                msg += f"; {conflict} sinh viên chưa thêm do đang thuộc lớp khác"
            flash(msg + ".", "success" if added or already else "warning")
            if conflicts:
                flash("Cần kiểm tra: " + "; ".join(conflicts[:8]), "warning")
        except Exception as exc:
            db.session.rollback()
            flash(f"Không đọc được danh sách sinh viên: {exc}", "danger")
        return redirect(class_management_url(user, school_class))

    @app.post("/classes/<int:class_id>/students/add")
    @login_required
    def add_class_student(class_id):
        user = current_user()
        school_class = db.session.get(SchoolClass, class_id) or abort(404)
        if not can_seed_school_class(user, school_class):
            flash("Chỉ GVCN được cập nhật danh sách đã có. GV dạy lớp chỉ được nhập danh sách ban đầu nếu lớp chưa có SV.", "warning")
            return redirect(url_for("my_classes_page") if SchoolClass.query.filter_by(homeroom_teacher_id=user.id).first() else url_for("teaching_classes_page"))
        code = normalize_student_code(request.form.get("student_code", ""))
        name = (request.form.get("full_name") or "").strip()
        effective_text = (request.form.get("effective_date") or "").strip()
        if not code or not name:
            flash("Vui lòng nhập mã sinh viên và họ tên.", "warning")
            return redirect(class_management_url(user, school_class))
        try:
            effective_date = date.fromisoformat(effective_text) if effective_text else utcnow().date()
            result = add_student_to_school_class(school_class, code, name, effective_date)
            if result.startswith("other:"):
                raise ValueError(f"Sinh viên đang thuộc lớp {result.split(':',1)[1]} tại ngày hiệu lực. Hãy cập nhật trạng thái ở lớp cũ trước.")
            db.session.commit()
            flash("Đã bổ sung sinh viên vào lớp." if result == "added" else "Sinh viên đã có trong lớp ở thời điểm này.", "success" if result == "added" else "info")
        except Exception as exc:
            db.session.rollback()
            flash(f"Không thêm được sinh viên: {exc}", "danger")
        return redirect(class_management_url(user, school_class))

    @app.post("/classes/<int:class_id>/students/<int:student_id>/edit")
    @login_required
    def edit_student_name(class_id, student_id):
        user = current_user()
        school_class = db.session.get(SchoolClass, class_id) or abort(404)
        student = db.session.get(Student, student_id) or abort(404)
        if school_class.homeroom_teacher_id != user.id:
            abort(403)
        if not ClassEnrollment.query.filter_by(school_class_id=class_id, student_id=student_id).first():
            abort(404)
        name = re.sub(r"\s+", " ", (request.form.get("full_name") or "").strip())
        if len(name) < 2 or len(name) > 255:
            flash("Họ tên sinh viên không hợp lệ.", "warning")
        else:
            audit_edit(user, "student", student.id, "full_name", student.full_name, name)
            student.full_name = name
            db.session.commit()
            flash("Đã sửa họ tên sinh viên. Mã SV và lịch sử quét/xem giữ nguyên.", "success")
        return redirect(url_for("my_classes_page") + f"#lop-{class_id}")

    @app.post("/classes/<int:class_id>/students/<int:student_id>/status")
    @login_required
    def change_student_status(class_id, student_id):
        user = current_user()
        school_class = db.session.get(SchoolClass, class_id) or abort(404)
        student = db.session.get(Student, student_id) or abort(404)
        if not can_manage_school_class(user, school_class):
            flash("Bạn chỉ được cập nhật sinh viên của lớp mình chủ nhiệm.", "warning")
            return redirect(url_for("classes_page"))
        status = (request.form.get("status") or "").strip()
        effective_text = (request.form.get("effective_date") or "").strip()
        note = (request.form.get("note") or "").strip()
        if status not in STUDENT_STATUS_LABELS:
            flash("Vui lòng chọn trạng thái hợp lệ.", "warning")
            return redirect(class_management_url(user, school_class))
        try:
            effective_date = date.fromisoformat(effective_text) if effective_text else utcnow().date()
            active_here = next((e for e in student.enrollments if e.school_class_id == school_class.id and enrollment_active_on(e, effective_date)), None)
            if status == "dang_hoc":
                if active_here:
                    flash("Sinh viên đang ở trạng thái Đang học trong lớp tại ngày này.", "info")
                    return redirect(class_management_url(user, school_class))
                for e in active_enrollments_for_student(student, effective_date):
                    if e.school_class_id != school_class.id:
                        raise ValueError(f"Sinh viên đang thuộc lớp {e.school_class.name} tại ngày hiệu lực.")
                db.session.add(ClassEnrollment(school_class_id=school_class.id, student_id=student.id, started_on=effective_date, note=note or None))
            else:
                if not active_here:
                    raise ValueError("Không tìm thấy khoảng thời gian sinh viên đang học tại lớp ở ngày hiệu lực đã chọn.")
                if effective_date < active_here.started_on:
                    raise ValueError("Ngày thay đổi không được trước ngày sinh viên bắt đầu vào lớp.")
                active_here.ended_on = effective_date
                active_here.end_status = status
                if note:
                    active_here.note = note
            db.session.commit()
            flash(f"Đã cập nhật trạng thái sinh viên thành “{STUDENT_STATUS_LABELS[status]}” từ ngày {effective_date.strftime('%d/%m/%Y')}.", "success")
        except Exception as exc:
            db.session.rollback()
            flash(f"Không cập nhật được trạng thái: {exc}", "danger")
        return redirect(class_management_url(user, school_class))

    @app.post("/classes/<int:class_id>/students/<int:student_id>/delete")
    @login_required
    def delete_class_student(class_id, student_id):
        """Xóa một SV được nhập nhầm/thừa khỏi lớp và loại bỏ số liệu thống kê phát sinh từ SV đó trong lớp.

        Đây là thao tác xóa dữ liệu sai, không dùng cho các trường hợp bảo lưu/nghỉ học/chuyển lớp.
        Với biến động thật của sĩ số phải dùng chức năng Thay đổi trạng thái để giữ lịch sử theo thời gian.
        """
        user = current_user()
        school_class = db.session.get(SchoolClass, class_id) or abort(404)
        student = db.session.get(Student, student_id) or abort(404)

        # Chỉ GVCN của chính lớp được xóa SV khỏi danh sách. Quản trị viên không thao tác
        # danh sách chi tiết SV trên giao diện quản trị lớp.
        if school_class.homeroom_teacher_id != user.id:
            flash("Bạn chỉ được xóa sinh viên khỏi lớp mình đang chủ nhiệm.", "warning")
            return redirect(url_for("classes_page"))

        enrollments = ClassEnrollment.query.filter_by(
            school_class_id=school_class.id, student_id=student.id
        ).all()
        if not enrollments:
            flash("Sinh viên không có trong danh sách của lớp này.", "warning")
            return redirect(class_management_url(user, school_class))

        code = student.student_code
        try:
            # Xóa toàn bộ dữ liệu quét/xem phát sinh dưới mã SV này trong chính lớp đang xóa.
            # Nhờ đó các chỉ số ở Tổng quan/Thống kê khai thác được tính lại ngay từ dữ liệu còn lại.
            watch_sessions = WatchSession.query.filter_by(
                participant_code=code, class_code=school_class.code
            ).all()
            watch_session_ids = [row.id for row in watch_sessions]
            watch_event_count = 0
            if watch_session_ids:
                watch_event_count = WatchEvent.query.filter(
                    WatchEvent.watch_session_id.in_(watch_session_ids)
                ).count()
                WatchEvent.query.filter(
                    WatchEvent.watch_session_id.in_(watch_session_ids)
                ).delete(synchronize_session=False)

            watch_count = len(watch_sessions)
            if watch_session_ids:
                WatchSession.query.filter(
                    WatchSession.id.in_(watch_session_ids)
                ).delete(synchronize_session=False)

            scan_count = ScanLog.query.filter_by(
                participant_code=code, class_code=school_class.code
            ).count()
            ScanLog.query.filter_by(
                participant_code=code, class_code=school_class.code
            ).delete(synchronize_session=False)
            # Bản ghi kiểm thử quét QR của SV nhập nhầm cũng phải được trừ khỏi nghiệm thu.
            FunctionalTestRun.query.filter_by(
                participant_code=code, participant_class_code=school_class.code
            ).delete(synchronize_session=False)

            # Xóa lịch sử SV thuộc lớp bị nhập nhầm.
            ClassEnrollment.query.filter_by(
                school_class_id=school_class.id, student_id=student.id
            ).delete(synchronize_session=False)

            # Nếu SV không còn thuộc bất kỳ lớp nào trong hệ thống thì xóa luôn hồ sơ SV và
            # dữ liệu nhận diện cũ còn sót lại. Nếu SV còn lịch sử ở lớp khác thì giữ hồ sơ đó.
            remaining_enrollment_count = ClassEnrollment.query.filter_by(student_id=student.id).count()
            removed_master = False
            if remaining_enrollment_count == 0:
                # Phòng trường hợp dữ liệu rất cũ chưa có class_code, xóa sạch log theo mã SV
                # vì hồ sơ SV này không còn được sử dụng ở bất kỳ lớp nào.
                remaining_watch_ids = [row.id for row in WatchSession.query.filter_by(participant_code=code).all()]
                if remaining_watch_ids:
                    WatchEvent.query.filter(WatchEvent.watch_session_id.in_(remaining_watch_ids)).delete(synchronize_session=False)
                    WatchSession.query.filter(WatchSession.id.in_(remaining_watch_ids)).delete(synchronize_session=False)
                ScanLog.query.filter_by(participant_code=code).delete(synchronize_session=False)
                FunctionalTestRun.query.filter_by(participant_code=code).delete(synchronize_session=False)
                Participant.query.filter_by(participant_code=code).delete(synchronize_session=False)
                db.session.delete(student)
                removed_master = True

            db.session.commit()
            extra = " Hồ sơ sinh viên cũng đã được xóa khỏi hệ thống vì không còn thuộc lớp nào khác." if removed_master else ""
            flash(
                f"Đã xóa sinh viên {code} khỏi lớp “{school_class.name}” và loại bỏ {scan_count} lượt quét, "
                f"{watch_count} phiên xem cùng dữ liệu sự kiện xem liên quan khỏi thống kê.{extra}",
                "success",
            )
        except Exception as exc:
            db.session.rollback()
            flash(f"Không xóa được sinh viên: {exc}", "danger")

        return redirect(class_management_url(user, school_class))

    @app.route("/research", methods=["GET", "POST"])
    @login_required
    def research_page():
        user = current_user()
        if user.role != "admin":
            flash("Bạn không có quyền truy cập chức năng này. Chức năng chỉ dành cho Quản trị viên.", "warning")
            return redirect(url_for("index"))

        if request.method == "POST":
            form_type = request.form.get("form_type")
            try:
                if form_type == "manual_experiment":
                    task_key = (request.form.get("task_key") or "").strip()
                    operations_count = int(request.form.get("operations_count") or 0)
                    elapsed_seconds = float(request.form.get("elapsed_seconds") or 0)
                    if task_key not in RESEARCH_COMPARISON_TASK_MAP:
                        raise ValueError("Vui lòng chọn một nhiệm vụ đối chứng trong danh mục chuẩn của hệ thống.")
                    if operations_count < 0 or elapsed_seconds < 0:
                        raise ValueError("Vui lòng nhập đầy đủ dữ liệu đối chứng thủ công.")
                    db.session.add(ExperimentRun(
                        owner_id=user.id,
                        task_key=task_key,
                        mode="manual",
                        operations_count=operations_count,
                        elapsed_seconds=elapsed_seconds,
                        system_processing_seconds=None,
                        success=request.form.get("success", "1") == "1",
                        note=(request.form.get("note") or "").strip(),
                    ))
                    flash("Đã lưu phép đo thủ công để làm dữ liệu đối chứng.", "success")
                elif form_type == "survey_links":
                    save_research_setting("lecturer_form_url", request.form.get("lecturer_form_url"))
                    save_research_setting("student_form_url", request.form.get("student_form_url"))
                    save_research_setting("survey_results_url", request.form.get("survey_results_url"))
                    flash("Đã lưu liên kết Google Forms/Google Sheets.", "success")
                else:
                    raise ValueError("Biểu mẫu không hợp lệ.")
                db.session.commit()
            except Exception as exc:
                db.session.rollback()
                flash(f"Không lưu được dữ liệu: {exc}", "danger")
            return redirect(url_for("research_page"))

        metrics = acceptance_metrics(user)
        experiments = ExperimentRun.query.order_by(ExperimentRun.created_at.desc()).limit(200).all()
        manual_runs = [x for x in experiments if x.mode == "manual"]
        webapp_runs = [x for x in experiments if x.mode == "webapp"]

        # Chỉ so sánh theo danh mục nhiệm vụ chuẩn. Nhờ vậy mã đối chứng thủ công
        # luôn trùng chính xác với mã Web App tự ghi và không phát sinh mã nhập sai.
        comparison_rows = []
        task_catalog = []
        for task in RESEARCH_COMPARISON_TASKS:
            key = task["key"]
            m = [x for x in manual_runs if x.task_key == key and x.success]
            w = [x for x in webapp_runs if x.task_key == key and x.success]
            m_ops = (sum(x.operations_count for x in m) / len(m)) if m else None
            w_ops = (sum(x.operations_count for x in w) / len(w)) if w else None
            m_time = (sum(x.elapsed_seconds for x in m) / len(m)) if m else None
            w_time = (sum(x.elapsed_seconds for x in w) / len(w)) if w else None
            reduction_ops = round(100 * (m_ops - w_ops) / m_ops, 1) if m_ops and w_ops is not None else None
            reduction_time = round(100 * (m_time - w_time) / m_time, 1) if m_time and w_time is not None else None
            comparison_rows.append({
                "task_key": key,
                "task_name": task["name"],
                "manual_task": task["manual_task"],
                "manual_ops": m_ops,
                "webapp_ops": w_ops,
                "manual_time": m_time,
                "webapp_time": w_time,
                "reduction_ops": reduction_ops,
                "reduction_time": reduction_time,
                "manual_count": len(m),
                "webapp_count": len(w),
            })
            if m and w:
                status = "Đủ dữ liệu để so sánh"
                status_class = "success"
            elif w and not m:
                status = "Cần đo quy trình thủ công"
                status_class = "warning"
            elif m and not w:
                status = "Cần thực hiện trên Web App"
                status_class = "warning"
            else:
                status = "Chưa có dữ liệu"
                status_class = "secondary"
            task_catalog.append({
                **task,
                "manual_count": len(m),
                "webapp_count": len(w),
                "status": status,
                "status_class": status_class,
            })

        survey_links = {
            "lecturer_form_url": research_setting("lecturer_form_url"),
            "student_form_url": research_setting("student_form_url"),
            "survey_results_url": research_setting("survey_results_url"),
        }
        return render_template(
            "research.html", metrics=metrics, experiments=experiments,
            webapp_runs=webapp_runs, manual_runs=manual_runs,
            comparison_rows=comparison_rows, task_catalog=task_catalog,
            survey_links=survey_links,
        )

    @app.post("/admin/catalog/rename")
    @login_required
    def rename_shared_catalog():
        user = current_user()
        if user.role != "admin":
            abort(403)
        field = (request.form.get("field") or "").strip()
        old = (request.form.get("old_value") or "").strip()
        new = re.sub(r"\s+", " ", (request.form.get("new_value") or "").strip())
        if field not in {"department_name", "faculty_name"} or not old or not new or len(new) > 255:
            flash("Thông tin chuẩn hóa khoa/bộ môn không hợp lệ.", "warning")
            return redirect(url_for("lecturers_page"))
        matched = Course.query.filter(getattr(Course, field) == old).all()
        for course in matched:
            audit_edit(user, "course", course.id, field, old, new)
            setattr(course, field, new)
            if field == "faculty_name":
                course.unit_name = new
        db.session.commit()
        flash(f"Đã chuẩn hóa {len(matched)} môn học/mô đun theo tên mới.", "success")
        return redirect(url_for("lecturers_page"))

    @app.get("/admin/lecturers")
    @login_required
    def lecturers_page():
        user = current_user()
        if user.role != "admin":
            abort(403)

        users = User.query.order_by(User.role.desc(), User.name.asc(), User.email.asc()).all()
        rows = []
        for account in users:
            channels = YouTubeChannel.query.filter_by(user_id=account.id).order_by(YouTubeChannel.created_at.asc()).all()
            channel_names = [c.title.strip() for c in channels if (c.title or "").strip()]
            if account.role == "admin":
                position = ("Quản trị viên kiêm GVCN" if
                            SchoolClass.query.filter_by(homeroom_teacher_id=account.id, active=True).first()
                            else "Quản trị viên kiêm giảng viên")
            elif SchoolClass.query.filter_by(homeroom_teacher_id=account.id, active=True).first():
                position = "Giảng viên kiêm GVCN"
            else:
                position = "Giảng viên"
            rows.append({
                "user": account,
                "channel_name": " / ".join(channel_names) if channel_names else "Chưa đồng bộ kênh",
                "position": position,
            })

        return render_template("lecturers.html", rows=rows)

    @app.post("/admin/lecturers/<int:lecturer_id>/assign-code")
    @login_required
    def assign_lecturer_code_admin(lecturer_id):
        if current_user().role != "admin":
            abort(403)
        flash("Mã giảng viên được hệ thống cấp tự động sau khi người dùng hoàn tất khai báo lần đầu.", "info")
        return redirect(url_for("lecturers_page"))

    @app.post("/admin/lecturers/<int:lecturer_id>/professional-group")
    @login_required
    def update_lecturer_professional_group(lecturer_id):
        admin = current_user()
        if admin.role != "admin":
            abort(403)
        lecturer = db.session.get(User, lecturer_id) or abort(404)
        # Quản trị viên cũng là giảng viên; được cập nhật nhóm chuyên môn của mình.
        group = (request.form.get("professional_group") or "").strip()
        if group not in {"it", "other", ""}:
            flash("Nhóm chuyên môn không hợp lệ.", "warning")
            return redirect(url_for("lecturers_page"))
        lecturer.professional_group = group or None
        db.session.commit()
        label = {"it": "Công nghệ thông tin", "other": "Khác CNTT", "": "Chưa xác định"}.get(group, "Chưa xác định")
        flash(f"Đã cập nhật nhóm chuyên môn của {lecturer.name or lecturer.email}: {label}.", "success")
        return redirect(url_for("lecturers_page"))

    @app.post("/admin/lecturers/<int:lecturer_id>/delete")
    @login_required
    def delete_lecturer_account(lecturer_id):
        admin = current_user()
        if admin.role != "admin":
            abort(403)
        lecturer = db.session.get(User, lecturer_id) or abort(404)
        if lecturer.role != "lecturer":
            flash("Chỉ có thể xóa tài khoản Giảng viên tại chức năng này. Tài khoản Quản trị viên được bảo vệ.", "warning")
            return redirect(url_for("lecturers_page"))
        # Không được cascade-delete QR đã in và nguồn học liệu bàn giao!
        # Trường hợp này cần lưu giữ tài khoản nguồn dù GV không còn giảng dạy.
        handed_off_qr = (QrCode.query.join(Course, Course.id == QrCode.course_id)
            .filter(QrCode.owner_id == lecturer.id, Course.owner_id != lecturer.id).first())
        handed_off_playlist = (CoursePlaylist.query.join(Course, Course.id == CoursePlaylist.course_id)
            .join(Playlist, Playlist.id == CoursePlaylist.playlist_id)
            .join(YouTubeChannel, YouTubeChannel.id == Playlist.channel_id)
            .filter(YouTubeChannel.user_id == lecturer.id, Course.owner_id != lecturer.id).first())
        handed_off_video = (CourseVideo.query.join(Course, Course.id == CourseVideo.course_id)
            .join(Video, Video.id == CourseVideo.video_id)
            .join(YouTubeChannel, YouTubeChannel.id == Video.channel_id)
            .filter(YouTubeChannel.user_id == lecturer.id, Course.owner_id != lecturer.id).first())
        if handed_off_qr or handed_off_playlist or handed_off_video:
            flash("Chưa thể xóa giảng viên này: mã QR hoặc học liệu của tài khoản vẫn đang phục vụ môn đã bàn giao. "
                  "Giữ tài khoản/kênh YouTube để bảo toàn QR đã phát hành và lịch sử khai thác.", "warning")
            return redirect(url_for("lecturers_page"))
        active_course_qrs = (QrCode.query.join(Course, Course.id == QrCode.course_id)
            .filter(Course.owner_id == lecturer.id, QrCode.active.is_(True)).first())
        if active_course_qrs:
            flash("Giảng viên này đang phụ trách môn có QR còn hoạt động. "
                  "Vui lòng bàn giao môn sang giảng viên khác trước khi xóa tài khoản, để mã QR đã in không bị mất.", "warning")
            return redirect(url_for("lecturers_page"))
        try:
            display = lecturer.name or lecturer.email
            email = lecturer.email

            # Lớp và sinh viên là dữ liệu dùng chung của Nhà trường nên không xóa theo GV.
            # Khi xóa một GVCN, lớp được giữ lại nhưng chuyển về trạng thái chưa giao GVCN.
            SchoolClass.query.filter_by(homeroom_teacher_id=lecturer.id).update(
                {SchoolClass.homeroom_teacher_id: None}, synchronize_session=False
            )

            # Xóa toàn bộ dữ liệu nghiệp vụ/thử nghiệm thuộc GV trước khi xóa tài khoản.
            delete_lecturer_business_data(lecturer)
            OAuthAccount.query.filter_by(user_id=lecturer.id).delete(synchronize_session=False)
            db.session.delete(lecturer)
            db.session.commit()
            flash(
                f"Đã xóa giảng viên {display} ({email}) khỏi Web App cùng toàn bộ dữ liệu thuộc giảng viên này. "
                "Các lớp và danh sách sinh viên dùng chung được giữ lại; lớp từng do giảng viên chủ nhiệm hiện ở trạng thái chưa giao GVCN.",
                "success",
            )
        except Exception as exc:
            db.session.rollback()
            flash(f"Không thể xóa giảng viên: {exc}", "danger")
        return redirect(url_for("lecturers_page"))

    def delete_lecturer_business_data(lecturer: User):
        """Xóa dữ liệu nghiệp vụ/thử nghiệm của một GV nhưng giữ tài khoản, họ tên, email, mã GV và quyền GVCN."""
        qr_ids = [x[0] for x in db.session.query(QrCode.id).filter(QrCode.owner_id == lecturer.id).all()]
        course_ids = [x[0] for x in db.session.query(Course.id).filter(Course.owner_id == lecturer.id).all()]
        channel_ids = [x[0] for x in db.session.query(YouTubeChannel.id).filter(YouTubeChannel.user_id == lecturer.id).all()]

        if qr_ids:
            watch_ids = [x[0] for x in db.session.query(WatchSession.id).filter(WatchSession.qr_id.in_(qr_ids)).all()]
            if watch_ids:
                WatchEvent.query.filter(WatchEvent.watch_session_id.in_(watch_ids)).delete(synchronize_session=False)
            WatchSession.query.filter(WatchSession.qr_id.in_(qr_ids)).delete(synchronize_session=False)
            ScanLog.query.filter(ScanLog.qr_id.in_(qr_ids)).delete(synchronize_session=False)
            QrTargetHistory.query.filter(QrTargetHistory.qr_id.in_(qr_ids)).delete(synchronize_session=False)

        if course_ids:
            Participant.query.filter(Participant.course_id.in_(course_ids)).delete(synchronize_session=False)
            CourseVideo.query.filter(CourseVideo.course_id.in_(course_ids)).delete(synchronize_session=False)
            CoursePlaylist.query.filter(CoursePlaylist.course_id.in_(course_ids)).delete(synchronize_session=False)
            CourseClass.query.filter(CourseClass.course_id.in_(course_ids)).delete(synchronize_session=False)
            ClassGroup.query.filter(ClassGroup.course_id.in_(course_ids)).delete(synchronize_session=False)

        QrCode.query.filter_by(owner_id=lecturer.id).delete(synchronize_session=False)
        Course.query.filter_by(owner_id=lecturer.id).delete(synchronize_session=False)

        if channel_ids:
            HubEvent.query.filter(HubEvent.channel_id.in_(channel_ids)).delete(synchronize_session=False)
            # HubSubscription/Playlist/Video/PlaylistVideo được xóa theo cascade khi xóa kênh.
        YouTubeChannel.query.filter_by(user_id=lecturer.id).delete(synchronize_session=False)

        ConnectedGoogleAccount.query.filter_by(user_id=lecturer.id).delete(synchronize_session=False)
        EditHistory.query.filter_by(changed_by=lecturer.id).delete(synchronize_session=False)
        ExperimentRun.query.filter_by(owner_id=lecturer.id).delete(synchronize_session=False)
        FunctionalTestRun.query.filter_by(owner_id=lecturer.id).delete(synchronize_session=False)
        SurveyResult.query.filter_by(owner_id=lecturer.id).delete(synchronize_session=False)

    @app.post("/admin/lecturers/<int:lecturer_id>/reset-data")
    @login_required
    def reset_lecturer_data(lecturer_id):
        # Giữ tương thích với liên kết cũ. Từ v3.9, quản lý/xóa GV thực hiện tại trang Giảng viên.
        if current_user().role != "admin":
            abort(403)
        flash("Chức năng quản lý giảng viên đã được chuyển sang giao diện Giảng viên.", "info")
        return redirect(url_for("lecturers_page"))

    def reset_business_code_sequences():
        """Đặt lại bộ đếm ID của các bảng dùng ID để sinh mã nghiệp vụ.

        Không đặt lại sequence users vì tài khoản quản trị hiện tại được giữ lại;
        mã GV được cấp bằng bộ đếm riêng trong allocate_lecturer_code().
        """
        table_names = [Course.__tablename__, SchoolClass.__tablename__, ClassGroup.__tablename__]
        dialect = db.engine.dialect.name
        if dialect == "postgresql":
            for table_name in table_names:
                db.session.execute(text(
                    "SELECT setval(pg_get_serial_sequence(:table_name, 'id'), 1, false)"
                ), {"table_name": table_name})
        elif dialect == "sqlite":
            for table_name in table_names:
                db.session.execute(text("DELETE FROM sqlite_sequence WHERE name = :name"), {"name": table_name})

    @app.post("/admin/reset-all-data")
    @login_required
    def reset_all_data():
        user = current_user()
        if user.role != "admin":
            abort(403)
        try:
            admin_id = user.id
            # Xóa toàn bộ dữ liệu nghiệp vụ/thử nghiệm nhưng giữ tài khoản + phiên OAuth
            # của chính quản trị viên đang thao tác để không phải đăng nhập lại.
            # Tất cả tài khoản GV bị xóa nên toàn bộ mã GV đã cấp cũng được làm mới từ đầu.
            keep_models = {User, OAuthAccount}
            for _, model in reversed(backup_table_specs()):
                if model in keep_models:
                    continue
                db.session.execute(model.__table__.delete())
            OAuthAccount.query.filter(OAuthAccount.user_id != admin_id).delete(synchronize_session=False)
            User.query.filter(User.id != admin_id).delete(synchronize_session=False)
            user.lecturer_code = None
            user.profile_completed = False
            user.professional_group = None
            reset_business_code_sequences()
            db.session.commit()
            session["user_id"] = admin_id
            session.permanent = True
            flash("Đã xóa toàn bộ dữ liệu đã lưu, toàn bộ mã GV đã cấp và đặt lại bộ đếm mã nghiệp vụ. Mã GV, mã lớp và mã môn học/ mô đun mới sẽ được tạo lại từ đầu; phiên quản trị hiện tại vẫn được giữ.", "success")
            return redirect(url_for("profile_setup"))
        except Exception as exc:
            db.session.rollback()
            flash(f"Không thể xóa toàn bộ dữ liệu: {exc}", "danger")
            return redirect(url_for("index"))

    @app.route("/admin/data", methods=["GET", "POST"])
    @login_required
    def data_management():
        user = current_user()
        if user.role != "admin":
            abort(403)

        if request.method == "POST":
            upload = request.files.get("backup_file")
            confirm = (request.form.get("confirm") or "").strip()
            if not upload or not upload.filename:
                flash("Vui lòng chọn tệp sao lưu ZIP do chính hệ thống tạo.", "warning")
                return redirect(url_for("data_management"))
            if confirm != "KHÔI PHỤC":
                flash("Vui lòng nhập chính xác KHÔI PHỤC để xác nhận thao tác.", "warning")
                return redirect(url_for("data_management"))
            try:
                admin_email = user.email
                restore_from_zip(upload.stream)
                # Sau restore, user hiện tại có thể đã đổi ID hoặc không còn trong bộ dữ liệu khôi phục.
                restored_admin = User.query.filter_by(email=admin_email).first()
                record_functional_test(
                    "RESTORE", "Khôi phục dữ liệu từ bản sao lưu", True,
                    owner_id=(restored_admin.id if restored_admin else None), note=upload.filename or ""
                )
                db.session.commit()
                session.clear()
                flash("Khôi phục dữ liệu thành công. Vui lòng đăng nhập lại.", "success")
                return redirect(url_for("index"))
            except Exception as exc:
                db.session.rollback()
                db.session.add(ErrorLog(module="restore", action="Khôi phục dữ liệu", course_module="Hệ thống", actor_user_id=user.id, message="Khôi phục bản sao lưu thất bại", details=str(exc)))
                record_functional_test("RESTORE", "Khôi phục dữ liệu từ bản sao lưu", False, owner_id=user.id, note=str(exc))
                db.session.commit()
                flash(f"Khôi phục thất bại: {exc}", "danger")
                return redirect(url_for("data_management"))

        summary = {
            "videos": db.session.query(Video.id).join(YouTubeChannel).filter(
                Video.uploader_channel_id == YouTubeChannel.channel_id).count(),
            "courses": Course.query.filter_by(active=True).count(),
            "qrs": QrCode.query.filter_by(active=True).count(),
            "scans": ScanLog.query.count(),
        }
        return render_template("data_management.html", summary=summary)

    @app.get("/admin/backup.zip")
    @login_required
    def backup_zip():
        user = current_user()
        if user.role != "admin":
            abort(403)
        try:
            bio = build_backup_zip()
            record_functional_test("BACKUP", "Tạo bản sao lưu ZIP", True, owner_id=user.id)
            db.session.commit()
            return send_file(bio, mimetype="application/zip", as_attachment=True, download_name=f"sao_luu_ma_qr_{utcnow().date().isoformat()}.zip")
        except Exception as exc:
            db.session.rollback()
            record_functional_test("BACKUP", "Tạo bản sao lưu ZIP", False, owner_id=user.id, note=str(exc))
            db.session.commit()
            raise

    @app.route("/admin/restore", methods=["GET", "POST"])
    @login_required
    def restore_backup():
        # Giữ tương thích với liên kết cũ; giao diện quản trị mới dùng /admin/data.
        if current_user().role != "admin":
            abort(403)
        if request.method == "POST":
            flash("Giao diện khôi phục đã được chuyển sang mục Sao lưu & khôi phục.", "info")
        return redirect(url_for("data_management"))

    @app.get("/errors")
    @login_required
    def error_logs_page():
        if current_user().role != "admin":
            abort(403)
        logs = ErrorLog.query.order_by(ErrorLog.created_at.desc()).limit(300).all()
        action_fallback = {
            "oauth": "Đăng nhập hệ thống",
            "youtube": "Đồng bộ kênh/video",
            "websub": "Đồng bộ video",
            "reconcile": "Đồng bộ video",
            "restore": "Khôi phục dữ liệu",
            "backup": "Sao lưu dữ liệu",
            "qr": "Tạo/điều chỉnh mã QR",
            "scan": "Quét mã QR",
            "viewer": "Phát video",
        }
        student_codes = {e.actor_student_code for e in logs if e.actor_student_code}
        class_codes = {e.actor_class_code for e in logs if e.actor_class_code}
        student_by_code = {s.student_code: s for s in Student.query.filter(Student.student_code.in_(student_codes)).all()} if student_codes else {}
        class_by_code = {c.code: c for c in SchoolClass.query.filter(SchoolClass.code.in_(class_codes)).all()} if class_codes else {}
        error_rows = []
        for e in logs:
            actor = "Không xác định (bản ghi cũ)"
            if e.actor_user_id:
                teacher = e.actor or db.session.get(User, e.actor_user_id)
                actor = f"{teacher.name} ({teacher.email})" if teacher else "Giảng viên (tài khoản đã xóa)"
            elif e.actor_student_code:
                st = student_by_code.get(e.actor_student_code)
                cl = class_by_code.get(e.actor_class_code)
                digits = re.sub(r"\D", "", e.actor_student_code)
                short = (digits or e.actor_student_code)[-3:]
                actor = f"{st.full_name if st else 'Sinh viên'} ({short} – {cl.name if cl else (e.actor_class_code or 'Chưa rõ lớp')})"
            elif (e.module or "").lower() in {"websub", "reconcile"}:
                actor = "Tác vụ đồng bộ tự động"
            elif (e.module or "").lower() == "oauth":
                actor = "Người đăng nhập chưa xác định"
            error_rows.append({
                "log": e, "actor_name": actor,
                "action": e.action or action_fallback.get((e.module or "").lower(), "Xử lý hệ thống"),
                "course_module": e.course_module or "Hệ thống",
            })
        return render_template("errors.html", error_rows=error_rows)

    @app.post("/tasks/reconcile")
    @csrf.exempt
    def task_reconcile():
        key = request.headers.get("X-Task-Key") or request.args.get("key")
        if not secrets.compare_digest(str(key or ""), str(current_app.config["RECONCILE_KEY"])):
            abort(403)
        total = 0
        results = []
        for user in User.query.all():
            if not user.oauth_account:
                continue
            try:
                count = reconcile_user(user)
                total += count
                results.append({"user": user.email, "videos": count, "ok": True})
            except Exception as exc:
                db.session.add(ErrorLog(module="reconcile", action="Đồng bộ video", course_module="Hệ thống", message=f"Đối soát lỗi: {user.email}", details=str(exc)))
                db.session.commit()
                results.append({"user": user.email, "ok": False, "error": str(exc)})
        return jsonify({"ok": True, "total_videos": total, "results": results})



BACKUP_VERSION = "4.4"

def backup_table_specs():
    """Thứ tự từ bảng cha đến bảng con, dùng chung cho sao lưu, khôi phục, xóa sạch."""
    return [
        ("users", User),
        ("oauth_accounts", OAuthAccount),
        ("connected_google_accounts", ConnectedGoogleAccount),
        ("youtube_channels", YouTubeChannel),
        ("playlists", Playlist),
        ("videos", Video),
        ("playlist_videos", PlaylistVideo),
        ("courses", Course),
        ("course_transfers", CourseTransfer),
        ("class_groups", ClassGroup),
        ("school_classes", SchoolClass),
        ("homeroom_changes", HomeroomChange),
        ("students", Student),
        ("class_enrollments", ClassEnrollment),
        ("course_classes", CourseClass),
        ("participants", Participant),
        ("qr_codes", QrCode),
        ("course_videos", CourseVideo),
        ("course_playlists", CoursePlaylist),
        ("qr_target_history", QrTargetHistory),
        ("scan_logs", ScanLog),
        ("watch_sessions", WatchSession),
        ("watch_events", WatchEvent),
        ("hub_subscriptions", HubSubscription),
        ("hub_events", HubEvent),
        ("error_logs", ErrorLog),
        ("experiment_runs", ExperimentRun),
        ("functional_test_runs", FunctionalTestRun),
        ("research_settings", ResearchSetting),
        ("survey_results", SurveyResult),
        ("edit_history", EditHistory),
    ]

def build_backup_zip():
    bio = io.BytesIO()
    with zipfile.ZipFile(bio, "w", zipfile.ZIP_DEFLATED) as zf:
        manifest = {
            "product": current_app.config.get("APP_NAME", "TẠO MÃ QR TỰ ĐỘNG TỪ VIDEO YOUTUBE"),
            "backup_version": BACKUP_VERSION,
            "created_at": utcnow().isoformat(),
            "warning": "OAuth token được lưu ở dạng mã hóa; không chia sẻ tệp backup công khai."
        }
        zf.writestr("manifest.json", json.dumps(manifest, ensure_ascii=False, indent=2))
        for name, model in backup_table_specs():
            data = [serialize_model(row) for row in model.query.all()]
            zf.writestr(f"{name}.json", json.dumps(data, ensure_ascii=False, indent=2, default=str))
    bio.seek(0)
    return bio

def restore_from_zip(fileobj):
    raw = fileobj.read()
    if len(raw) > 50 * 1024 * 1024:
        raise ValueError("Tệp backup vượt quá 50 MB.")
    with zipfile.ZipFile(io.BytesIO(raw), "r") as zf:
        names = set(zf.namelist())
        if "manifest.json" not in names:
            raise ValueError("Không có manifest.json; đây không phải backup hợp lệ.")
        manifest = json.loads(zf.read("manifest.json"))
        if manifest.get("backup_version") not in {"2.0", "2.1", "3.0", "4.3", BACKUP_VERSION}:
            raise ValueError("Phiên bản backup không tương thích.")
        payload = {}
        for name, model in backup_table_specs():
            fn = f"{name}.json"
            if fn not in names:
                # Các bảng bổ sung ở phiên bản sau là tùy chọn để backup cũ vẫn khôi phục được.
                if name in {"research_settings", "course_playlists", "school_classes", "students", "class_enrollments", "course_classes", "connected_google_accounts", "homeroom_changes", "edit_history", "course_transfers"}:
                    payload[name] = []
                    continue
                raise ValueError(f"Thiếu {fn} trong backup.")
            rows = json.loads(zf.read(fn))
            columns = {c.name: c for c in model.__table__.columns}
            cleaned = []
            for row in rows:
                item = {}
                for k, v in row.items():
                    col = columns.get(k)
                    if not col:
                        continue
                    if v is not None and isinstance(v, str):
                        if isinstance(col.type, DateTime):
                            v = datetime.fromisoformat(v.replace("Z", "+00:00"))
                        elif col.type.__class__.__name__ == "Date":
                            v = date.fromisoformat(v[:10])
                    item[k] = v
                cleaned.append(item)
            payload[name] = cleaned

    # Xóa theo thứ tự ngược để tôn trọng khóa ngoại, sau đó chèn lại theo thứ tự cha -> con.
    for _, model in reversed(backup_table_specs()):
        db.session.execute(model.__table__.delete())
    db.session.flush()
    for name, model in backup_table_specs():
        rows = payload[name]
        if rows:
            db.session.execute(model.__table__.insert(), rows)
    db.session.commit()
    migrate_legacy_classes()

    # Đồng bộ lại sequence PostgreSQL sau khi khôi phục ID tường minh.
    if db.engine.dialect.name == "postgresql":
        for _, model in backup_table_specs():
            table = model.__tablename__
            if "id" in model.__table__.columns:
                db.session.execute(text(
                    f"SELECT setval(pg_get_serial_sequence('{table}','id'), COALESCE((SELECT MAX(id) FROM {table}), 1), (SELECT MAX(id) IS NOT NULL FROM {table}))"
                ))
        db.session.commit()

def acceptance_metrics(user):
    """Tính chỉ tiêu nghiệm thu.

    Khi người dùng là quản trị viên, số liệu được tổng hợp trên toàn hệ thống để
    phản ánh thực nghiệm của nhiều giảng viên/bộ môn. Nhánh người dùng thường vẫn
    được giữ để hàm có thể tái sử dụng khi cần.
    """
    scan_q = ScanLog.query.join(QrCode, ScanLog.qr_id == QrCode.id)
    if user.role != "admin":
        scan_q = scan_q.filter(managed_qr_scope(user.id))
    scans_total = scan_q.count()
    scans_success = scan_q.filter(ScanLog.success.is_(True)).count()
    scan_rate = round(100 * scans_success / scans_total, 1) if scans_total else None

    tests_q = FunctionalTestRun.query.filter(FunctionalTestRun.source == "automatic")
    if user.role != "admin":
        tests_q = tests_q.filter(FunctionalTestRun.owner_id == user.id)
    tests_total = tests_q.count()
    tests_pass = tests_q.filter(FunctionalTestRun.passed.is_(True)).count()
    test_rate = round(100 * tests_pass / tests_total, 1) if tests_total else None

    login_q = tests_q.filter(FunctionalTestRun.case_code == "AUTH_LOGIN")
    login_total = login_q.count()
    login_pass = login_q.filter(FunctionalTestRun.passed.is_(True)).count()
    login_rate = round(100 * login_pass / login_total, 1) if login_total else None

    sync_q = tests_q.filter(FunctionalTestRun.case_code == "YT_SYNC")
    sync_total = sync_q.count()
    sync_pass = sync_q.filter(FunctionalTestRun.passed.is_(True)).count()
    sync_rate = round(100 * sync_pass / sync_total, 1) if sync_total else None

    exp_q = ExperimentRun.query if user.role == "admin" else ExperimentRun.query.filter_by(owner_id=user.id)
    comparison_keys = list(RESEARCH_COMPARISON_TASK_MAP.keys())
    manual = exp_q.filter(ExperimentRun.mode == "manual", ExperimentRun.success.is_(True), ExperimentRun.task_key.in_(comparison_keys)).all()
    webapp = exp_q.filter(ExperimentRun.mode == "webapp", ExperimentRun.success.is_(True), ExperimentRun.task_key.in_(comparison_keys)).all()
    # Chỉ so sánh các nhiệm vụ có dữ liệu ở CẢ hai phía để tránh trộn các công việc khác nhau.
    matched_keys = {x.task_key for x in manual} & {x.task_key for x in webapp}
    manual_matched = [x for x in manual if x.task_key in matched_keys]
    webapp_matched = [x for x in webapp if x.task_key in matched_keys]
    manual_ops = sum(x.operations_count for x in manual_matched) / len(manual_matched) if manual_matched else None
    app_ops = sum(x.operations_count for x in webapp_matched) / len(webapp_matched) if webapp_matched else None
    manual_time = sum(x.elapsed_seconds for x in manual_matched) / len(manual_matched) if manual_matched else None
    app_time = sum(x.elapsed_seconds for x in webapp_matched) / len(webapp_matched) if webapp_matched else None

    qr_perf_q = QrCode.query.filter(QrCode.processing_seconds.isnot(None))
    if user.role != "admin":
        qr_perf_q = qr_perf_q.filter(managed_qr_scope(user.id))
    qr_perf = [q.processing_seconds for q in qr_perf_q.all() if q.processing_seconds is not None]
    processing_values = qr_perf or [x.system_processing_seconds for x in webapp if x.system_processing_seconds is not None]
    avg_processing = sum(processing_values) / len(processing_values) if processing_values else None

    reduction_ops = round(100 * (manual_ops - app_ops) / manual_ops, 1) if manual_ops and app_ops is not None else None
    reduction_time = round(100 * (manual_time - app_time) / manual_time, 1) if manual_time and app_time is not None else None

    # Khảo sát được thu bằng Google Forms và thống kê trong Google Sheets để tránh nhập tay từng phiếu.
    # Web App chỉ liên kết tới nguồn khảo sát; hai chỉ tiêu hài lòng được đánh giá từ Sheets khi nghiệm thu.
    lecturer_sat = None
    student_sat = None

    participant_base = db.session.query(ScanLog.participant_code).join(QrCode, ScanLog.qr_id == QrCode.id)
    if user.role != "admin":
        participant_base = participant_base.filter(managed_qr_scope(user.id))
    # Chỉ tính GV đã thực sự tạo môn học/QR hoặc gán học liệu, không tính đăng nhập
    # quản trị thuần túy. Mỗi GV được tính tối đa một lần kể cả kiêm Admin/GVCN.
    # Sinh viên tham gia được tính theo mã SV thực tế đã quét QR.
    if user.role == "admin":
        teacher_ids = {x[0] for x in db.session.query(Course.owner_id).all()}
        teacher_ids.update(x[0] for x in db.session.query(QrCode.owner_id).all())
        participating_lecturers_q = User.query.filter(User.profile_completed.is_(True), User.id.in_(teacher_ids))
        gv_count = participating_lecturers_q.count()
        gv_non_it_count = participating_lecturers_q.filter(User.professional_group == "other").count()
    else:
        gv_count = 1
        gv_non_it_count = 1 if user.professional_group == "other" else 0
    sv_count = participant_base.filter(ScanLog.role == "SV").distinct().count()

    video_q = Video.query.join(YouTubeChannel).filter(
        Video.uploader_channel_id == YouTubeChannel.channel_id)
    if user.role != "admin":
        video_q = video_q.filter(YouTubeChannel.user_id == user.id)

    course_q = Course.query.filter_by(active=True) if user.role == "admin" else Course.query.filter_by(owner_id=user.id, active=True)
    course_count = course_q.count()

    # Chỉ tiêu nghiệm thu yêu cầu số bộ môn, vì vậy đếm tên bộ môn duy nhất.
    # Dữ liệu cũ chưa có department_name được phép dùng unit_name làm giá trị dự phòng.
    department_rows = course_q.with_entities(Course.department_name, Course.unit_name).all()
    department_keys = set()
    for department, legacy_unit in department_rows:
        department = (department or legacy_unit or "").strip()
        if department:
            department_keys.add(department.casefold())
    department_count = len(department_keys)

    metrics = [
        metric("Đăng nhập vào hệ thống", login_rate, "=", 100, "%"),
        metric("Đồng bộ video thành công", sync_rate, "=", 100, "%"),
        metric("Tỷ lệ QR quét thành công", scan_rate, ">=", 98, "%"),
        metric("Thời gian xử lý hệ thống trung bình/video", round(avg_processing, 2) if avg_processing is not None else None, "<=", 5, " giây"),
        metric("Giảm số thao tác so với thủ công", reduction_ops, ">=", 30, "%"),
        metric("Giảm tổng thời gian so với thủ công", reduction_time, ">=", 30, "%"),
        metric("Tỷ lệ ca kiểm thử chức năng đạt", test_rate, ">=", 95, "%"),
        metric("Số môn học/ mô đun thực nghiệm", course_count, ">=", 3, ""),
        metric("Số bộ môn", department_count, ">=", 2, ""),
        metric("Số sinh viên tham gia", sv_count, ">=", 50, ""),
        metric("Số video/tình huống kiểm thử", video_q.count(), ">=", 100, ""),
        metric("Mức hài lòng giảng viên", lecturer_sat, ">=", 80, "%", external=True),
        metric("Mức hài lòng sinh viên", student_sat, ">=", 80, "%", external=True),
    ]
    gv_metric = metric("Số giảng viên tham gia", gv_count, ">=", 3, "")
    gv_metric["target_display"] = ">= 3 (GV CN khác CNTT: >= 1)"
    gv_metric["value_display"] = f"{gv_count} ({gv_non_it_count})"
    gv_metric["passed"] = gv_count >= 3 and gv_non_it_count >= 1
    metrics.insert(9, gv_metric)
    return metrics


def metric(name, value, op, target, suffix, external=False):
    passed = None
    if value is not None:
        if op == ">=":
            passed = value >= target
        elif op == "<=":
            passed = value <= target
        elif op == "=":
            passed = value == target
        else:
            passed = None
    return {"name": name, "value": value, "op": op, "target": target, "suffix": suffix, "passed": passed, "external": external}


def serialize_model(obj):
    out = {}
    for col in obj.__table__.columns:
        value = getattr(obj, col.name)
        if hasattr(value, "isoformat"):
            value = value.isoformat()
        out[col.name] = value
    return out
